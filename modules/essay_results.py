# modules/essay_results.py
# EMIS — Essay / Theory Results Extension

import csv
import json
import re
import time
from pathlib import Path
from datetime import datetime
from threading import Lock

from flask import jsonify, request, session
from openpyxl import load_workbook

from modules.class_config import SUPPORTED_CLASSES, normalize_class_level, normalize_class_arm
from modules.excel_manager import read_results, get_preferred_excel_path, normalize_result_term, result_term_label
from modules.result_sync import queue_emis_event_safely


BASE_DIR = Path(__file__).resolve().parent.parent
RESULTS_DIR = BASE_DIR / "RESULTS"
DATABASE_DIR = BASE_DIR / "static" / "data" / "database"
SUBJECTS_ROOT = BASE_DIR / "static" / "subjects"
MASTER_CSV = DATABASE_DIR / "students2026.csv"

CLASS_DATABASES = {
    "JSS1": DATABASE_DIR / "JSS1_Students.xlsx", "JSS2": DATABASE_DIR / "JSS2_Students.xlsx", "JSS3": DATABASE_DIR / "JSS3_Students.xlsx",
    "SS1": DATABASE_DIR / "SS1_Students.xlsx", "SS2": DATABASE_DIR / "SS2_Students.xlsx", "SS3": DATABASE_DIR / "SS3_Students.xlsx",
}

VALID_TERMS = {"FIRST", "SECOND", "THIRD"}

DEFAULT_OBJECTIVE_MAX = 60.0
DEFAULT_ESSAY_MAX = 40.0
DEFAULT_TOTAL_MAX = 100.0

ESSAY_FILE_NAME = "essay_scores.json"
ESSAY_LOCK = Lock()

STAFF_SESSION_TIMEOUT_SECONDS = 20 * 60
STAFF_SESSION_ACTIVITY_KEY = "_staff_last_activity"


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):
    return str(value or "").strip()


def normalize_class(value):
    return normalize_class_level(value)


def is_jss_class(value):
    return clean_text(normalize_class(value)).upper().startswith("JSS")


def is_ss_class(value):
    return clean_text(normalize_class(value)).upper().startswith("SS")


def normalize_term(value):
    return normalize_result_term(value)


def term_label(value):
    return {"FIRST": "1st Term", "SECOND": "2nd Term", "THIRD": "3rd Term"}.get(normalize_term(value), "")


def score_float(value, default=None):
    if value is None or str(value).strip() == "": return default

    try: return round(float(str(value).replace("%", "").strip()), 2)
    except (TypeError, ValueError): return default


def score_output(value):
    value = score_float(value)

    if value is None: return None
    return int(value) if float(value).is_integer() else value


def normalize_admission(value):
    return clean_text(value).upper()


def normalize_name(value):
    return " ".join(clean_text(value).upper().split())


def normalize_subject_key(value):
    value = clean_text(value).lower()
    value = re.sub(r"\.json$", "", value)
    value = re.sub(r"_(jss|ss)[123]$", "", value)
    value = re.sub(r"[^a-z0-9]+", "", value)
    return value


def subject_folder_name(value):
    value = clean_text(value).lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")
    return value or "subject"


def now_string():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ============================================================
# STAFF SESSION
# ============================================================

def can_manage_results():
    role = clean_text(session.get("user_type")).lower()

    if role not in {"admin", "teacher"}: return False

    now = time.time()
    last_activity = session.get(STAFF_SESSION_ACTIVITY_KEY)

    if last_activity is not None:
        try:
            if now - float(last_activity) > STAFF_SESSION_TIMEOUT_SECONDS:
                print(f"[ESSAY SESSION] {role} session expired after 20 minutes inactivity.")
                session.clear()
                return False
        except (TypeError, ValueError):
            pass

    session.permanent = True
    session[STAFF_SESSION_ACTIVITY_KEY] = now
    session.modified = True

    return True


# ============================================================
# STUDENT HELPERS
# ============================================================

def build_student_name(row):
    direct = clean_text(row.get("Student Name") or row.get("student_name") or row.get("full_name") or row.get("Name"))

    if direct: return direct

    last_name = clean_text(row.get("Last_name") or row.get("Last Name") or row.get("last_name"))
    first_name = clean_text(row.get("First_name") or row.get("First Name") or row.get("first_name"))
    other_names = clean_text(row.get("Other_names") or row.get("Other Names") or row.get("other_names"))

    return " ".join(part for part in (last_name, first_name, other_names) if part)


def normalize_student_row(row, class_level=""):
    class_level = normalize_class(class_level or row.get("Class_category") or row.get("Class Category") or row.get("Class Level") or row.get("Class"))

    raw_arm = row.get("Class") or row.get("Class_arm") or row.get("Class Arm") or class_level
    class_arm = normalize_class_arm(raw_arm, class_level) or clean_text(raw_arm).upper() or class_level

    return {
        "admission_number": clean_text(row.get("Admission_number") or row.get("Admission No") or row.get("admission_number") or row.get("student_id")),
        "student_name": build_student_name(row),
        "class_level": class_level,
        "class_category": class_level,
        "class_arm": class_arm,
        "sex": clean_text(row.get("Sex") or row.get("sex")),
    }


# ============================================================
# CLASS DATABASE
# ============================================================

def read_class_xlsx(class_level):
    class_level = normalize_class(class_level)
    path = CLASS_DATABASES.get(class_level)

    if not path or not path.exists(): return []

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) < 2:
            wb.close()
            return []

        headers = [clean_text(value) for value in rows[0]]
        students = []

        for values in rows[1:]:
            values = list(values) + [None] * max(0, len(headers) - len(values))
            student = normalize_student_row(dict(zip(headers, values[:len(headers)])), class_level)

            if student["admission_number"] or student["student_name"]: students.append(student)

        wb.close()
        return students

    except Exception as error:
        print(f"ESSAY STUDENT XLSX READ ERROR [{path}]:", error)
        return []


def read_master_csv(class_level):
    class_level = normalize_class(class_level)

    if not MASTER_CSV.exists(): return []

    try:
        students = []

        with MASTER_CSV.open("r", encoding="utf-8-sig", newline="") as file:
            for row in csv.DictReader(file):
                row_class = normalize_class(row.get("Class_category") or row.get("Class") or "")

                if row_class != class_level: continue

                student = normalize_student_row(row, class_level)

                if student["admission_number"] or student["student_name"]: students.append(student)

        return students

    except Exception as error:
        print("ESSAY MASTER CSV READ ERROR:", error)
        return []


def get_class_roster(class_level, arm=""):
    class_level = normalize_class(class_level)
    students = read_class_xlsx(class_level) or read_master_csv(class_level)

    unique = {}

    for student in students:
        key = normalize_admission(student.get("admission_number")) or normalize_name(student.get("student_name"))

        if key: unique[key] = student

    students = list(unique.values())

    if arm and clean_text(arm).lower() != "all":
        normalized_arm = normalize_class_arm(arm, class_level) or clean_text(arm).upper()
        students = [student for student in students if clean_text(student.get("class_arm")).upper() == clean_text(normalized_arm).upper()]

    students.sort(key=lambda student: (clean_text(student.get("student_name")).lower(), clean_text(student.get("admission_number")).lower()))

    return students


# ============================================================
# RESULT DIRECTORY
# ============================================================

def get_subject_result_dir(year, class_level, subject, term=""):
    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return None

    raw_term = clean_text(term)
    term = normalize_term(raw_term)

    if is_jss_class(class_level) and not term: return None
    if raw_term and not term: return None

    try:
        excel_path = get_preferred_excel_path(class_level, subject, str(year), term)
        return excel_path.parent
    except Exception:
        root = RESULTS_DIR / str(year) / "CLASS" / class_level

        if term: root = root / term

        return root / subject_folder_name(subject)


def find_subject_result_dir(year, class_level, subject, term="", create=False):
    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return None

    raw_term = clean_text(term)
    term = normalize_term(raw_term)

    if is_jss_class(class_level) and not term: return None
    if raw_term and not term: return None

    root = RESULTS_DIR / str(year) / "CLASS" / class_level

    if term: root = root / term

    requested_key = normalize_subject_key(subject)

    if root.exists():
        for folder in root.iterdir():
            if folder.is_dir() and normalize_subject_key(folder.name) == requested_key:
                return folder

    preferred = get_subject_result_dir(year, class_level, subject, term)

    if preferred and create: preferred.mkdir(parents=True, exist_ok=True)

    return preferred


# ============================================================
# ESSAY STORE
# ============================================================

def get_essay_score_path(year, class_level, subject, term="", create=False):
    folder = find_subject_result_dir(year, class_level, subject, term, create=create)
    return folder / ESSAY_FILE_NAME if folder else None


def empty_essay_store(year, class_level, subject, term=""):
    class_level = normalize_class(class_level)
    term = normalize_term(term)

    return {
        "version": 1,
        "year": str(year),
        "class_level": class_level,
        "class_category": class_level,
        "term": term,
        "term_label": term_label(term),
        "subject": clean_text(subject),
        "objective_max": DEFAULT_OBJECTIVE_MAX,
        "essay_max": DEFAULT_ESSAY_MAX,
        "total_max": DEFAULT_TOTAL_MAX,
        "scores": {},
    }


def read_essay_store(year, class_level, subject, term=""):
    class_level = normalize_class(class_level)
    term = normalize_term(term)
    path = get_essay_score_path(year, class_level, subject, term)

    if not path or not path.exists(): return empty_essay_store(year, class_level, subject, term)

    try:
        with path.open("r", encoding="utf-8-sig") as file:
            data = json.load(file)

        if not isinstance(data, dict): data = empty_essay_store(year, class_level, subject, term)
        if not isinstance(data.get("scores"), dict): data["scores"] = {}

        # Physical path/request is authoritative over stale metadata.
        data["version"] = data.get("version") or 1
        data["year"] = str(year)
        data["class_level"] = class_level
        data["class_category"] = class_level
        data["term"] = term
        data["term_label"] = term_label(term)
        data["subject"] = clean_text(data.get("subject") or subject)

        data.setdefault("objective_max", DEFAULT_OBJECTIVE_MAX)
        data.setdefault("essay_max", DEFAULT_ESSAY_MAX)
        data.setdefault("total_max", DEFAULT_TOTAL_MAX)

        return data

    except Exception as error:
        print(f"ESSAY STORE READ ERROR [{path}]:", error)
        return empty_essay_store(year, class_level, subject, term)


def write_essay_store(year, class_level, subject, term, data):
    path = get_essay_score_path(year, class_level, subject, term, create=True)

    if not path: raise ValueError("Could not resolve essay score path")

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")

    with ESSAY_LOCK:
        with temporary.open("w", encoding="utf-8") as file:
            json.dump(data, file, indent=2, ensure_ascii=False)

        temporary.replace(path)

    return path


# ============================================================
# EXAM JSON
# ============================================================

def get_subject_json_directory(year, class_level, term=""):
    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return None

    directory = SUBJECTS_ROOT / str(year) / "subjects-json" / class_level
    normalized_term = normalize_term(term)

    if is_jss_class(class_level):
        if not normalized_term: return None
        return directory / normalized_term

    if is_ss_class(class_level) and normalized_term:
        return directory / normalized_term

    return directory


def get_subject_json_directories(year, class_level, term=""):
    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return []

    root = SUBJECTS_ROOT / str(year) / "subjects-json" / class_level
    normalized_term = normalize_term(term)

    if is_jss_class(class_level):
        return [root / normalized_term] if normalized_term else []

    if is_ss_class(class_level):
        # Hybrid SS: selected term first, then root/general fallback.
        return [root / normalized_term, root] if normalized_term else [root]

    return [root]


def find_subject_json(year, class_level, subject, term=""):
    requested_key = normalize_subject_key(subject)
    seen = set()

    for directory in get_subject_json_directories(year, class_level, term):
        key = str(directory)

        if key in seen:
            continue

        seen.add(key)

        if not directory.exists():
            continue

        for path in sorted(directory.glob("*.json")):
            if path.name.lower() == "pushed_subjects.json": continue

            try:
                with path.open("r", encoding="utf-8-sig") as file:
                    data = json.load(file)

                if not isinstance(data, dict): continue

                subject_name = data.get("subject") or path.stem

                if normalize_subject_key(path.stem) == requested_key or normalize_subject_key(subject_name) == requested_key:
                    return path, data

            except Exception as error:
                print(f"ESSAY SUBJECT JSON READ ERROR [{path}]:", error)

    return None, {}


# ============================================================
# ESSAY CONFIG
# ============================================================

def get_exam_essay_config(year, class_level, subject, term=""):
    path, data = find_subject_json(year, class_level, subject, term)
    essay = data.get("essay") if isinstance(data, dict) and isinstance(data.get("essay"), dict) else {}
    questions = essay.get("questions", []) if isinstance(essay.get("questions", []), list) else []
    essay_defined = len(questions) > 0

    # Manual Essay/Theory scoring belongs to Admin Results.
    # It must remain available even when the online CBT JSON has no essay block.
    essay_max = score_float(essay.get("max_score"), DEFAULT_ESSAY_MAX)
    if essay_max is None or essay_max <= 0 or essay_max >= DEFAULT_TOTAL_MAX: essay_max = DEFAULT_ESSAY_MAX

    objective_max = round(DEFAULT_TOTAL_MAX - essay_max, 2)

    return {
        "essay_available": True, "essay_defined": essay_defined, "essay_questions": len(questions),
        "essay_title": clean_text(essay.get("title")) or "Manual Essay / Theory",
        "objective_max": objective_max, "essay_max": essay_max, "total_max": DEFAULT_TOTAL_MAX,
        "json_path": str(path) if path else "", "term": normalize_term(term), "term_label": term_label(term),
    }

# ============================================================
# OBJECTIVE RESULTS
# ============================================================

def read_objective_results(year, class_level, subject, term=""):
    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return []

    try:
        return read_results(class_level, subject, str(year), normalize_term(term))
    except Exception as error:
        print(f"ESSAY OBJECTIVE READ ERROR [{year}/{class_level}/{term or 'GENERAL'}/{subject}]:", error)
        return []


def build_objective_lookup(records):
    lookup = {}

    for row in records:
        admission = normalize_admission(row.get("Admission No") or row.get("Admission_number") or row.get("admission_number"))
        name = normalize_name(row.get("Student Name") or row.get("student_name") or row.get("full_name"))

        if admission: lookup[f"ADM::{admission}"] = row
        if name: lookup[f"NAME::{name}"] = row

    return lookup


def find_objective_record(lookup, admission="", student_name=""):
    admission = normalize_admission(admission)
    student_name = normalize_name(student_name)

    if admission and f"ADM::{admission}" in lookup: return lookup[f"ADM::{admission}"]
    if student_name and f"NAME::{student_name}" in lookup: return lookup[f"NAME::{student_name}"]

    return None


# ============================================================
# OBJECTIVE COMPONENT
# ============================================================

def calculate_objective_component(row, objective_max):
    if not isinstance(row, dict): return None, None, None

    correct = score_float(row.get("Correct") or row.get("correct"))
    total_questions = score_float(row.get("Total") or row.get("total") or row.get("total_questions"))

    if correct is not None and total_questions is not None and total_questions > 0:
        return round((correct / total_questions) * float(objective_max), 2), correct, total_questions

    percentage = score_float(row.get("Score (%)") or row.get("Score Number") or row.get("score_percentage") or row.get("percentage"))

    if percentage is not None:
        return round((percentage / 100.0) * float(objective_max), 2), correct, total_questions

    return None, correct, total_questions


# ============================================================
# SAVED ESSAY SCORE
# ============================================================

def find_saved_score(store, admission="", student_name=""):
    scores = store.get("scores", {}) if isinstance(store, dict) else {}

    if not isinstance(scores, dict): return None, None

    admission_key = normalize_admission(admission)
    name_key = normalize_name(student_name)

    if admission_key and admission_key in scores:
        entry = scores.get(admission_key)
        return score_float(entry.get("score")) if isinstance(entry, dict) else score_float(entry), entry

    for _, entry in scores.items():
        if not isinstance(entry, dict): continue

        entry_admission = normalize_admission(entry.get("admission_number"))
        entry_name = normalize_name(entry.get("student_name"))

        if admission_key and entry_admission == admission_key: return score_float(entry.get("score")), entry
        if not admission_key and name_key and entry_name == name_key: return score_float(entry.get("score")), entry

    return None, None


# ============================================================
# COMBINED SCORE
# ============================================================

def build_score_components(objective_row, essay_score, config):
    essay_available = bool(config.get("essay_available"))
    objective_max = score_float(config.get("objective_max"), DEFAULT_OBJECTIVE_MAX)
    essay_max = score_float(config.get("essay_max"), DEFAULT_ESSAY_MAX if essay_available else 0.0)

    objective_score, correct, total_questions = calculate_objective_component(objective_row, objective_max)

    has_objective = objective_score is not None
    has_essay = essay_score is not None

    if not essay_available:
        final_score = objective_score
        result_state = "COMPLETE" if has_objective else "PENDING"

    elif has_objective and has_essay:
        final_score = round(objective_score + essay_score, 2)
        result_state = "COMPLETE"

    elif has_objective:
        final_score = None
        result_state = "AWAITING ESSAY"

    elif has_essay:
        final_score = None
        result_state = "AWAITING OBJECTIVE"

    else:
        final_score = None
        result_state = "PENDING"

    final_status = "PASS" if result_state == "COMPLETE" and final_score is not None and final_score >= 50 else "FAIL" if result_state == "COMPLETE" and final_score is not None else ""

    return {
        "has_objective": has_objective,
        "objective_score": score_output(objective_score),
        "objective_max": score_output(objective_max),
        "objective_correct": score_output(correct),
        "objective_total": score_output(total_questions),

        "essay_available": essay_available,
        "has_essay": has_essay,
        "essay_score": score_output(essay_score),
        "essay_max": score_output(essay_max),

        "final_score": score_output(final_score),
        "combined_score": score_output(final_score),
        "total_score": score_output(final_score),
        "total_max": 100,

        "result_state": result_state,
        "final_status": final_status,
    }


# ============================================================
# ENRICH ADMIN RESULTS
# ============================================================

def enrich_results_with_essay(records, class_level, subject, year, term=""):
    if not isinstance(records, list): return records

    class_level = normalize_class(class_level)

    if class_level not in SUPPORTED_CLASSES: return records

    requested_term = normalize_term(term)
    cache = {}

    for row in records:
        if not isinstance(row, dict): continue

        # Explicit API/folder term wins. Otherwise use the row's own term.
        row_term = requested_term or normalize_term(row.get("Term") or row.get("term"))

        # JSS should always have a term at this stage. Old malformed rows
        # without one simply cannot resolve a theory store safely.
        if is_jss_class(class_level) and not row_term:
            continue

        if row_term not in cache:
            cache[row_term] = (
                get_exam_essay_config(year, class_level, subject, row_term),
                read_essay_store(year, class_level, subject, row_term),
            )

        config, store = cache[row_term]

        admission = row.get("Admission No") or row.get("Admission_number") or row.get("admission_number")
        student_name = row.get("Student Name") or row.get("student_name") or row.get("full_name")

        essay_score, _ = find_saved_score(store, admission, student_name)
        components = build_score_components(row, essay_score, config)

        row.update(components)

        if row_term:
            row["Term"] = row_term
            row["Term Label"] = result_term_label(row_term)

        row["Essay Available"] = components["essay_available"]
        row["Essay Score"] = components["essay_score"] if components["essay_score"] is not None else ""
        row["Essay Max"] = components["essay_max"]

        row["Objective Score"] = components["objective_score"] if components["objective_score"] is not None else ""
        row["Objective Max"] = components["objective_max"]

        row["Final Score"] = components["final_score"] if components["final_score"] is not None else ""
        row["Final Max"] = 100
        row["Result State"] = components["result_state"]
        row["Final Status"] = components["final_status"]

    return records


# ============================================================
# ESSAY ROSTER
# ============================================================

def build_essay_roster(year, class_level, subject, term="", arm=""):
    class_level = normalize_class(class_level)
    term = normalize_term(term)

    config = get_exam_essay_config(year, class_level, subject, term)
    store = read_essay_store(year, class_level, subject, term)

    objective_records = read_objective_results(year, class_level, subject, term)
    objective_lookup = build_objective_lookup(objective_records)

    students = get_class_roster(class_level, arm)
    roster = []

    for student in students:
        admission = student.get("admission_number", "")
        student_name = student.get("student_name", "")

        objective = find_objective_record(objective_lookup, admission, student_name)
        essay_score, saved_entry = find_saved_score(store, admission, student_name)

        components = build_score_components(objective, essay_score, config)

        roster.append({
            **student,
            **components,
            "term": term,
            "term_label": term_label(term),
            "essay_saved": saved_entry is not None,
            "essay_updated_at": clean_text(saved_entry.get("updated_at")) if isinstance(saved_entry, dict) else "",
        })

    return roster, config




# ============================================================
# SYNC HELPERS
# ============================================================

def essay_sync_entity_key(year, class_level, subject, term=""):
    return "|".join([clean_text(year), clean_text(class_level).upper(), normalize_term(term), normalize_subject_key(subject)])


def queue_essay_store_sync(year, class_level, subject, term, store):
    payload = {
        "year": clean_text(year), "class_level": normalize_class(class_level), "subject": clean_text(subject),
        "term": normalize_term(term), "store": dict(store or {}),
    }
    return queue_emis_event_safely(
        "essay", "replace_store", payload,
        entity_key=essay_sync_entity_key(payload["year"], payload["class_level"], payload["subject"], payload["term"]),
    )


def apply_essay_sync_event(action, payload, event=None):
    if clean_text(action).lower() != "replace_store": raise ValueError(f"Unsupported essay sync action: {action}")
    if not isinstance(payload, dict): raise ValueError("Essay sync payload must be an object.")

    year = clean_text(payload.get("year"))
    class_level = normalize_class(payload.get("class_level") or payload.get("class"))
    subject = clean_text(payload.get("subject"))
    raw_term = clean_text(payload.get("term")); term = normalize_term(raw_term)
    store = payload.get("store")

    if not year or not year.isdigit(): raise ValueError("Essay sync requires a valid result year.")
    if class_level not in SUPPORTED_CLASSES: raise ValueError("Essay sync requires a valid class.")
    if not subject: raise ValueError("Essay sync requires a subject.")
    if raw_term and not term: raise ValueError("Essay sync received an invalid term.")
    if is_jss_class(class_level) and not term: raise ValueError("Essay sync requires a term for JSS.")
    if not isinstance(store, dict): raise ValueError("Essay synchronized store must be an object.")

    normalized_store = dict(store)
    normalized_store["year"] = year
    normalized_store["class_level"] = class_level
    normalized_store["class_category"] = class_level
    normalized_store["term"] = term
    normalized_store["term_label"] = term_label(term)
    normalized_store["subject"] = clean_text(normalized_store.get("subject") or subject)
    if not isinstance(normalized_store.get("scores"), dict): normalized_store["scores"] = {}
    path = write_essay_store(year, class_level, subject, term, normalized_store)
    return {"year": year, "class_level": class_level, "subject": subject, "term": term, "score_count": len(normalized_store["scores"]), "path": str(path)}


# ============================================================
# ROUTES
# ============================================================

def register_essay_routes(api_bp):

    @api_bp.route("/api/results/essay", methods=["GET"])
    def get_essay_scores():
        if not can_manage_results(): return jsonify({"error": "Unauthorized", "students": []}), 403

        year = clean_text(request.args.get("year"))
        class_level = normalize_class(request.args.get("class") or request.args.get("class_level"))
        subject = clean_text(request.args.get("subject"))
        raw_term = clean_text(request.args.get("term"))
        term = normalize_term(raw_term)
        arm = clean_text(request.args.get("arm"))

        if not year or not year.isdigit(): return jsonify({"error": "Invalid or missing year", "students": []}), 400
        if class_level not in SUPPORTED_CLASSES: return jsonify({"error": "Invalid or missing class", "students": []}), 400
        if not subject: return jsonify({"error": "Subject is required", "students": []}), 400
        if raw_term and raw_term.lower() != "all" and not term: return jsonify({"error": "Invalid academic term", "students": []}), 400
        if is_jss_class(class_level) and not term: return jsonify({"error": "Term is required for JSS essay scores", "students": []}), 400

        try:
            students, config = build_essay_roster(year, class_level, subject, term, arm)

            return jsonify({
                "success": True,

                "year": year,
                "class": class_level,
                "class_level": class_level,
                "arm": arm,

                "term": term,
                "term_label": term_label(term),

                "supports_term": True,
                "requires_term": is_jss_class(class_level),
                "term_mode": "required" if is_jss_class(class_level) else "term" if term else "general",

                "subject": subject,

                "essay_available": config["essay_available"],
                "essay_questions": config["essay_questions"],
                "essay_title": config["essay_title"],

                "objective_max": score_output(config["objective_max"]),
                "essay_max": score_output(config["essay_max"]),
                "total_max": 100,

                "student_count": len(students),
                "students": students,
            }), 200

        except Exception as error:
            print("GET ESSAY SCORES ERROR:", error)
            return jsonify({"error": "Failed to load essay scores", "details": str(error), "students": []}), 500


    @api_bp.route("/api/results/essay/save", methods=["POST"])
    def save_essay_scores():
        if not can_manage_results(): return jsonify({"error": "Unauthorized"}), 403

        data = request.get_json(silent=True) or {}

        year = clean_text(data.get("year"))
        class_level = normalize_class(data.get("class") or data.get("class_level"))
        subject = clean_text(data.get("subject"))

        raw_term = clean_text(data.get("term"))
        term = normalize_term(raw_term)

        scores = data.get("scores", [])

        if not year or not year.isdigit(): return jsonify({"error": "Invalid or missing year"}), 400
        if class_level not in SUPPORTED_CLASSES: return jsonify({"error": "Invalid or missing class"}), 400
        if not subject: return jsonify({"error": "Subject is required"}), 400
        if raw_term and raw_term.lower() != "all" and not term: return jsonify({"error": "Invalid academic term"}), 400
        if is_jss_class(class_level) and not term: return jsonify({"error": "Term is required for JSS essay scores"}), 400
        if not isinstance(scores, list) or not scores: return jsonify({"error": "No essay scores supplied"}), 400

        config = get_exam_essay_config(year, class_level, subject, term)

        if not config["essay_available"]:
            return jsonify({"error": "This examination does not contain an essay / theory section"}), 400

        essay_max = score_float(config.get("essay_max"), DEFAULT_ESSAY_MAX)

        if essay_max is None or essay_max <= 0:
            return jsonify({"error": "Invalid essay maximum score"}), 400

        roster = get_class_roster(class_level)
        valid_students = {}

        for student in roster:
            admission_key = normalize_admission(student.get("admission_number"))

            if admission_key: valid_students[admission_key] = student

        store = read_essay_store(year, class_level, subject, term)

        store.update({
            "version": 1,
            "year": year,

            "class_level": class_level,
            "class_category": class_level,

            "term": term,
            "term_label": term_label(term),

            "subject": subject,

            "objective_max": score_output(config["objective_max"]),
            "essay_max": score_output(essay_max),
            "total_max": 100,

            "updated_at": now_string(),
        })

        if not isinstance(store.get("scores"), dict): store["scores"] = {}

        saved = 0
        cleared = 0
        rejected = []

        for item in scores:
            if not isinstance(item, dict): continue

            admission = clean_text(item.get("admission_number") or item.get("Admission_number") or item.get("Admission No"))
            admission_key = normalize_admission(admission)

            student_name = clean_text(item.get("student_name") or item.get("Student Name"))
            class_arm = clean_text(item.get("class_arm") or item.get("Class Arm")).upper()

            if not admission_key:
                rejected.append({"student_name": student_name, "reason": "Missing admission number"})
                continue

            roster_student = valid_students.get(admission_key)

            if not roster_student:
                rejected.append({
                    "admission_number": admission,
                    "student_name": student_name,
                    "reason": "Student was not found in the class database",
                })
                continue

            value = item.get("score")

            # Blank value removes an existing theory score.
            if value is None or str(value).strip() == "":
                if admission_key in store["scores"]:
                    store["scores"].pop(admission_key, None)
                    cleared += 1

                continue

            score = score_float(value)

            if score is None:
                rejected.append({
                    "admission_number": admission,
                    "student_name": student_name,
                    "reason": "Invalid score",
                })
                continue

            if score < 0 or score > essay_max:
                rejected.append({
                    "admission_number": admission,
                    "student_name": student_name,
                    "reason": f"Essay score must be between 0 and {score_output(essay_max)}",
                })
                continue

            normalized_arm = normalize_class_arm(
                roster_student.get("class_arm") or class_arm,
                class_level,
            ) or roster_student.get("class_arm") or class_arm

            store["scores"][admission_key] = {
                "admission_number": roster_student.get("admission_number") or admission,
                "student_name": roster_student.get("student_name") or student_name,
                "class_arm": normalized_arm,

                "score": score_output(score),

                "updated_at": now_string(),
                "updated_by": clean_text(session.get("username") or session.get("user_name") or session.get("user_type") or "teacher"),
            }

            saved += 1

        if not saved and not cleared and rejected:
            return jsonify({
                "error": "No valid essay scores were saved",
                "rejected": rejected,
            }), 400

        path = write_essay_store(year, class_level, subject, term, store)
        sync_info = queue_essay_store_sync(year, class_level, subject, term, store)

        return jsonify({
            "success": True,
            "message": "Essay scores saved successfully.",

            "saved_count": saved,
            "cleared_count": cleared,

            "rejected_count": len(rejected),
            "rejected": rejected,

            "year": year,
            "class": class_level,

            "term": term,
            "term_label": term_label(term),

            "supports_term": True,
            "requires_term": is_jss_class(class_level),
            "term_mode": "required" if is_jss_class(class_level) else "term" if term else "general",

            "subject": subject,
            "essay_max": score_output(essay_max),

            "path": str(path),
            "sync": sync_info,
        }), 200