# modules/api_routes.py
from flask import Blueprint, jsonify, request, session, Response, stream_with_context, send_file
from pathlib import Path
from collections import deque
from threading import Lock
from datetime import datetime
from io import BytesIO
import json, time, uuid, sqlite3, re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.page import PageMargins

from modules.excel_manager import read_results, get_excel_path, EXPECTED_HEADERS, repair_missing_headers
from modules.class_config import SUPPORTED_CLASSES, normalize_class_level
from modules.student_results import normalize_result_year, academic_session_from_year
from modules.academic_settings import get_academic_settings as get_global_academic_settings, save_academic_settings, YEAR_MIN as ACADEMIC_YEAR_MIN, YEAR_MAX as ACADEMIC_YEAR_MAX
from modules.result_sync import queue_emis_event_safely
from push import get_latest_year, get_active_term_for_target, valid_target_year
from modules.essay_results import register_essay_routes, enrich_results_with_essay


api_bp = Blueprint("api_bp", __name__)
register_essay_routes(api_bp)

BASE_DIR = Path(__file__).resolve().parent.parent
RESULTS_DIR = BASE_DIR / "RESULTS"
SUBJECT_LIBRARY_DIR = BASE_DIR / "static" / "subjects"

VALID_TERMS = ("FIRST", "SECOND", "THIRD")
STAFF_SESSION_TIMEOUT_SECONDS = 2 * 60 * 60
STAFF_SESSION_ACTIVITY_KEY = "_staff_last_activity"

ADMIN_NOTIFICATION_LIMIT = 1000
ADMIN_NOTIFICATIONS = deque(maxlen=ADMIN_NOTIFICATION_LIMIT)
ADMIN_NOTIFICATION_LOCK = Lock()
ADMIN_NOTIFICATION_SEQUENCE = 0


# ============================================================
# HELPERS
# ============================================================

def normalize_class(value):
    return normalize_class_level(value)


def is_jss_class(value):
    return str(normalize_class(value) or "").upper().startswith("JSS")


def is_ss_class(value):
    return str(normalize_class(value) or "").upper().startswith("SS")


def supports_result_terms(value):
    return normalize_class(value) in SUPPORTED_CLASSES


def normalize_term(value):
    raw = " ".join(str(value or "").strip().upper().replace("_", " ").replace("-", " ").split())
    aliases = {
        "FIRST": "FIRST", "FIRST TERM": "FIRST", "TERM 1": "FIRST", "TERM ONE": "FIRST",
        "1": "FIRST", "1ST": "FIRST", "1ST TERM": "FIRST",

        "SECOND": "SECOND", "SECOND TERM": "SECOND", "TERM 2": "SECOND", "TERM TWO": "SECOND",
        "2": "SECOND", "2ND": "SECOND", "2ND TERM": "SECOND",

        "THIRD": "THIRD", "THIRD TERM": "THIRD", "TERM 3": "THIRD", "TERM THREE": "THIRD",
        "3": "THIRD", "3RD": "THIRD", "3RD TERM": "THIRD",
    }
    return aliases.get(raw, "")


def term_label(value):
    return {
        "FIRST": "FIRST TERM",
        "SECOND": "SECOND TERM",
        "THIRD": "THIRD TERM",
    }.get(normalize_term(value), "")


def short_term_label(value):
    return {
        "FIRST": "1st Term",
        "SECOND": "2nd Term",
        "THIRD": "3rd Term",
    }.get(normalize_term(value), "")


def clean_records(records):
    return [{key: value if value is not None else "" for key, value in row.items()} for row in records]


def normalize_score_value(value):
    try:
        return int(float(str(value or "").replace("%", "").strip()))
    except (TypeError, ValueError):
        return 0


def normalize_subject_folder(subject):
    return str(subject or "").strip()


def subject_key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def subject_matches(filter_value, folder_value, record_subject=""):
    if not filter_value or str(filter_value).lower() == "all":
        return True

    wanted = subject_key(filter_value)
    candidates = {subject_key(folder_value), subject_key(record_subject)}

    if wanted in candidates:
        return True

    raw_filter = str(filter_value).lower()
    raw_folder = str(folder_value).lower()

    return raw_filter in raw_folder or raw_folder in raw_filter


def can_view_results():
    role = str(session.get("user_type") or "").strip().lower()

    if role not in {"admin", "teacher"}:
        return False

    now = time.time()
    last_activity = session.get(STAFF_SESSION_ACTIVITY_KEY)

    if last_activity is not None:
        try:
            elapsed = now - float(last_activity)

            if elapsed > STAFF_SESSION_TIMEOUT_SECONDS:
                print(
                    f"[RESULT SESSION] {role} session expired after "
                    f"{round(elapsed / 3600, 2)} hour(s) inactivity."
                )
                session.clear()
                return False

        except (TypeError, ValueError):
            pass

    # Any genuine staff API activity refreshes inactivity timer.
    session.permanent = True
    session[STAFF_SESSION_ACTIVITY_KEY] = now
    session.modified = True

    return True


def get_requested_term(class_cat=""):
    class_cat = normalize_class(class_cat)

    if not class_cat:
        return ""

    raw = request.args.get("term")

    if raw is not None:
        raw = str(raw or "").strip()

        if raw.lower() == "all":
            return ""

        term = normalize_term(raw)

        if term:
            return term

    term = normalize_term(session.get("selected_term"))

    if term:
        return term

    # Global academic settings are the normal staff-module default.
    try:
        term = normalize_term((get_global_academic_settings() or {}).get("current_term"))
        if term:
            return term
    except Exception as error:
        print("GLOBAL TERM RESOLUTION ERROR:", error)

    # Class/arm push settings remain a final compatibility fallback.
    try:
        return normalize_term(get_active_term_for_target(class_cat, class_cat))
    except Exception as error:
        print("ACTIVE TERM RESOLUTION ERROR:", error)
        return ""


def normalize_result_record(row, class_cat="", subject_name="", year="", term=""):
    class_raw = (
        row.get("Class Level")
        or row.get("Class Category")
        or row.get("Class")
        or class_cat
    )

    arm_raw = (
        row.get("Class Arm")
        or row.get("Class")
        or class_raw
    )

    class_level = normalize_class(class_raw) or str(class_raw or "").upper().strip()
    class_arm = str(arm_raw or class_level).upper().strip()
    row_term = normalize_term(row.get("Term") or term)

    row["Year"] = row.get("Year") or year
    row["Class Level"] = class_level
    row["Class Arm"] = class_arm
    row["Class"] = class_arm or class_level
    row["Class Category"] = class_level

    row["Term"] = row_term
    row["Term Label"] = term_label(row_term) if row_term else ""

    row["Subject"] = (
        row.get("Subject")
        or str(subject_name or "").replace("_", " ").upper()
    )

    row["Subject Folder"] = subject_name or row.get("Subject Folder") or ""
    row["Score Number"] = normalize_score_value(row.get("Score (%)"))

    return row


def read_excel_file_direct(excel_path: Path):
    if not excel_path or not excel_path.exists():
        return []

    wb = load_workbook(excel_path)
    ws = wb.active

    repair_missing_headers(ws)

    rows = list(ws.iter_rows(values_only=True))
    wb.save(excel_path)

    if len(rows) < 2:
        return []

    headers = list(rows[0])
    results = []

    for row in rows[1:]:
        values = list(row) + [None] * max(0, len(headers) - len(row))
        item = dict(zip(headers, values[:len(headers)]))

        if item.get("Student Name") or item.get("Admission No"):
            results.append(item)

    return results


def get_class_result_root(year, class_cat):
    return RESULTS_DIR / str(year) / "CLASS" / str(class_cat)


def get_term_result_path(class_cat, subject, year, term=""):
    class_cat = normalize_class(class_cat)
    subject = normalize_subject_folder(subject)
    term = normalize_term(term)

    root = get_class_result_root(year, class_cat)

    if term:
        return root / term / subject / "results.xlsx"

    if is_jss_class(class_cat):
        return None

    return root / subject / "results.xlsx"


def get_legacy_result_path(class_cat, subject, year):
    try:
        return get_excel_path(normalize_class(class_cat), subject, year)

    except Exception:
        return (
            get_class_result_root(year, normalize_class(class_cat))
            / str(subject)
            / "results.xlsx"
        )


def find_existing_result_path(class_cat, subject, year, term=""):
    class_cat = normalize_class(class_cat)
    term = normalize_term(term)

    term_path = get_term_result_path(class_cat, subject, year, term)

    if term_path and term_path.exists():
        return term_path

    legacy = get_legacy_result_path(class_cat, subject, year)

    if not term and legacy and legacy.exists():
        return legacy

    if is_jss_class(class_cat) and term and legacy and legacy.exists():
        return legacy

    return term_path or legacy


def read_results_term_aware(class_cat, subject, year, term=""):
    class_cat = normalize_class(class_cat)
    term = normalize_term(term)

    term_path = get_term_result_path(class_cat, subject, year, term)

    if term and term_path and term_path.exists():
        return read_excel_file_direct(term_path)

    legacy = get_legacy_result_path(class_cat, subject, year)

    if legacy and legacy.exists():
        rows = read_excel_file_direct(legacy)

        if term and is_ss_class(class_cat):
            return [
                row
                for row in rows
                if normalize_term(row.get("Term")) == term
            ]

        return rows

    if term and is_ss_class(class_cat):
        return []

    try:
        return read_results(class_cat, subject, year, term)

    except Exception as error:
        print("TERM-AWARE RESULT READ FALLBACK ERROR:", error)
        return []


def get_terms_for_class(year, class_cat):
    class_cat = normalize_class(class_cat)

    if not class_cat:
        return []

    terms = set()

    result_root = get_class_result_root(year, class_cat)
    library_root = (
        SUBJECT_LIBRARY_DIR
        / str(year)
        / "subjects-json"
        / class_cat
    )

    for root in (result_root, library_root):
        if not root.exists():
            continue

        for folder in root.iterdir():
            if not folder.is_dir():
                continue

            term = normalize_term(folder.name)

            if term:
                terms.add(term)

    # Also support transitional flat SS files that already contain Term values.
    if result_root.exists():
        for folder in result_root.iterdir():
            if not folder.is_dir() or normalize_term(folder.name):
                continue

            path = folder / "results.xlsx"

            if not path.exists():
                continue

            try:
                for row in read_excel_file_direct(path):
                    term = normalize_term(row.get("Term"))

                    if term:
                        terms.add(term)

            except Exception as error:
                print(f"TERM DISCOVERY ERROR [{path}]:", error)

    return sorted(
        terms,
        key=lambda value: VALID_TERMS.index(value)
        if value in VALID_TERMS
        else 999,
    )


def get_library_subjects(year, class_cat, term=""):
    class_cat = normalize_class(class_cat)
    term = normalize_term(term)

    base = (
        SUBJECT_LIBRARY_DIR
        / str(year)
        / "subjects-json"
        / class_cat
    )

    if is_jss_class(class_cat):
        if not term:
            return []

        roots = [base / term]

    elif is_ss_class(class_cat):
        roots = [base / term, base] if term else [base]

    else:
        roots = [base]

    subjects = {}

    for root in roots:
        if not root.exists():
            continue

        for path in root.glob("*.json"):
            if path.name.lower() == "pushed_subjects.json":
                continue

            try:
                with path.open("r", encoding="utf-8-sig") as file:
                    data = json.load(file)

                subject = str(
                    data.get("subject")
                    or path.stem
                ).strip()

                key = subject_key(subject)

                if key and key not in subjects:
                    subjects[key] = subject

            except Exception as error:
                print(f"Could not read subject JSON [{path}]: {error}")

    return sorted(subjects.values(), key=str.lower)


def get_subject_folders(year, class_cat, term=""):
    class_cat = normalize_class(class_cat)
    term = normalize_term(term)

    root = get_class_result_root(year, class_cat)

    if not root.exists():
        return []

    subjects = set()

    if term:
        term_root = root / term

        if term_root.exists():
            subjects.update(
                folder.name
                for folder in term_root.iterdir()
                if folder.is_dir()
                and not normalize_term(folder.name)
            )

        for folder in root.iterdir():
            if not folder.is_dir() or normalize_term(folder.name):
                continue

            path = folder / "results.xlsx"

            if not path.exists():
                continue

            # JSS keeps legacy flat-result compatibility.
            if is_jss_class(class_cat):
                subjects.add(folder.name)
                continue

            # SS flat files only belong to a term when the row says so.
            try:
                if any(
                    normalize_term(row.get("Term")) == term
                    for row in read_excel_file_direct(path)
                ):
                    subjects.add(folder.name)

            except Exception:
                pass

        return sorted(subjects, key=str.lower)

    return sorted(
        [
            folder.name
            for folder in root.iterdir()
            if folder.is_dir()
            and not normalize_term(folder.name)
        ],
        key=str.lower,
    )


def iter_result_files():
    if not RESULTS_DIR.exists():
        return

    for year_folder in RESULTS_DIR.iterdir():
        class_root = year_folder / "CLASS"

        if not year_folder.is_dir() or not class_root.exists():
            continue

        for class_folder in class_root.iterdir():
            if not class_folder.is_dir():
                continue

            class_cat = normalize_class(class_folder.name)

            if class_cat not in SUPPORTED_CLASSES:
                continue

            for child in class_folder.iterdir():
                if not child.is_dir():
                    continue

                term = normalize_term(child.name)

                # FIRST / SECOND / THIRD
                if term:
                    for subject_folder in child.iterdir():
                        path = subject_folder / "results.xlsx"

                        if subject_folder.is_dir() and path.exists():
                            yield {
                                "year": year_folder.name,
                                "class": class_cat,
                                "term": term,
                                "subject": subject_folder.name,
                                "path": path,
                            }

                    continue

                # General / legacy flat structure
                path = child / "results.xlsx"

                if path.exists():
                    yield {
                        "year": year_folder.name,
                        "class": class_cat,
                        "term": "",
                        "subject": child.name,
                        "path": path,
                    }


def get_output_headers():
    headers = list(EXPECTED_HEADERS)

    if "Term" not in headers:
        if "Subject" in headers:
            headers.insert(
                headers.index("Subject") + 1,
                "Term",
            )
        else:
            headers.append("Term")

    return headers


def enrich_records_with_essay(records, class_cat, subject, year, folder_term=""):
    if not records:
        return records

    groups = {}

    for row in records:
        term = normalize_term(
            row.get("Term")
            or folder_term
        )

        groups.setdefault(term, []).append(row)

    for term, group in groups.items():
        try:
            enrich_results_with_essay(
                group,
                class_cat,
                subject,
                year,
                term,
            )

        except Exception as error:
            print(
                f"ESSAY RESULT ENRICH ERROR "
                f"[{year}/{class_cat}/{term or 'GENERAL'}/{subject}]: "
                f"{error}"
            )

    return records


def candidate_result_paths(class_cat, subject, year, term=""):
    class_cat = normalize_class(class_cat)
    term = normalize_term(term)

    candidates = []

    if term:
        path = get_term_result_path(
            class_cat,
            subject,
            year,
            term,
        )

        if path and path.exists():
            candidates.append((path, term))

        legacy = get_legacy_result_path(
            class_cat,
            subject,
            year,
        )

        if legacy and legacy.exists():
            candidates.append((legacy, ""))

    elif is_jss_class(class_cat):
        for possible_term in VALID_TERMS:
            path = get_term_result_path(
                class_cat,
                subject,
                year,
                possible_term,
            )

            if path and path.exists():
                candidates.append(
                    (path, possible_term)
                )

        legacy = get_legacy_result_path(
            class_cat,
            subject,
            year,
        )

        if legacy and legacy.exists():
            candidates.append((legacy, ""))

    else:
        # Blank SS term means General / legacy only.
        legacy = get_legacy_result_path(
            class_cat,
            subject,
            year,
        )

        if legacy and legacy.exists():
            candidates.append((legacy, ""))

    unique = []
    seen = set()

    for path, actual_term in candidates:
        key = str(path.resolve())

        if key in seen:
            continue

        seen.add(key)
        unique.append(
            (path, actual_term)
        )

    return unique


def delete_sqlite_result(cursor, admission, name, year, class_cat, subject, term):
    identity_column = "admission_number" if admission else "full_name"
    identity_value = admission or name

    cursor.execute(
        f"""
        DELETE FROM student_results
        WHERE LOWER(TRIM(COALESCE({identity_column}, ''))) = LOWER(TRIM(?))
          AND TRIM(COALESCE(year, '')) = TRIM(?)
          AND UPPER(TRIM(COALESCE(class_level, ''))) = UPPER(TRIM(?))
          AND UPPER(TRIM(COALESCE(subject, ''))) = UPPER(TRIM(?))
          AND UPPER(TRIM(COALESCE(term, ''))) = UPPER(TRIM(?))
        """,
        (
            identity_value,
            year,
            class_cat,
            subject,
            term,
        ),
    )

    return max(cursor.rowcount, 0)


# ============================================================
# REAL-TIME ADMIN NOTIFICATIONS
# ============================================================

def push_admin_notification(event_type, message="", payload=None):
    global ADMIN_NOTIFICATION_SEQUENCE

    payload = payload if isinstance(payload, dict) else {}

    with ADMIN_NOTIFICATION_LOCK:
        ADMIN_NOTIFICATION_SEQUENCE += 1

        event = {
            "id": uuid.uuid4().hex,
            "sequence": ADMIN_NOTIFICATION_SEQUENCE,
            "type": str(event_type or "notification").strip().lower(),
            "message": str(message or "").strip(),
            "payload": payload,
            "created_at": int(time.time() * 1000),
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
        }

        ADMIN_NOTIFICATIONS.append(event)

    return event


def student_notification_payload(data):
    student = (
        session.get("student")
        if isinstance(session.get("student"), dict)
        else {}
    )

    name = str(
        data.get("student_name")
        or session.get("student_name")
        or student.get("full_name")
        or "Student"
    ).strip()

    admission = str(
        data.get("admission_number")
        or session.get("admission_number")
        or student.get("admission_number")
        or ""
    ).strip()

    payload = {
        "student_name": name,
        "full_name": name,
        "admission_number": admission,
        "student_id": admission,

        "class_category": (
            data.get("class_category")
            or session.get("class_category")
            or student.get("class_category")
            or ""
        ),

        "class_level": (
            data.get("class_level")
            or session.get("class_level")
            or student.get("class_level")
            or ""
        ),

        "class_arm": (
            data.get("class_arm")
            or session.get("class_arm")
            or student.get("class_arm")
            or ""
        ),

        "subject": (
            data.get("subject")
            or session.get("selected_subject")
            or ""
        ),

        "year": str(
            data.get("year")
            or session.get("selected_year")
            or ""
        ),

        "term": normalize_term(
            data.get("term")
            or session.get("selected_term")
            or ""
        ),

        "term_label": data.get("term_label") or "",
    }

    return name, admission, payload


@api_bp.route("/api/notifications/notify/exam_start", methods=["POST"])
def notify_exam_start():
    if str(session.get("user_type", "")).lower() != "student":
        return jsonify({"error": "Unauthorized"}), 403

    data = request.get_json(silent=True) or {}
    name, admission, payload = student_notification_payload(data)

    payload["started_at"] = data.get("started_at") or ""

    event = push_admin_notification(
        "exam_start",
        f"{name} started {payload['subject'] or 'an examination'}.",
        payload,
    )

    print(
        "[NOTIFICATION] EXAM START:",
        event["sequence"],
        admission,
        payload["subject"],
    )

    return jsonify({
        "status": "ok",
        "notification": event,
    }), 200


@api_bp.route("/api/notifications/notify/exam_end", methods=["POST"])
def notify_exam_end():
    if str(session.get("user_type", "")).lower() != "student":
        return jsonify({"error": "Unauthorized"}), 403

    data = request.get_json(silent=True) or {}
    name, admission, payload = student_notification_payload(data)

    status = str(
        data.get("status")
        or "completed"
    ).strip().lower()

    payload.update({
        "score": data.get("score", ""),
        "total_questions": data.get("total_questions", ""),
        "flagged": data.get("flagged", 0),
        "submitted_at": data.get("submitted_at") or "",
        "status": status,
    })

    event_type = (
        "timeout"
        if status in {"timeout", "timed_out"}
        else "exam_end"
    )

    message = (
        f"{name}'s examination timed out."
        if event_type == "timeout"
        else f"{name} submitted {payload['subject'] or 'an examination'}."
    )

    event = push_admin_notification(
        event_type,
        message,
        payload,
    )

    print(
        "[NOTIFICATION] EXAM END:",
        event["sequence"],
        admission,
        payload["subject"],
        status,
    )

    return jsonify({
        "status": "ok",
        "notification": event,
    }), 200


@api_bp.route("/api/notifications/fetch")
def fetch_admin_notifications():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized",
            "notifications": [],
        }), 403

    try:
        since = int(
            request.args.get(
                "since_sequence",
                0,
            )
            or 0
        )

    except (TypeError, ValueError):
        since = 0

    with ADMIN_NOTIFICATION_LOCK:
        if since:
            notifications = [
                dict(item)
                for item in ADMIN_NOTIFICATIONS
                if int(item.get("sequence", 0) or 0) > since
            ]
        else:
            notifications = [
                dict(item)
                for item in ADMIN_NOTIFICATIONS
            ]

    notifications.sort(
        key=lambda item: int(
            item.get("sequence", 0)
            or 0
        )
    )

    latest = max(
        (
            int(item.get("sequence", 0) or 0)
            for item in notifications
        ),
        default=since,
    )

    return jsonify({
        "notifications": notifications,
        "count": len(notifications),
        "latest_sequence": latest,
    }), 200


@api_bp.route("/api/notifications/stream")
def stream_admin_notifications():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    try:
        starting_sequence = int(
            request.args.get(
                "since_sequence",
                0,
            )
            or 0
        )

    except (TypeError, ValueError):
        starting_sequence = 0

    @stream_with_context
    def event_stream():
        last_sequence = starting_sequence

        yield (
            f'event: connected\n'
            f'data: {json.dumps({"status": "connected", "sequence": last_sequence})}\n\n'
        )

        while True:
            with ADMIN_NOTIFICATION_LOCK:
                events = [
                    dict(item)
                    for item in ADMIN_NOTIFICATIONS
                    if int(item.get("sequence", 0) or 0) > last_sequence
                ]

            events.sort(
                key=lambda item: int(
                    item.get("sequence", 0)
                    or 0
                )
            )

            for item in events:
                last_sequence = max(
                    last_sequence,
                    int(item.get("sequence", 0) or 0),
                )

                yield (
                    f"event: notification\n"
                    f"data: {json.dumps(item, ensure_ascii=False)}\n\n"
                )

            yield ": keepalive\n\n"
            time.sleep(1)

    return Response(
        event_stream(),
        mimetype="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


# ============================================================
# RESULT DISCOVERY / LOADING
# ============================================================

@api_bp.route("/api/results/years")
def get_years():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    years = (
        sorted(
            [
                folder.name
                for folder in RESULTS_DIR.iterdir()
                if folder.is_dir()
            ],
            reverse=True,
        )
        if RESULTS_DIR.exists()
        else []
    )

    return jsonify({
        "years": years
    })


@api_bp.route("/api/results/classes")
def get_classes_for_year():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    root = (
        RESULTS_DIR
        / year
        / "CLASS"
    )

    classes = (
        sorted([
            folder.name
            for folder in root.iterdir()
            if folder.is_dir()
            and normalize_class(folder.name) in SUPPORTED_CLASSES
        ])
        if year and root.exists()
        else []
    )

    return jsonify({
        "classes": classes
    })


@api_bp.route("/api/results/terms")
def get_terms_for_year_and_class():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    class_cat = normalize_class(
        request.args.get(
            "class",
            "",
        )
    )

    terms = (
        get_terms_for_class(
            year,
            class_cat,
        )
        if year and class_cat
        else []
    )

    return jsonify({
        "terms": terms,
        "class": class_cat,
        "supports_term": supports_result_terms(class_cat),
        "requires_term": is_jss_class(class_cat),
    })


@api_bp.route("/api/results/subjects")
def get_subjects_for_class_and_year():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    class_cat = normalize_class(
        request.args.get(
            "class",
            "",
        )
    )

    if not year or not class_cat:
        return jsonify({
            "subjects": []
        })

    term = get_requested_term(
        class_cat
    )

    if is_jss_class(class_cat) and not term:
        return jsonify({
            "subjects": [],
            "year": year,
            "class": class_cat,
            "term": "",
            "supports_term": True,
            "requires_term": True,
        })

    subjects = {}

    source_subjects = [
        *get_subject_folders(
            year,
            class_cat,
            term,
        ),

        *get_library_subjects(
            year,
            class_cat,
            term,
        ),
    ]

    for subject in source_subjects:
        key = subject_key(subject)

        if key and key not in subjects:
            subjects[key] = subject

    return jsonify({
        "subjects": sorted(
            subjects.values(),
            key=str.lower,
        ),

        "year": year,
        "class": class_cat,
        "term": term,

        "supports_term": supports_result_terms(class_cat),
        "requires_term": is_jss_class(class_cat),
    })


@api_bp.route("/api/results/load")
def load_excel_results():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized",
            "results": [],
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    class_cat = normalize_class(
        request.args.get(
            "class",
            "",
        )
    )

    subject = request.args.get(
        "subject",
        "",
    ).strip()

    if not year or not class_cat or not subject:
        return jsonify({
            "error": "Missing parameters",
            "results": [],
        }), 400

    term = get_requested_term(
        class_cat
    )

    if is_jss_class(class_cat) and not term:
        return jsonify({
            "error": "Term is required for JSS results",
            "results": [],
        }), 400

    try:
        records = read_results_term_aware(
            class_cat,
            subject,
            year,
            term,
        )

        for row in records:
            normalize_result_record(
                row,
                class_cat,
                subject,
                year,
                term,
            )

        enrich_records_with_essay(
            records,
            class_cat,
            subject,
            year,
            term,
        )

        return jsonify({
            "results": clean_records(records),
            "year": year,
            "class": class_cat,
            "subject": subject,
            "term": term,
            "term_label": term_label(term) if term else "",
            "supports_term": supports_result_terms(class_cat),
            "requires_term": is_jss_class(class_cat),
            "count": len(records),
        }), 200

    except Exception as error:
        print(
            "RESULT LOAD ERROR:",
            error,
        )

        return jsonify({
            "error": "Failed to load results",
            "details": str(error),
            "results": [],
            "year": year,
            "class": class_cat,
            "subject": subject,
            "term": term,
        }), 500


@api_bp.route("/api/results/exists")
def excel_exists():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    class_cat = normalize_class(
        request.args.get(
            "class",
            "",
        )
    )

    subject = request.args.get(
        "subject",
        "",
    ).strip()

    if not year or not class_cat or not subject:
        return jsonify({
            "exists": False
        })

    term = get_requested_term(
        class_cat
    )

    if is_jss_class(class_cat) and not term:
        return jsonify({
            "exists": False,
            "error": "Term is required for JSS results",
        })

    term_path = get_term_result_path(
        class_cat,
        subject,
        year,
        term,
    )

    path = term_path
    exists = bool(
        term_path
        and term_path.exists()
    )

    if not exists:
        legacy = get_legacy_result_path(
            class_cat,
            subject,
            year,
        )

        if legacy and legacy.exists():
            if term and is_ss_class(class_cat):
                try:
                    exists = any(
                        normalize_term(row.get("Term")) == term
                        for row in read_excel_file_direct(legacy)
                    )

                except Exception:
                    exists = False

                if exists:
                    path = legacy

            elif not term or is_jss_class(class_cat):
                path = legacy
                exists = True

    return jsonify({
        "exists": exists,
        "path": str(path) if path and exists else "",
        "year": year,
        "class": class_cat,
        "subject": subject,
        "term": term,
    })


@api_bp.route("/api/results")
def api_get_results():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized",
            "results": [],
        }), 403

    year = request.args.get(
        "year",
        "",
    ).strip()

    class_raw = request.args.get(
        "class",
        "",
    ).strip()

    subject = request.args.get(
        "subject",
        "",
    ).strip()

    term_raw = request.args.get(
        "term",
        "",
    ).strip()

    if (
        not year
        or year.lower() == "all"
        or not class_raw
        or class_raw.lower() == "all"
        or not subject
        or subject.lower() == "all"
        or term_raw.lower() == "all"
    ):
        return api_get_all_results()

    class_cat = normalize_class(
        class_raw
    )

    if not class_cat:
        return jsonify({
            "error": "Invalid class",
            "results": [],
        }), 400

    term = get_requested_term(
        class_cat
    )

    if is_jss_class(class_cat) and not term:
        return jsonify({
            "error": "Term is required for JSS results",
            "results": [],
        }), 400

    try:
        records = read_results_term_aware(
            class_cat,
            subject,
            year,
            term,
        )

        for row in records:
            normalize_result_record(
                row,
                class_cat,
                subject,
                year,
                term,
            )

        enrich_records_with_essay(
            records,
            class_cat,
            subject,
            year,
            term,
        )

        return jsonify({
            "results": clean_records(records),
            "year": year,
            "class": class_cat,
            "subject": subject,
            "term": term,
        }), 200

    except Exception as error:
        print(
            "RESULT READ ERROR:",
            error,
        )

        return jsonify({
            "error": "Failed to read results",
            "details": str(error),
            "results": [],
        }), 500


@api_bp.route("/api/results/all")
def api_get_all_results():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized",
            "results": [],
        }), 403

    year_filter = request.args.get(
        "year",
        "all",
    ).strip()

    class_raw = request.args.get(
        "class",
        "all",
    ).strip()

    subject_filter = request.args.get(
        "subject",
        "all",
    ).strip()

    term_raw = request.args.get(
        "term",
        "all",
    ).strip()

    class_filter = (
        normalize_class(class_raw)
        if class_raw.lower() != "all"
        else "all"
    )

    term_filter = (
        normalize_term(term_raw)
        if term_raw.lower() != "all"
        else "all"
    )

    if not RESULTS_DIR.exists():
        return jsonify({
            "results": [],

            "summary": {
                "total": 0,
                "years": [],
                "classes": [],
                "terms": [],
                "subjects": [],
            },
        }), 200

    all_results = []

    years_found = set()
    classes_found = set()
    terms_found = set()
    subjects_found = set()

    for info in iter_result_files():
        year = info["year"]
        class_cat = info["class"]
        folder_term = info["term"]
        subject_name = info["subject"]
        path = info["path"]

        if (
            year_filter.lower() != "all"
            and year != year_filter
        ):
            continue

        if (
            class_filter != "all"
            and class_cat != class_filter
        ):
            continue

        if (
            term_filter != "all"
            and folder_term
            and folder_term != term_filter
        ):
            continue

        try:
            records = read_excel_file_direct(
                path
            )

        except Exception as error:
            print(
                f"Could not read {path}: "
                f"{error}"
            )

            continue

        for row in records:
            normalize_result_record(
                row,
                class_cat,
                subject_name,
                year,
                folder_term,
            )

        # FIRST/SECOND/THIRD filtering works for both JSS and SS.
        if term_filter != "all":
            records = [
                row
                for row in records
                if normalize_term(
                    row.get("Term")
                ) == term_filter
            ]

        if subject_filter.lower() != "all":
            records = [
                row
                for row in records
                if subject_matches(
                    subject_filter,
                    subject_name,
                    row.get("Subject"),
                )
            ]

        if not records:
            continue

        enrich_records_with_essay(
            records,
            class_cat,
            subject_name,
            year,
            folder_term
            or (
                term_filter
                if term_filter != "all"
                else ""
            ),
        )

        all_results.extend(records)

        years_found.add(year)
        classes_found.add(class_cat)
        subjects_found.add(subject_name)

        if folder_term:
            terms_found.add(folder_term)

        for row in records:
            row_term = normalize_term(
                row.get("Term")
            )

            if row_term:
                terms_found.add(
                    row_term
                )

    all_results.sort(
        key=lambda row: str(
            row.get(
                "Submitted At",
                "",
            )
        ),
        reverse=True,
    )

    return jsonify({
        "results": clean_records(
            all_results
        ),

        "summary": {
            "total": len(all_results),

            "years": sorted(
                years_found,
                reverse=True,
            ),

            "classes": sorted(
                classes_found
            ),

            "terms": sorted(
                terms_found,
                key=lambda value: (
                    VALID_TERMS.index(value)
                    if value in VALID_TERMS
                    else 999
                ),
            ),

            "subjects": sorted(
                subjects_found,
                key=str.lower,
            ),
        },
    }), 200


# ============================================================
# RESULT DELETION
# ============================================================

@api_bp.route("/api/results/delete", methods=["POST"])
def delete_excel_results():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    data = request.get_json(
        silent=True
    ) or {}

    delete_list = data.get(
        "delete_items",
        [],
    )

    if not isinstance(delete_list, list) or not delete_list:
        return jsonify({
            "error": "No items to delete"
        }), 400

    grouped = {}
    skipped = []

    for item in delete_list:
        if not isinstance(item, dict):
            continue

        year = str(
            item.get("Year")
            or item.get("year")
            or data.get("year")
            or ""
        ).strip()

        class_cat = normalize_class(
            item.get("Class Level")
            or item.get("Class Category")
            or item.get("class_level")
            or item.get("class_category")
            or item.get("Class")
            or item.get("class")
            or data.get("class_category")
            or data.get("class")
        )

        subject = str(
            item.get("Subject Folder")
            or item.get("subject_folder")
            or item.get("Subject")
            or item.get("subject")
            or data.get("subject")
            or ""
        ).strip()

        term = normalize_term(
            item.get("Term")
            or item.get("term")
            or data.get("term")
        )

        student_name = str(
            item.get("Student Name")
            or item.get("full_name")
            or item.get("student_name")
            or ""
        ).strip().upper()

        admission = str(
            item.get("Admission No")
            or item.get("admission_number")
            or item.get("student_id")
            or ""
        ).strip().upper()

        if (
            class_cat
            and is_jss_class(class_cat)
            and not term
        ):
            term = get_requested_term(
                class_cat
            )

        if not year or not class_cat or not subject:
            skipped.append({
                "reason": "Missing year, class or subject",
                "year": year,
                "class": class_cat,
                "subject": subject,
                "admission": admission,
            })

            continue

        if not student_name and not admission:
            skipped.append({
                "reason": "Missing student identity",
                "year": year,
                "class": class_cat,
                "subject": subject,
            })

            continue

        key = (
            year,
            class_cat,
            term,
            subject,
        )

        grouped.setdefault(
            key,
            set(),
        ).add(
            (
                student_name,
                admission,
            )
        )

    if not grouped:
        return jsonify({
            "error": "No valid result records were supplied for deletion",
            "deleted": 0,
            "skipped": skipped,
        }), 400

    deleted_count = 0
    sqlite_deleted = 0
    touched_files = 0

    sqlite_path = (
        BASE_DIR
        / "database.db"
    )

    for (
        year,
        class_cat,
        requested_term,
        subject,
    ), targets in grouped.items():

        candidates = candidate_result_paths(
            class_cat,
            subject,
            year,
            requested_term,
        )

        if not candidates:
            skipped.append({
                "reason": "Result file not found",
                "year": year,
                "class": class_cat,
                "term": requested_term,
                "subject": subject,
            })

            continue

        for path, folder_term in candidates:
            try:
                results = read_excel_file_direct(
                    path
                )

            except Exception as error:
                print(
                    f"DELETE READ ERROR "
                    f"[{path}]: {error}"
                )

                skipped.append({
                    "reason": "Could not read result file",
                    "file": str(path),
                })

                continue

            updated = []
            removed = []

            for row in results:
                row_name = str(
                    row.get("Student Name")
                    or ""
                ).strip().upper()

                row_admission = str(
                    row.get("Admission No")
                    or ""
                ).strip().upper()

                row_term = normalize_term(
                    row.get("Term")
                    or folder_term
                )

                # Old flat JSS can inherit selected term.
                # Blank SS remains General / Legacy.
                if (
                    requested_term
                    and not row_term
                    and is_jss_class(class_cat)
                ):
                    row_term = requested_term

                if (
                    requested_term
                    and row_term != requested_term
                ):
                    updated.append(row)
                    continue

                match = any(
                    (
                        target_admission
                        and row_admission == target_admission
                    )
                    or (
                        not target_admission
                        and target_name
                        and row_name == target_name
                    )
                    for target_name, target_admission in targets
                )

                if match:
                    removed.append(row)
                else:
                    updated.append(row)

            if not removed:
                continue

            wb = Workbook()
            ws = wb.active
            ws.title = "Results"

            headers = get_output_headers()
            ws.append(headers)

            for row in updated:
                normalized = normalize_result_record(
                    row,
                    class_cat,
                    subject,
                    year,
                    folder_term,
                )

                ws.append([
                    normalized.get(
                        header,
                        "",
                    )
                    for header in headers
                ])

            path.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

            wb.save(path)

            deleted_count += len(removed)
            touched_files += 1

            if sqlite_path.exists():
                conn = None

                try:
                    conn = sqlite3.connect(
                        sqlite_path
                    )

                    cursor = conn.cursor()

                    for row in removed:
                        admission = str(
                            row.get("Admission No")
                            or ""
                        ).strip()

                        name = str(
                            row.get("Student Name")
                            or ""
                        ).strip()

                        stored_subject = str(
                            row.get("Subject")
                            or subject
                        ).strip()

                        db_term = normalize_term(
                            row.get("Term")
                            or folder_term
                            or requested_term
                        )

                        sqlite_deleted += delete_sqlite_result(
                            cursor,
                            admission,
                            name,
                            year,
                            class_cat,
                            stored_subject,
                            db_term,
                        )

                    conn.commit()

                except Exception as error:
                    print(
                        "SQLITE RESULT DELETE ERROR:",
                        error,
                    )

                    if conn:
                        conn.rollback()

                finally:
                    if conn:
                        conn.close()

            # Exact term-folder result found.
            # Do not unnecessarily continue into legacy file.
            if (
                requested_term
                and folder_term == requested_term
            ):
                break

    if deleted_count == 0:
        return jsonify({
            "status": "not_found",
            "error": "No matching result record was found to delete.",
            "deleted": 0,
            "sqlite_deleted": sqlite_deleted,
            "files_updated": touched_files,
            "skipped": skipped,
        }), 404

    return jsonify({
        "status": "ok",
        "message": "Result records permanently deleted.",
        "deleted": deleted_count,
        "sqlite_deleted": sqlite_deleted,
        "files_updated": touched_files,
        "skipped": skipped,
    }), 200


# ============================================================
# GLOBAL RESULT SEARCH
# ============================================================

@api_bp.route("/api/results/search_admission")
def search_admission():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    q = request.args.get(
        "q",
        "",
    ).strip().lower()

    if len(q) < 2 or not RESULTS_DIR.exists():
        return jsonify({
            "results": []
        })

    matches = []

    for info in iter_result_files():
        try:
            records = read_excel_file_direct(
                info["path"]
            )

        except Exception as error:
            print(
                f"Search could not read "
                f"{info['path']}: {error}"
            )

            continue

        for row in records:
            normalize_result_record(
                row,
                info["class"],
                info["subject"],
                info["year"],
                info["term"],
            )

        enrich_records_with_essay(
            records,
            info["class"],
            info["subject"],
            info["year"],
            info["term"],
        )

        for row in records:
            admission = str(
                row.get("Admission No")
                or ""
            ).strip().lower()

            name = str(
                row.get("Student Name")
                or ""
            ).strip().lower()

            if (
                q in admission
                or q in name
            ):
                matches.append(row)

    matches.sort(
        key=lambda row: str(
            row.get(
                "Submitted At",
                "",
            )
        ),
        reverse=True,
    )

    return jsonify({
        "results": clean_records(
            matches
        )
    })


# ============================================================
# ACADEMIC SETTINGS — GLOBAL PERSISTENT DEFAULT
#
# One admin-controlled academic year/session/term is stored locally and used
# as the default context across Attendance, CA/Test, Results and Report Sheets.
# Explicit page selections can still override the default for historical work.
# Local changes are queued through the existing offline-first sync engine so
# the deployed EMIS receives the same active academic period.
# ============================================================

@api_bp.route("/api/academic-settings", methods=["GET"])
def api_get_academic_settings():
    if not can_view_results():
        return jsonify({"error": "Unauthorized"}), 403

    settings = get_global_academic_settings()

    return jsonify({
        "success": True,
        "settings": settings,
        "year_min": ACADEMIC_YEAR_MIN,
        "year_max": ACADEMIC_YEAR_MAX,
        "persistent_source": "local_academic_settings",
    })


@api_bp.route("/api/academic-settings", methods=["POST"])
def api_update_academic_settings():
    if str(session.get("user_type") or "").strip().lower() != "admin":
        return jsonify({"success": False, "error": "Admin access is required to change the global academic period."}), 403

    data = request.get_json(silent=True) or {}
    raw_year = data.get("current_year") or data.get("year") or ""
    raw_session = data.get("current_session") or data.get("academic_session") or data.get("session") or ""
    raw_term = data.get("current_term") or data.get("term") or ""

    try:
        settings = save_academic_settings(
            current_year=raw_year,
            current_session=raw_session,
            current_term=raw_term,
            updated_by=session.get("username") or session.get("admin_username") or "EMIS Admin",
            source="local",
        )
    except ValueError as error:
        return jsonify({"success": False, "error": str(error)}), 400

    # Keep the current staff browser aligned immediately. These session values
    # are conveniences only; the JSON setting above remains authoritative.
    session["selected_year"] = settings["current_year"]
    session["academic_session"] = settings["current_session"]
    session["selected_term"] = settings["current_term"]

    sync_payload = {
        "current_year": settings["current_year"],
        "current_session": settings["current_session"],
        "current_term": settings["current_term"],
        "updated_at": settings["updated_at"],
        "updated_by": settings["updated_by"],
    }
    sync_info = queue_emis_event_safely("academic_settings", "set", sync_payload, entity_key="global")

    return jsonify({
        "success": True,
        "message": "Global academic period updated successfully.",
        "settings": settings,
        "sync": sync_info,
    })


# ============================================================
# EXCEL EXPORT
# ============================================================

@api_bp.route("/api/results/export/excel", methods=["POST"])
def export_results_excel():
    if not can_view_results():
        return jsonify({
            "error": "Unauthorized"
        }), 403

    data = request.get_json(
        silent=True
    ) or {}

    records = data.get(
        "results",
        [],
    )

    if not isinstance(records, list) or not records:
        return jsonify({
            "error": "No results supplied for export"
        }), 400

    filters = (
        data.get("filters", {})
        if isinstance(
            data.get("filters"),
            dict,
        )
        else {}
    )

    generated_at = datetime.now().strftime(
        "%d %B %Y, %I:%M %p"
    )

    def value(row, *keys, default=""):
        for key in keys:
            if row.get(key) not in (None, ""):
                return row.get(key)

        return default

    def export_term(raw, class_level=""):
        label = short_term_label(raw)

        if label:
            return label

        if is_ss_class(class_level):
            return "General / Legacy"

        return ""

    def numeric_score(row):
        raw = value(
            row,
            "Score (%)",
            "Score Number",
            "score",
            "score_percentage",
            "percentage",
            default=0,
        )

        try:
            return float(
                str(raw)
                .replace("%", "")
                .strip()
            )

        except (TypeError, ValueError):
            return 0.0

    def result_status(row):
        explicit = str(
            value(
                row,
                "Status",
                "status",
                default="",
            )
        ).strip().upper()

        if explicit in {
            "PASS",
            "FAIL",
        }:
            return explicit

        return (
            "PASS"
            if numeric_score(row) >= 50
            else "FAIL"
        )

    wb = Workbook()
    ws = wb.active

    ws.title = "Examination Results"
    ws.sheet_view.showGridLines = False

    dark = "17324D"
    teal = "0F766E"
    teal_light = "E7F5F3"
    header_fill = "334155"
    border_colour = "D5DEE8"
    soft_fill = "F8FAFC"
    green = "15803D"
    green_fill = "DCFCE7"
    red = "B91C1C"
    red_fill = "FEE2E2"
    gold = "B7791F"
    gold_fill = "FEF3C7"
    white = "FFFFFF"
    text = "17202A"
    muted = "64748B"

    thin = Side(
        style="thin",
        color=border_colour,
    )

    cell_border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    headers = [
        "S/N",
        "Student Name",
        "Admission No",
        "Year",
        "Class Level",
        "Class Arm",
        "Term",
        "Subject",
        "Score (%)",
        "Correct",
        "Total",
        "Status",
        "Time Taken",
        "Submitted At",
        "Session",
    ]

    total_columns = len(headers)

    # TITLE
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_columns,
    )

    title = ws.cell(
        1,
        1,
        "EMIS CBT — EXAMINATION RESULTS",
    )

    title.font = Font(
        name="Calibri",
        size=18,
        bold=True,
        color=white,
    )

    title.fill = PatternFill(
        "solid",
        fgColor=dark,
    )

    title.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[1].height = 32

    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=total_columns,
    )

    subtitle = ws.cell(
        2,
        1,
        "Official Examination Results Report",
    )

    subtitle.font = Font(
        name="Calibri",
        size=12,
        bold=True,
        color=teal,
    )

    subtitle.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[2].height = 23

    # REPORT INFORMATION
    term_filter_raw = filters.get(
        "term"
    )

    term_filter = (
        short_term_label(
            term_filter_raw
        )
        or str(
            term_filter_raw
            or "All Terms"
        )
    )

    report_info = [
        (
            "Academic Year",
            str(
                filters.get("year")
                or "All Years"
            ),
        ),

        (
            "Class",
            str(
                filters.get("class")
                or "All Classes"
            ),
        ),

        (
            "Term",
            term_filter,
        ),

        (
            "Subject",
            str(
                filters.get("subject")
                or "All Subjects"
            ),
        ),

        (
            "Session",
            str(
                filters.get("session")
                or "All Sessions"
            ),
        ),

        (
            "Generated",
            generated_at,
        ),
    ]

    for index, (
        label,
        content,
    ) in enumerate(report_info):

        column = (
            1
            + (
                index % 3
            )
            * 5
        )

        row = (
            4
            + (
                index // 3
            )
        )

        ws.cell(
            row,
            column,
            label,
        ).font = Font(
            name="Calibri",
            size=10,
            bold=True,
            color=muted,
        )

        ws.cell(
            row,
            column + 1,
            content,
        ).font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=text,
        )

        ws.merge_cells(
            start_row=row,
            start_column=column + 1,
            end_row=row,
            end_column=column + 3,
        )

    # SUMMARY
    scores = [
        numeric_score(row)
        for row in records
    ]

    pass_count = sum(
        1
        for row in records
        if result_status(row) == "PASS"
    )

    average_score = (
        sum(scores) / len(scores)
        if scores
        else 0
    )

    highest_score = (
        max(scores)
        if scores
        else 0
    )

    summary_items = [
        (
            "TOTAL RESULTS",
            len(records),
            teal,
            teal_light,
        ),

        (
            "PASS",
            pass_count,
            green,
            green_fill,
        ),

        (
            "FAIL",
            len(records) - pass_count,
            red,
            red_fill,
        ),

        (
            "AVERAGE SCORE",
            f"{average_score:.1f}%",
            gold,
            gold_fill,
        ),

        (
            "HIGHEST SCORE",
            f"{highest_score:.1f}%",
            teal,
            teal_light,
        ),
    ]

    for index, (
        label,
        summary_value,
        font_colour,
        fill_colour,
    ) in enumerate(summary_items):

        start = (
            1
            + index * 3
        )

        end = min(
            3 + index * 3,
            total_columns,
        )

        ws.merge_cells(
            start_row=7,
            start_column=start,
            end_row=7,
            end_column=end,
        )

        ws.merge_cells(
            start_row=8,
            start_column=start,
            end_row=8,
            end_column=end,
        )

        label_cell = ws.cell(
            7,
            start,
            label,
        )

        value_cell = ws.cell(
            8,
            start,
            summary_value,
        )

        label_cell.font = Font(
            name="Calibri",
            size=9,
            bold=True,
            color=muted,
        )

        value_cell.font = Font(
            name="Calibri",
            size=15,
            bold=True,
            color=font_colour,
        )

        label_cell.fill = PatternFill(
            "solid",
            fgColor=fill_colour,
        )

        value_cell.fill = PatternFill(
            "solid",
            fgColor=fill_colour,
        )

        label_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        value_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for row_num in (7, 8):
            for col_num in range(
                start,
                end + 1,
            ):
                ws.cell(
                    row_num,
                    col_num,
                ).border = cell_border

    ws.row_dimensions[7].height = 19
    ws.row_dimensions[8].height = 28

    # TABLE HEADER
    header_row = 11

    for column, heading in enumerate(
        headers,
        1,
    ):
        cell = ws.cell(
            header_row,
            column,
            heading,
        )

        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=white,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=header_fill,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = cell_border

    ws.row_dimensions[
        header_row
    ].height = 28

    # RESULT ROWS
    for serial, row in enumerate(
        records,
        1,
    ):
        excel_row = (
            header_row
            + serial
        )

        class_level = str(
            value(
                row,
                "Class Level",
                "Class Category",
                "class_level",
                "class_category",
                default="",
            )
        ).strip()

        score = numeric_score(
            row
        )

        status = result_status(
            row
        )

        values = [
            serial,

            str(
                value(
                    row,
                    "Student Name",
                    "full_name",
                    "student_name",
                    default="Unknown Student",
                )
            ).strip(),

            str(
                value(
                    row,
                    "Admission No",
                    "admission_number",
                    "student_id",
                    default="",
                )
            ).strip(),

            str(
                value(
                    row,
                    "Year",
                    "year",
                    default="",
                )
            ).strip(),

            class_level,

            str(
                value(
                    row,
                    "Class Arm",
                    "class_arm",
                    "Class",
                    "class",
                    default=class_level,
                )
            ).strip(),

            export_term(
                value(
                    row,
                    "Term",
                    "term",
                    "Term Label",
                    default="",
                ),
                class_level,
            ),

            str(
                value(
                    row,
                    "Subject",
                    "subject",
                    "Subject Folder",
                    "subject_folder",
                    default="",
                )
            )
            .replace("_", " ")
            .upper()
            .strip(),

            score / 100,

            value(
                row,
                "Correct",
                "correct",
                default=0,
            ),

            value(
                row,
                "Total",
                "total",
                default=0,
            ),

            status,

            value(
                row,
                "Time Taken",
                "time_taken",
                "timeTaken",
                default="",
            ),

            value(
                row,
                "Submitted At",
                "submitted_at",
                "submittedAt",
                default="",
            ),

            value(
                row,
                "Session",
                "session",
                "Academic Session",
                "academic_session",
                default="",
            ),
        ]

        row_fill = PatternFill(
            "solid",
            fgColor=(
                soft_fill
                if serial % 2 == 0
                else white
            ),
        )

        for column, content in enumerate(
            values,
            1,
        ):
            cell = ws.cell(
                excel_row,
                column,
                content,
            )

            cell.font = Font(
                name="Calibri",
                size=11,
                color=text,
            )

            cell.fill = row_fill
            cell.border = cell_border

            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
            )

        # Student name
        ws.cell(
            excel_row,
            2,
        ).font = Font(
            name="Calibri",
            size=12,
            bold=True,
            color=dark,
        )

        # Centre common fields
        for column in [
            1,
            3,
            4,
            5,
            6,
            7,
            10,
            11,
            12,
        ]:
            ws.cell(
                excel_row,
                column,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # Subject
        ws.cell(
            excel_row,
            8,
        ).font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=text,
        )

        # Score
        score_cell = ws.cell(
            excel_row,
            9,
        )

        score_cell.number_format = "0.0%"

        score_cell.font = Font(
            name="Calibri",
            size=13,
            bold=True,
            color=(
                green
                if score >= 50
                else red
            ),
        )

        score_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        score_cell.fill = PatternFill(
            "solid",
            fgColor=(
                green_fill
                if score >= 50
                else red_fill
            ),
        )

        # Status
        status_cell = ws.cell(
            excel_row,
            12,
        )

        status_cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color=(
                green
                if status == "PASS"
                else red
            ),
        )

        status_cell.fill = PatternFill(
            "solid",
            fgColor=(
                green_fill
                if status == "PASS"
                else red_fill
            ),
        )

        ws.row_dimensions[
            excel_row
        ].height = 25

    first_data_row = (
        header_row
        + 1
    )

    last_data_row = (
        header_row
        + len(records)
    )

    table_ref = (
        f"A{header_row}:"
        f"O{last_data_row}"
    )

    table = Table(
        displayName="EMISResultsTable",
        ref=table_ref,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )

    ws.add_table(
        table
    )

    widths = {
        "A": 7,
        "B": 28,
        "C": 16,
        "D": 10,
        "E": 13,
        "F": 15,
        "G": 15,
        "H": 24,
        "I": 13,
        "J": 10,
        "K": 10,
        "L": 12,
        "M": 14,
        "N": 22,
        "O": 16,
    }

    for column, width in widths.items():
        ws.column_dimensions[
            column
        ].width = width

    ws.freeze_panes = (
        f"A{first_data_row}"
    )

    ws.auto_filter.ref = (
        table_ref
    )

    ws.sheet_view.zoomScale = 90
    ws.sheet_view.zoomScaleNormal = 90

    ws.page_setup.paperSize = (
        ws.PAPERSIZE_A4
    )

    ws.page_setup.orientation = (
        ws.ORIENTATION_LANDSCAPE
    )

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.45,
        bottom=0.45,
        header=0.2,
        footer=0.2,
    )

    ws.print_title_rows = (
        f"{header_row}:{header_row}"
    )

    ws.print_area = (
        f"A1:O{last_data_row}"
    )

    ws.oddHeader.center.text = (
        "&BEMIS CBT — EXAMINATION RESULTS"
    )

    ws.oddHeader.center.size = 10

    ws.oddFooter.left.text = "EMIS CBT"
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "Generated &D"

    ws.sheet_properties.outlinePr.summaryBelow = True

    output = BytesIO()

    wb.save(
        output
    )

    output.seek(0)

    def safe_name(value, fallback):
        return (
            str(value or fallback)
            .replace("/", "-")
            .replace(" ", "_")
        )

    filename = (
        f"EMIS_Results_"
        f"{safe_name(filters.get('year'), 'all_years')}_"
        f"{safe_name(filters.get('class'), 'all_classes')}_"
        f"{safe_name(filters.get('term'), 'all_terms')}_"
        f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )