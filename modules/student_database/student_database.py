# MODULE: Student Database Manager — Promotion Manager database and synchronization engine

import csv
import shutil
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook

from modules.class_config import (
    SUPPORTED_CLASSES, CLASS_ARMS_BY_LEVEL, STUDENT_CSV_FILE,
    normalize_class_level, normalize_class_arm, is_valid_class, is_valid_class_arm,
)


BASE_DIR = Path(__file__).resolve().parent.parent.parent
DATABASE_DIR = BASE_DIR / "static" / "data" / "database"
MASTER_CSV = BASE_DIR / STUDENT_CSV_FILE
BACKUP_DIR = DATABASE_DIR / "backups" / "student_database"


# ============================================================
# CLASS EXCEL DATABASES
# ============================================================

CLASS_DATABASE_FILES = {
    "JSS1": DATABASE_DIR / "JSS1_Students.xlsx",
    "JSS2": DATABASE_DIR / "JSS2_Students.xlsx",
    "JSS3": DATABASE_DIR / "JSS3_Students.xlsx",
    "SS1": DATABASE_DIR / "SS1_Students.xlsx",
    "SS2": DATABASE_DIR / "SS2_Students.xlsx",
    "SS3": DATABASE_DIR / "SS3_Students.xlsx",
}


# ============================================================
# ACTIVE STUDENT DATABASE HEADERS
#
# students2026.csv is the authoritative ACTIVE student roster.
#
# Graduated students are archived by graduation_manager.py.
# Deleted / withdrawn students are removed from the active
# roster and their admission numbers are NOT automatically
# released.
# ============================================================

STUDENT_HEADERS = [
    "Admission_number", "Last_name", "First_name", "Other_names",
    "Phone", "Sex", "Class", "Class_category",
]


# ============================================================
# BASIC HELPERS
# ============================================================

def ensure_database_directories():
    DATABASE_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)


def clean(value):
    return str(value or "").strip()


def normalize_admission(value):
    return clean(value).lower()


def admission_sort_key(row):
    value = normalize_admission(row.get("Admission_number"))

    if value.startswith("std") and value[3:].isdigit():
        return 0, int(value[3:])

    return 1, value


def normalize_sex(value):
    value = clean(value).upper()

    if value in {"M", "MALE"}:
        return "Male"

    if value in {"F", "FEMALE"}:
        return "Female"

    return clean(value)


def full_name(row):
    return " ".join(
        part for part in [
            clean(row.get("Last_name")),
            clean(row.get("First_name")),
            clean(row.get("Other_names")),
        ]
        if part
    )


# ============================================================
# CLASS HELPERS
# ============================================================

def normalize_student_level(value, fallback=""):
    level = normalize_class_level(value) or normalize_class_level(fallback)
    return level if level in SUPPORTED_CLASSES else ""


def normalize_student_arm(value, class_level=""):
    level = normalize_student_level(class_level or value)

    if not level:
        return ""

    arm = normalize_class_arm(value, level)

    if arm and arm in CLASS_ARMS_BY_LEVEL.get(level, []):
        return arm

    return ""


def validate_student_class(class_category, class_arm):
    level = normalize_student_level(class_category or class_arm)

    if not level or not is_valid_class(level):
        raise ValueError("Invalid class category.")

    arm = normalize_student_arm(class_arm, level)

    if not arm:
        available = ", ".join(CLASS_ARMS_BY_LEVEL.get(level, []))
        raise ValueError(f"Invalid class arm for {level}. Choose one of: {available}.")

    if not is_valid_class_arm(arm):
        raise ValueError(f"Invalid class arm: {arm}.")

    return level, arm


# ============================================================
# NORMALIZE ACTIVE STUDENT RECORD
#
# IMPORTANT:
# Automatic class-arm assignment does NOT happen here.
#
# promotion_manager.py resolves an automatic class arm before
# calling add_student() / edit_student().
#
# This database layer deliberately remains strict so invalid
# or incomplete class data can never be written to the master
# CSV or Excel databases.
# ============================================================

def normalize_student_record(row, fallback_class=""):
    row = row or {}

    admission = normalize_admission(row.get("Admission_number"))
    raw_category = clean(row.get("Class_category")) or clean(fallback_class)
    raw_class = clean(row.get("Class")) or raw_category
    level = normalize_student_level(raw_category or raw_class, fallback_class)

    if not level:
        raise ValueError(
            f"Invalid class category for student {admission or '(no admission number)'}."
        )

    arm = normalize_student_arm(raw_class, level)

    if not arm:
        available = ", ".join(CLASS_ARMS_BY_LEVEL.get(level, []))
        raise ValueError(
            f"Invalid class arm '{raw_class}' for {level}. Choose one of: {available}."
        )

    return {
        "Admission_number": admission,
        "Last_name": clean(row.get("Last_name")),
        "First_name": clean(row.get("First_name")),
        "Other_names": clean(row.get("Other_names")),
        "Phone": clean(row.get("Phone")),
        "Sex": normalize_sex(row.get("Sex")),
        "Class": arm,
        "Class_category": level,
    }


# ============================================================
# STUDENT RECORD VALIDATION
# ============================================================

def validate_student_record(row, require_admission=True):
    student = normalize_student_record(row)

    if require_admission and not student["Admission_number"]:
        raise ValueError("Admission number is required.")

    if not student["Last_name"]:
        raise ValueError("Last name is required.")

    if not student["First_name"]:
        raise ValueError("First name is required.")

    if student["Sex"] not in {"Male", "Female"}:
        raise ValueError("Sex must be Male or Female.")

    validate_student_class(student["Class_category"], student["Class"])
    return student


# ============================================================
# READ MASTER CSV
# ============================================================

def read_master_students():
    ensure_database_directories()

    if not MASTER_CSV.exists():
        return []

    rows = []

    with MASTER_CSV.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)

        for row in reader:
            admission = normalize_admission(row.get("Admission_number"))

            if not admission:
                continue

            try:
                rows.append(normalize_student_record(row))

            except ValueError:
                # Preserve invalid raw records rather than silently
                # correcting them. database_health_check() can then
                # report exactly what is wrong.
                rows.append({
                    "Admission_number": admission,
                    "Last_name": clean(row.get("Last_name")),
                    "First_name": clean(row.get("First_name")),
                    "Other_names": clean(row.get("Other_names")),
                    "Phone": clean(row.get("Phone")),
                    "Sex": normalize_sex(row.get("Sex")),
                    "Class": clean(row.get("Class")),
                    "Class_category": clean(row.get("Class_category")),
                })

    return rows


# ============================================================
# MASTER DATABASE VALIDATION
# ============================================================

def validate_master_rows(rows):
    normalized_rows = []
    seen = set()
    duplicates = []

    for row in rows:
        admission = normalize_admission(row.get("Admission_number")) or "(no admission number)"

        try:
            student = validate_student_record(row)
        except ValueError as error:
            raise ValueError(f"Student {admission}: {error}") from error

        admission = student["Admission_number"]

        if admission in seen:
            duplicates.append(admission)
        else:
            seen.add(admission)

        normalized_rows.append(student)

    if duplicates:
        duplicate_text = ", ".join(sorted(set(duplicates)))
        raise ValueError(f"Duplicate admission number(s) detected: {duplicate_text}.")

    normalized_rows.sort(key=admission_sort_key)
    return normalized_rows

# ============================================================
# WRITE MASTER CSV
# ============================================================

def write_master_students(rows):
    ensure_database_directories()

    normalized_rows = validate_master_rows(rows)
    temp_path = MASTER_CSV.with_name(f"{MASTER_CSV.stem}.tmp.csv")

    with temp_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(
            file, fieldnames=STUDENT_HEADERS, extrasaction="ignore"
        )

        writer.writeheader()
        writer.writerows(normalized_rows)

    temp_path.replace(MASTER_CSV)
    return normalized_rows


# ============================================================
# READ ONE CLASS EXCEL DATABASE
# ============================================================

def read_class_workbook(class_category):
    level = normalize_student_level(class_category)

    if not level:
        raise ValueError("Invalid class category.")

    path = CLASS_DATABASE_FILES[level]

    if not path.exists():
        return []

    workbook = load_workbook(path, data_only=True, read_only=True)

    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            return []

        headers = [clean(value) for value in rows[0]]
        students = []

        for values in rows[1:]:
            row = {
                headers[index]: values[index] if index < len(values) else ""
                for index in range(len(headers))
            }

            if not normalize_admission(row.get("Admission_number")):
                continue

            students.append(normalize_student_record(row, level))

        return students

    finally:
        workbook.close()


# ============================================================
# WRITE ONE CLASS EXCEL DATABASE
# ============================================================

def write_class_workbook(class_category, rows):
    level = normalize_student_level(class_category)

    if not level:
        raise ValueError("Invalid class category.")

    ensure_database_directories()
    normalized_rows = []

    for row in rows:
        student = validate_student_record(row)

        if student["Class_category"] != level:
            raise ValueError(
                f"Student {student['Admission_number']} does not belong to {level}."
            )

        normalized_rows.append(student)

    normalized_rows.sort(key=admission_sort_key)

    target_path = CLASS_DATABASE_FILES[level]
    temp_path = target_path.with_name(f"{target_path.stem}.tmp.xlsx")

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = f"{level} Students"
    worksheet.append(STUDENT_HEADERS)

    for student in normalized_rows:
        worksheet.append([student.get(header, "") for header in STUDENT_HEADERS])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:H{max(1, len(normalized_rows) + 1)}"

    widths = {
        "A": 18, "B": 22, "C": 22, "D": 24,
        "E": 18, "F": 12, "G": 18, "H": 18,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    workbook.save(temp_path)
    workbook.close()

    temp_path.replace(target_path)
    return normalized_rows


# ============================================================
# GET MASTER STUDENTS FOR ONE CLASS
# ============================================================

def get_students_by_class(class_category, class_arm=""):
    level = normalize_student_level(class_category)

    if not level:
        raise ValueError("Invalid class category.")

    requested_arm = normalize_student_arm(class_arm, level) if clean(class_arm) else ""

    if clean(class_arm) and not requested_arm:
        raise ValueError(f"Invalid class arm for {level}: {class_arm}.")

    rows = []

    for student in read_master_students():
        student_level = normalize_student_level(
            student.get("Class_category") or student.get("Class")
        )

        if student_level != level:
            continue

        if requested_arm:
            student_arm = normalize_student_arm(student.get("Class"), level)

            if student_arm != requested_arm:
                continue

        rows.append(student)

    rows.sort(key=admission_sort_key)
    return rows


# ============================================================
# FIND ACTIVE STUDENT
# ============================================================

def find_student(admission_number):
    target = normalize_admission(admission_number)

    if not target:
        return None

    matches = [
        student
        for student in read_master_students()
        if normalize_admission(student.get("Admission_number")) == target
    ]

    if len(matches) != 1:
        return None

    return matches[0]


# ============================================================
# SYNCHRONIZE ONE CLASS WORKBOOK FROM MASTER CSV
# ============================================================

def sync_class_workbook(class_category, master_rows=None):
    level = normalize_student_level(class_category)

    if not level:
        raise ValueError("Invalid class category.")

    master_rows = read_master_students() if master_rows is None else master_rows

    class_rows = [
        row
        for row in master_rows
        if normalize_student_level(row.get("Class_category") or row.get("Class")) == level
    ]

    return write_class_workbook(level, class_rows)


# ============================================================
# SYNCHRONIZE ALL SIX CLASS WORKBOOKS
# ============================================================

def sync_all_class_workbooks(master_rows=None):
    master_rows = read_master_students() if master_rows is None else master_rows
    result = {}

    for level in SUPPORTED_CLASSES:
        result[level] = len(sync_class_workbook(level, master_rows))

    return result


# ============================================================
# BACKUP TRANSACTION SNAPSHOT
#
# Snapshot contains:
#
#   students2026.csv
#   JSS1_Students.xlsx
#   JSS2_Students.xlsx
#   JSS3_Students.xlsx
#   SS1_Students.xlsx
#   SS2_Students.xlsx
#   SS3_Students.xlsx
# ============================================================

def create_database_backup(reason="transaction"):
    ensure_database_directories()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")

    safe_reason = (
        clean(reason)
        .replace(" ", "_")
        .replace("/", "-")
        .replace("\\", "-")
        or "transaction"
    )

    snapshot_name = f"{timestamp}_{safe_reason}"
    snapshot_dir = BACKUP_DIR / snapshot_name

    snapshot_dir.mkdir(parents=True, exist_ok=False)

    if MASTER_CSV.exists():
        shutil.copy2(MASTER_CSV, snapshot_dir / MASTER_CSV.name)

    for path in CLASS_DATABASE_FILES.values():
        if path.exists():
            shutil.copy2(path, snapshot_dir / path.name)

    return snapshot_dir


# ============================================================
# RESTORE TRANSACTION SNAPSHOT
# ============================================================

def restore_database_backup(snapshot):
    snapshot_dir = Path(snapshot)

    if not snapshot_dir.exists() or not snapshot_dir.is_dir():
        raise ValueError("Backup snapshot does not exist.")

    master_backup = snapshot_dir / MASTER_CSV.name

    if master_backup.exists():
        shutil.copy2(master_backup, MASTER_CSV)
    elif MASTER_CSV.exists():
        MASTER_CSV.unlink()

    for path in CLASS_DATABASE_FILES.values():
        backup_file = snapshot_dir / path.name

        if backup_file.exists():
            shutil.copy2(backup_file, path)
        elif path.exists():
            path.unlink()

    return True


# ============================================================
# VERIFY MASTER CSV AGAINST ALL CLASS XLSX DATABASES
# ============================================================

def verify_database_sync(master_rows=None):
    master_rows = (
        read_master_students()
        if master_rows is None
        else validate_master_rows(master_rows)
    )

    result = {
        "valid": True,
        "classes": {},
        "errors": [],
    }

    for level in SUPPORTED_CLASSES:
        master_class = [
            row
            for row in master_rows
            if normalize_student_level(
                row.get("Class_category") or row.get("Class")
            ) == level
        ]

        try:
            excel_class = read_class_workbook(level)

        except Exception as error:
            result["valid"] = False

            result["classes"][level] = {
                "valid": False,
                "error": str(error),
            }

            result["errors"].append(f"{level}: {error}")
            continue

        master_map = {
            normalize_admission(row.get("Admission_number")): row
            for row in master_class
        }

        excel_map = {
            normalize_admission(row.get("Admission_number")): row
            for row in excel_class
        }

        missing_in_excel = sorted(set(master_map) - set(excel_map))
        extra_in_excel = sorted(set(excel_map) - set(master_map))
        mismatched = []

        for admission in sorted(set(master_map) & set(excel_map)):
            master_student = master_map[admission]
            excel_student = excel_map[admission]

            for field in STUDENT_HEADERS:
                master_value = clean(master_student.get(field))
                excel_value = clean(excel_student.get(field))

                if master_value != excel_value:
                    mismatched.append({
                        "admission_number": admission,
                        "field": field,
                        "master": master_value,
                        "excel": excel_value,
                    })

        class_valid = not missing_in_excel and not extra_in_excel and not mismatched

        if not class_valid:
            result["valid"] = False

        result["classes"][level] = {
            "valid": class_valid,
            "master_count": len(master_class),
            "excel_count": len(excel_class),
            "missing_in_excel": missing_in_excel,
            "extra_in_excel": extra_in_excel,
            "mismatched": mismatched,
        }

    return result


# ============================================================
# MASTER TRANSACTION COMMIT
#
# Transaction sequence:
#
#   1. Validate complete proposed master database.
#   2. Backup master CSV + all six class XLSX files.
#   3. Write students2026.csv.
#   4. Rebuild all six class XLSX files.
#   5. Verify master CSV against all class workbooks.
#   6. Roll back everything if any stage fails.
# ============================================================

def commit_master_transaction(rows, reason="student_database_update"):
    normalized_rows = validate_master_rows(rows)
    snapshot = create_database_backup(reason)

    try:
        written_rows = write_master_students(normalized_rows)
        class_counts = sync_all_class_workbooks(written_rows)
        verification = verify_database_sync(written_rows)

        if not verification["valid"]:
            error_text = "; ".join(verification.get("errors", []))

            if error_text:
                raise RuntimeError(
                    f"Database synchronization verification failed: {error_text}"
                )

            raise RuntimeError("Database synchronization verification failed.")

        return {
            "success": True,
            "backup": snapshot.name,
            "total_students": len(written_rows),
            "class_counts": class_counts,
            "verification": verification,
        }

    except Exception:
        restore_database_backup(snapshot)
        raise


# ============================================================
# ADD STUDENT
#
# IMPORTANT:
#
# Admission number and exact class arm MUST already be resolved
# before this function is called.
#
# promotion_manager.py:
#
#   • assigns automatic class arm when Class is blank
#   • reserves recycled / new admission number
#
# student_database.py then performs strict validation and the
# actual database transaction.
# ============================================================

def add_student(student, reason="add_student"):
    clean_student = validate_student_record(student)
    admission = clean_student["Admission_number"]
    rows = read_master_students()

    duplicate = any(
        normalize_admission(row.get("Admission_number")) == admission
        for row in rows
    )

    if duplicate:
        raise ValueError(
            f"Admission number {admission} already belongs to an active student."
        )

    rows.append(clean_student)

    transaction = commit_master_transaction(rows, f"{reason}_{admission}")

    return {
        "student": clean_student,
        "backup": transaction["backup"],
        "total_students": transaction["total_students"],
        "class_counts": transaction["class_counts"],
        "verification": transaction["verification"],
    }


# ============================================================
# EDIT STUDENT
#
# Admission number is immutable.
#
# Editable fields:
#
#   Last_name
#   First_name
#   Other_names
#   Phone
#   Sex
#   Class
#   Class_category
# ============================================================

def edit_student(original_admission, student, reason="edit_student"):
    original_admission = normalize_admission(original_admission)

    if not original_admission:
        raise ValueError("Original admission number is required.")

    rows = read_master_students()
    existing_index = None

    for index, row in enumerate(rows):
        if normalize_admission(row.get("Admission_number")) == original_admission:
            existing_index = index
            break

    if existing_index is None:
        raise ValueError("Student was not found.")

    incoming_admission = normalize_admission(
        student.get("Admission_number") or original_admission
    )

    if incoming_admission != original_admission:
        raise ValueError("Admission number cannot be changed while editing a student.")

    updated_student = dict(student)
    updated_student["Admission_number"] = original_admission
    updated_student = validate_student_record(updated_student)

    rows[existing_index] = updated_student

    transaction = commit_master_transaction(
        rows, f"{reason}_{original_admission}"
    )

    return {
        "student": updated_student,
        "backup": transaction["backup"],
        "total_students": transaction["total_students"],
        "class_counts": transaction["class_counts"],
        "verification": transaction["verification"],
    }


# ============================================================
# DELETE STUDENTS
#
# Used when a student leaves, transfers or is removed.
#
# IMPORTANT:
#
# Deleting a student DOES NOT release the admission number.
#
# Only successfully graduated SS3 students are released into
# the recycled admission-number pool.
# ============================================================

def delete_students(class_category, admissions, reason="delete_students"):
    level = normalize_student_level(class_category)

    if not level:
        raise ValueError("Invalid class category.")

    admission_set = {
        normalize_admission(value)
        for value in admissions
        if normalize_admission(value)
    }

    if not admission_set:
        raise ValueError("No students selected.")

    rows = read_master_students()
    kept = []
    deleted = []

    for row in rows:
        row_level = normalize_student_level(
            row.get("Class_category") or row.get("Class")
        )

        admission = normalize_admission(row.get("Admission_number"))

        if row_level == level and admission in admission_set:
            deleted.append(row)
        else:
            kept.append(row)

    if not deleted:
        raise ValueError("Selected students were not found in the requested class.")

    deleted_set = {
        normalize_admission(row.get("Admission_number"))
        for row in deleted
    }

    missing = sorted(admission_set - deleted_set)

    if missing:
        raise ValueError(f"Student(s) not found in {level}: {', '.join(missing)}.")

    transaction = commit_master_transaction(kept, f"{reason}_{level}")

    return {
        "deleted": deleted,
        "deleted_admissions": [row["Admission_number"] for row in deleted],
        "count": len(deleted),
        "backup": transaction["backup"],
        "total_students": transaction["total_students"],
        "class_counts": transaction["class_counts"],
        "verification": transaction["verification"],
    }


# ============================================================
# MOVE / PROMOTE STUDENTS
#
# This function performs DATABASE movement only.
#
# promotion_manager.py determines whether the promotion is
# allowed according to:
#
#   JSS1 → JSS2
#   JSS2 → JSS3
#   JSS3 → SS1
#   SS1  → SS2
#   SS2  → SS3
#
# SS3 graduation is handled by graduation_manager.py.
#
# destination_classes example:
#
# {
#     "std001": "JSS2A",
#     "std002": "JSS2B",
# }
# ============================================================

def move_students(source_class, destination_class, admissions, destination_classes, reason="promote_students"):
    source_level = normalize_student_level(source_class)
    destination_level = normalize_student_level(destination_class)

    if not source_level:
        raise ValueError("Invalid source class.")

    if not destination_level:
        raise ValueError("Invalid destination class.")

    if source_level == destination_level:
        raise ValueError(
            "Source and destination classes must be different for a promotion."
        )

    admission_set = {
        normalize_admission(value)
        for value in admissions
        if normalize_admission(value)
    }

    if not admission_set:
        raise ValueError("No students selected.")

    normalized_destination_map = {}

    for admission, class_arm in (destination_classes or {}).items():
        normalized_admission = normalize_admission(admission)

        if not normalized_admission:
            continue

        normalized_arm = normalize_student_arm(class_arm, destination_level)
        normalized_destination_map[normalized_admission] = normalized_arm

    rows = read_master_students()
    moved = []

    available_source_admissions = {
        normalize_admission(row.get("Admission_number"))
        for row in rows
        if normalize_student_level(
            row.get("Class_category") or row.get("Class")
        ) == source_level
    }

    missing = sorted(admission_set - available_source_admissions)

    if missing:
        raise ValueError(
            f"Student(s) not found in {source_level}: {', '.join(missing)}."
        )

    for admission in admission_set:
        destination_arm = normalized_destination_map.get(admission)

        if not destination_arm:
            raise ValueError(
                f"A destination class arm is required for {admission}."
            )

        validate_student_class(destination_level, destination_arm)

    for row in rows:
        admission = normalize_admission(row.get("Admission_number"))
        row_level = normalize_student_level(
            row.get("Class_category") or row.get("Class")
        )

        if admission not in admission_set or row_level != source_level:
            continue

        old_class = clean(row.get("Class"))
        new_class = normalized_destination_map[admission]

        row["Class_category"] = destination_level
        row["Class"] = new_class

        moved.append({
            "Admission_number": admission,
            "Name": full_name(row),
            "From": old_class,
            "To": new_class,
        })

    if not moved:
        raise ValueError("No students were moved.")

    transaction = commit_master_transaction(
        rows,
        f"{reason}_{source_level}_to_{destination_level}",
    )

    return {
        "moved": moved,
        "admissions": [item["Admission_number"] for item in moved],
        "count": len(moved),
        "source_class": source_level,
        "destination_class": destination_level,
        "backup": transaction["backup"],
        "total_students": transaction["total_students"],
        "class_counts": transaction["class_counts"],
        "verification": transaction["verification"],
    }


# ============================================================
# UPDATE STUDENT CLASS DIRECTLY
#
# Intended for controlled administrative corrections.
#
# Normal yearly student movement should always use
# promotion_manager.py.
# ============================================================

def update_student_class(admission_number, destination_class, destination_arm, reason="class_correction"):
    admission = normalize_admission(admission_number)

    destination_level, destination_arm = validate_student_class(
        destination_class, destination_arm
    )

    if not admission:
        raise ValueError("Admission number is required.")

    rows = read_master_students()
    student = None

    for row in rows:
        if normalize_admission(row.get("Admission_number")) == admission:
            student = row
            break

    if not student:
        raise ValueError("Student was not found.")

    source_level = normalize_student_level(
        student.get("Class_category") or student.get("Class")
    )

    if source_level == destination_level:
        updated_student = dict(student)
        updated_student["Class_category"] = destination_level
        updated_student["Class"] = destination_arm

        return edit_student(
            admission,
            updated_student,
            reason=reason,
        )

    return move_students(
        source_class=source_level,
        destination_class=destination_level,
        admissions=[admission],
        destination_classes={admission: destination_arm},
        reason=reason,
    )


# ============================================================
# INITIAL DATABASE SYNCHRONIZATION
#
# students2026.csv is treated as authoritative.
#
# This function rebuilds all six class workbooks from the
# master CSV.
# ============================================================

def initialize_class_databases(create_backup=True):
    ensure_database_directories()
    master_rows = read_master_students()

    if not master_rows:
        raise ValueError("students2026.csv contains no active student records.")

    backup = create_database_backup("initial_sync") if create_backup else None

    try:
        normalized_rows = validate_master_rows(master_rows)
        class_counts = sync_all_class_workbooks(normalized_rows)
        verification = verify_database_sync(normalized_rows)

        if not verification["valid"]:
            raise RuntimeError(
                "Initial class database synchronization failed verification."
            )

        return {
            "success": True,
            "master_students": len(normalized_rows),
            "class_counts": class_counts,
            "backup": backup.name if backup else None,
            "verification": verification,
        }

    except Exception:
        if backup:
            restore_database_backup(backup)

        raise


# ============================================================
# DATABASE SUMMARY
# ============================================================

def get_database_summary():
    rows = read_master_students()
    summary = {}

    for level in SUPPORTED_CLASSES:
        level_rows = [
            row
            for row in rows
            if normalize_student_level(
                row.get("Class_category") or row.get("Class")
            ) == level
        ]

        arms = {}

        for arm in CLASS_ARMS_BY_LEVEL.get(level, []):
            arms[arm] = sum(
                1
                for row in level_rows
                if normalize_student_arm(row.get("Class"), level) == arm
            )

        summary[level] = {
            "total": len(level_rows),
            "arms": arms,
            "available_arms": list(CLASS_ARMS_BY_LEVEL.get(level, [])),
        }

    return summary


# ============================================================
# DATABASE HEALTH CHECK
#
# Checks:
#
#   • master CSV exists
#   • duplicate admission numbers
#   • invalid student records
#   • all six class XLSX files
#   • CSV ↔ XLSX synchronization
# ============================================================

def database_health_check():
    ensure_database_directories()

    result = {
        "healthy": True,
        "master_csv_exists": MASTER_CSV.exists(),
        "master_total": 0,
        "duplicate_admissions": [],
        "invalid_students": [],
        "class_files": {},
        "sync": None,
    }

    if not MASTER_CSV.exists():
        result["healthy"] = False
        return result

    with MASTER_CSV.open("r", newline="", encoding="utf-8-sig") as file:
        raw_rows = list(csv.DictReader(file))

    admission_counts = {}

    for row in raw_rows:
        admission = normalize_admission(row.get("Admission_number"))

        if admission:
            admission_counts[admission] = admission_counts.get(admission, 0) + 1

        try:
            validate_student_record(row)

        except Exception as error:
            result["invalid_students"].append({
                "admission_number": admission,
                "error": str(error),
            })

    result["master_total"] = len([
        row for row in raw_rows
        if normalize_admission(row.get("Admission_number"))
    ])

    result["duplicate_admissions"] = sorted(
        admission
        for admission, count in admission_counts.items()
        if count > 1
    )

    for level, path in CLASS_DATABASE_FILES.items():
        result["class_files"][level] = {
            "exists": path.exists(),
            "path": str(path),
        }

        if not path.exists():
            result["healthy"] = False

    if result["duplicate_admissions"] or result["invalid_students"]:
        result["healthy"] = False

    try:
        result["sync"] = verify_database_sync()

        if not result["sync"]["valid"]:
            result["healthy"] = False

    except Exception as error:
        result["healthy"] = False
        result["sync"] = {
            "valid": False,
            "error": str(error),
        }

    return result