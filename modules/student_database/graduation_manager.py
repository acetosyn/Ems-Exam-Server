# MODULE: Graduation Manager — Archives graduating SS3 students and releases completed admission numbers

import csv
import json
import shutil
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook

from modules.class_config import normalize_class_level

from modules.student_database.student_database import (
    read_master_students, commit_master_transaction, normalize_admission, full_name,
)

from modules.student_database.admission_manager import release_admission_number


BASE_DIR = Path(__file__).resolve().parent.parent.parent
DATABASE_DIR = BASE_DIR / "static" / "data" / "database"

GRADUATES_DIR = DATABASE_DIR / "graduates"
GRADUATION_LOG_FILE = GRADUATES_DIR / "graduation_log.csv"
GRADUATE_BACKUP_DIR = DATABASE_DIR / "backups" / "graduates"

GRADUATE_HEADERS = [
    "Admission_number", "Last_name", "First_name", "Other_names", "Phone", "Sex",
    "Class", "Class_category", "Graduation_date", "Graduation_year", "Academic_session",
]

GRADUATION_LOG_HEADERS = [
    "Timestamp", "Admission_number", "Student_name", "Class", "Graduation_date",
    "Graduation_year", "Academic_session", "Archive_file",
]


# ============================================================
# BASIC HELPERS
# ============================================================

def ensure_graduation_directories():
    GRADUATES_DIR.mkdir(parents=True, exist_ok=True)


def clean(value):
    return str(value or "").strip()


def normalize_session(value):
    value = clean(value)

    if value:
        return value

    current_year = datetime.now().year
    return f"{current_year - 1}/{current_year}"


def normalize_graduation_date(value=None):
    if not value:
        return datetime.now().strftime("%Y-%m-%d")

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    value = clean(value)

    for format_string in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, format_string).strftime("%Y-%m-%d")
        except ValueError:
            continue

    raise ValueError("Invalid graduation date. Use YYYY-MM-DD.")


def get_graduation_year(graduation_date=None):
    if not graduation_date:
        return datetime.now().year

    if isinstance(graduation_date, datetime):
        return graduation_date.year

    normalized_date = normalize_graduation_date(graduation_date)
    return datetime.strptime(normalized_date, "%Y-%m-%d").year


# ============================================================
# ARCHIVE PATH
# ============================================================

def get_graduation_workbook_path(year):
    year = int(year)

    year_dir = GRADUATES_DIR / str(year)
    year_dir.mkdir(parents=True, exist_ok=True)

    return year_dir / f"graduated_students_{year}.xlsx"


# ============================================================
# READ GRADUATES
# ============================================================

def read_graduates(year=None):
    ensure_graduation_directories()
    files = []

    if year:
        path = get_graduation_workbook_path(year)

        if path.exists():
            files.append(path)
    else:
        files = sorted(GRADUATES_DIR.glob("*/graduated_students_*.xlsx"))

    graduates = []

    for path in files:
        workbook = load_workbook(path, data_only=True, read_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            workbook.close()
            continue

        headers = [clean(value) for value in rows[0]]

        for values in rows[1:]:
            row = {
                headers[index]: values[index] if index < len(values) else ""
                for index in range(len(headers))
            }

            admission = normalize_admission(row.get("Admission_number"))

            if not admission:
                continue

            row["Admission_number"] = admission
            graduates.append(row)

        workbook.close()

    graduates.sort(key=lambda row: clean(row.get("Graduation_date")), reverse=True)
    return graduates


# ============================================================
# CHECK GRADUATE
# ============================================================

def graduate_exists(admission_number):
    admission = normalize_admission(admission_number)

    if not admission:
        return False

    return any(
        normalize_admission(row.get("Admission_number")) == admission
        for row in read_graduates()
    )


# ============================================================
# WRITE GRADUATION WORKBOOK
# ============================================================

def write_graduation_workbook(year, rows):
    path = get_graduation_workbook_path(year)
    temp_path = path.with_name(f"{path.stem}.tmp.xlsx")

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = f"{year} Graduates"
    worksheet.append(GRADUATE_HEADERS)

    for row in rows:
        worksheet.append([row.get(header, "") for header in GRADUATE_HEADERS])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:K{max(1, len(rows) + 1)}"

    widths = {
        "A": 18, "B": 22, "C": 22, "D": 24, "E": 18, "F": 12,
        "G": 18, "H": 18, "I": 18, "J": 18, "K": 20,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    workbook.save(temp_path)
    workbook.close()

    temp_path.replace(path)
    return path


# ============================================================
# APPEND GRADUATES TO ARCHIVE
# ============================================================

def append_graduates_to_archive(students, graduation_date=None, academic_session=""):
    graduation_date = normalize_graduation_date(graduation_date)
    graduation_year = get_graduation_year(graduation_date)
    academic_session = normalize_session(academic_session)

    existing = read_graduates(graduation_year)

    existing_admissions = {
        normalize_admission(row.get("Admission_number"))
        for row in existing
    }

    archived = []

    for student in students:
        admission = normalize_admission(student.get("Admission_number"))

        if not admission:
            raise ValueError("Graduate admission number is required.")

        if admission in existing_admissions:
            raise ValueError(
                f"Student {admission} already exists in the "
                f"{graduation_year} graduation archive."
            )

        record = {
            "Admission_number": admission,
            "Last_name": clean(student.get("Last_name")),
            "First_name": clean(student.get("First_name")),
            "Other_names": clean(student.get("Other_names")),
            "Phone": clean(student.get("Phone")),
            "Sex": clean(student.get("Sex")),
            "Class": clean(student.get("Class")),
            "Class_category": clean(student.get("Class_category")),
            "Graduation_date": graduation_date,
            "Graduation_year": graduation_year,
            "Academic_session": academic_session,
        }

        existing.append(record)
        existing_admissions.add(admission)
        archived.append(record)

    existing.sort(key=lambda row: normalize_admission(row.get("Admission_number")))

    archive_path = write_graduation_workbook(graduation_year, existing)

    return {
        "archived": archived,
        "count": len(archived),
        "graduation_year": graduation_year,
        "graduation_date": graduation_date,
        "academic_session": academic_session,
        "archive_path": archive_path,
    }


# ============================================================
# GRADUATION ACTIVITY LOG
# ============================================================

def append_graduation_log(records, archive_path):
    ensure_graduation_directories()
    write_header = not GRADUATION_LOG_FILE.exists()

    with GRADUATION_LOG_FILE.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=GRADUATION_LOG_HEADERS)

        if write_header:
            writer.writeheader()

        for record in records:
            writer.writerow({
                "Timestamp": datetime.now().isoformat(timespec="seconds"),
                "Admission_number": normalize_admission(record.get("Admission_number")),
                "Student_name": full_name(record),
                "Class": clean(record.get("Class")),
                "Graduation_date": clean(record.get("Graduation_date")),
                "Graduation_year": clean(record.get("Graduation_year")),
                "Academic_session": clean(record.get("Academic_session")),
                "Archive_file": str(archive_path),
            })


# ============================================================
# VALIDATE STUDENTS FOR GRADUATION
# ============================================================

def validate_graduation_students(admissions):
    admission_set = {
        normalize_admission(value)
        for value in admissions
        if normalize_admission(value)
    }

    if not admission_set:
        raise ValueError("No students selected for graduation.")

    master_rows = read_master_students()
    selected = []

    for student in master_rows:
        admission = normalize_admission(student.get("Admission_number"))

        if admission not in admission_set:
            continue

        class_level = normalize_class_level(
            student.get("Class_category") or student.get("Class")
        )

        if class_level != "SS3":
            raise ValueError(f"Student {admission} is not in SS3 and cannot be graduated.")

        selected.append(student)

    found = {
        normalize_admission(student.get("Admission_number"))
        for student in selected
    }

    missing = sorted(admission_set - found)

    if missing:
        raise ValueError(
            "Student(s) not found in active "
            f"SS3 database: {', '.join(missing)}."
        )

    return selected


# ============================================================
# GRADUATE STUDENTS
#
# Sequence:
#
# 1. Validate selected students are active SS3 students.
# 2. Archive them in the graduation workbook.
# 3. Remove them from students2026.csv.
# 4. Synchronize all six class Excel databases.
# 5. Verify database synchronization.
# 6. Release their admission numbers.
# 7. Write graduation activity log.
#
# IMPORTANT:
# Admission numbers are released ONLY after successful
# database removal.
# ============================================================

def graduate_students(admissions, graduation_date=None, academic_session=""):
    students = validate_graduation_students(admissions)

    graduation_date = normalize_graduation_date(graduation_date)
    graduation_year = get_graduation_year(graduation_date)
    academic_session = normalize_session(academic_session)

    archive_path = get_graduation_workbook_path(graduation_year)
    archive_existed_before = archive_path.exists()
    archive_backup = None

    if archive_existed_before:
        archive_backup = archive_path.with_name(
            f"{archive_path.stem}.before_graduation.xlsx"
        )
        archive_backup.write_bytes(archive_path.read_bytes())

    archive_result = None
    transaction = None
    transaction_completed = False

    try:
        # ----------------------------------------------------
        # 1. ARCHIVE STUDENTS
        # ----------------------------------------------------
        archive_result = append_graduates_to_archive(
            students=students,
            graduation_date=graduation_date,
            academic_session=academic_session,
        )

        # ----------------------------------------------------
        # 2. REMOVE FROM ACTIVE DATABASE
        # ----------------------------------------------------
        graduating_admissions = {
            normalize_admission(student.get("Admission_number"))
            for student in students
        }

        remaining_students = [
            student
            for student in read_master_students()
            if normalize_admission(student.get("Admission_number"))
            not in graduating_admissions
        ]

        # ----------------------------------------------------
        # 3. MASTER DATABASE TRANSACTION
        #
        # This updates students2026.csv, rebuilds the six
        # class XLSX files and verifies synchronization.
        # ----------------------------------------------------
        transaction = commit_master_transaction(
            remaining_students,
            f"graduate_{graduation_year}",
        )

        transaction_completed = True

        # ----------------------------------------------------
        # 4. RELEASE ADMISSION NUMBERS
        #
        # Only completed SS3 graduates enter the recycled
        # admission-number pool.
        # ----------------------------------------------------
        released_numbers = []

        for student in students:
            admission = normalize_admission(student.get("Admission_number"))
            release_result = release_admission_number(admission)
            released_admission = normalize_admission(
                release_result.get("admission_number")
            )

            if not released_admission:
                raise RuntimeError(
                    f"Admission number release failed for {admission}."
                )

            released_numbers.append(released_admission)

        # ----------------------------------------------------
        # 5. GRADUATION LOG
        # ----------------------------------------------------
        append_graduation_log(
            archive_result["archived"],
            archive_result["archive_path"],
        )

        # ----------------------------------------------------
        # 6. REMOVE TEMPORARY ARCHIVE BACKUP
        # ----------------------------------------------------
        if archive_backup and archive_backup.exists():
            archive_backup.unlink()

        return {
            "success": True,
            "graduated": archive_result["archived"],
            "count": len(archive_result["archived"]),
            "graduation_date": graduation_date,
            "graduation_year": graduation_year,
            "academic_session": academic_session,
            "released_admission_numbers": released_numbers,
            "archive_file": str(archive_result["archive_path"]),
            "database_backup": transaction["backup"],
            "total_active_students": transaction["total_students"],
        }

    except Exception:
        # ----------------------------------------------------
        # DATABASE TRANSACTION FAILED
        #
        # commit_master_transaction() already restores the
        # active CSV/XLSX database itself when its internal
        # transaction fails.
        #
        # Therefore restore the graduation archive only when
        # the active database transaction was not completed.
        # ----------------------------------------------------
        if not transaction_completed:
            if archive_backup and archive_backup.exists():
                archive_backup.replace(archive_path)

            elif archive_path.exists() and not archive_existed_before:
                archive_path.unlink()

        raise


# ============================================================
# GET ONE GRADUATE
# ============================================================

def get_graduate_by_admission(admission_number):
    admission = normalize_admission(admission_number)

    if not admission:
        return None

    matches = [
        graduate
        for graduate in read_graduates()
        if normalize_admission(graduate.get("Admission_number")) == admission
    ]

    if len(matches) != 1:
        return None

    return matches[0]


# ============================================================
# SEARCH GRADUATES
# ============================================================

def search_graduates(search="", year=None):
    search = clean(search).lower()
    graduates = read_graduates(year)

    if not search:
        return graduates

    result = []

    for graduate in graduates:
        values = [
            graduate.get("Admission_number"),
            graduate.get("Last_name"),
            graduate.get("First_name"),
            graduate.get("Other_names"),
            graduate.get("Class"),
            graduate.get("Academic_session"),
            graduate.get("Graduation_year"),
        ]

        searchable = " ".join(clean(value).lower() for value in values)

        if search in searchable:
            result.append(graduate)

    return result


# ============================================================
# GRADUATION SUMMARY
# ============================================================

def get_graduation_summary():
    ensure_graduation_directories()

    graduates = read_graduates()
    years = {}

    for graduate in graduates:
        year = clean(graduate.get("Graduation_year"))

        if not year:
            continue

        years[year] = years.get(year, 0) + 1

    return {
        "total_graduates": len(graduates),
        "years": dict(sorted(years.items(), reverse=True)),
        "latest_year": max(years.keys()) if years else None,
    }

# ============================================================
# CLEAR GRADUATE ARCHIVE
#
# IMPORTANT SAFETY RULES:
#   • Clearing the archive NEVER restores students to SS3.
#   • Clearing the archive NEVER releases admission numbers.
#   • A verified filesystem backup is created before deletion.
#   • The graduation activity log is preserved.
# ============================================================

def _graduate_archive_files(year=None):
    ensure_graduation_directories()

    if year is not None:
        year = int(year)
        path = GRADUATES_DIR / str(year) / f"graduated_students_{year}.xlsx"
        return [path] if path.exists() else []

    return sorted(GRADUATES_DIR.glob("*/graduated_students_*.xlsx"))


def _verified_copy(source, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)

    if not destination.exists() or destination.stat().st_size != source.stat().st_size:
        raise IOError(f"Graduate archive backup verification failed for {source.name}.")

    return destination


def clear_graduate_archive(year=None, requested_by="Admin"):
    """
    Remove one graduation year or the full visible graduate archive.

    This function intentionally touches ONLY the archived graduate workbook(s).
    It does not edit students2026.csv, class XLSX files, the recycled admission
    number pool, or the historical graduation_log.csv.
    """
    ensure_graduation_directories()

    normalized_year = None
    if year not in (None, "", "all", "ALL"):
        try:
            normalized_year = int(year)
        except (TypeError, ValueError):
            raise ValueError("Invalid graduation year.")

        if normalized_year < 1900 or normalized_year > 9999:
            raise ValueError("Invalid graduation year.")

    archive_files = _graduate_archive_files(normalized_year)

    if not archive_files:
        scope = f"{normalized_year}" if normalized_year else "the graduate archive"
        raise ValueError(f"No archived graduates were found for {scope}.")

    records = read_graduates(normalized_year) if normalized_year else read_graduates()
    admissions = [
        normalize_admission(row.get("Admission_number"))
        for row in records
        if normalize_admission(row.get("Admission_number"))
    ]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    scope_slug = str(normalized_year) if normalized_year else "all_years"
    backup_dir = GRADUATE_BACKUP_DIR / f"archive_clear_{scope_slug}_{timestamp}"
    backup_dir.mkdir(parents=True, exist_ok=False)

    backed_up_files = []

    try:
        for source in archive_files:
            relative = source.relative_to(GRADUATES_DIR)
            destination = backup_dir / relative
            _verified_copy(source, destination)
            backed_up_files.append(str(destination))

        if GRADUATION_LOG_FILE.exists():
            log_backup = backup_dir / GRADUATION_LOG_FILE.name
            _verified_copy(GRADUATION_LOG_FILE, log_backup)
            backed_up_files.append(str(log_backup))

        manifest = {
            "action": "CLEAR_GRADUATE_ARCHIVE",
            "scope": "year" if normalized_year else "all_years",
            "year": normalized_year,
            "requested_by": clean(requested_by) or "Admin",
            "cleared_at": datetime.now().isoformat(timespec="seconds"),
            "record_count": len(records),
            "admissions": admissions,
            "source_files": [str(path) for path in archive_files],
            "backup_files": backed_up_files,
            "graduation_log_preserved": True,
            "active_student_database_changed": False,
            "admission_number_pool_changed": False,
        }
        manifest_path = backup_dir / "clear_manifest.json"
        manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

        # Delete only after every requested archive file has a verified copy.
        for source in archive_files:
            source.unlink()

            parent = source.parent
            try:
                if parent != GRADUATES_DIR and not any(parent.iterdir()):
                    parent.rmdir()
            except OSError:
                pass

    except Exception:
        # The source archives are not modified until backups finish, so a backup
        # failure leaves the live archive untouched. If deletion failed midway,
        # restore any missing source workbook from the verified backup.
        for source in archive_files:
            if source.exists():
                continue
            relative = source.relative_to(GRADUATES_DIR)
            backup_source = backup_dir / relative
            if backup_source.exists():
                source.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(backup_source, source)
        raise

    return {
        "success": True,
        "scope": "year" if normalized_year else "all_years",
        "year": normalized_year,
        "count": len(records),
        "admissions": admissions,
        "backup_directory": str(backup_dir),
        "backup_manifest": str(backup_dir / "clear_manifest.json"),
        "cleared_files": [str(path) for path in archive_files],
        "graduation_log_preserved": True,
        "active_student_database_changed": False,
        "admission_number_pool_changed": False,
    }

