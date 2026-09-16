# MODULE: Smart Student Import Manager — CSV/XLSX/DOCX/PDF/TXT extraction, normalization, validation and transactional import

import csv
import io
import re
import unicodedata
from difflib import SequenceMatcher
from openpyxl import load_workbook

from modules.class_config import CLASS_ARMS_BY_LEVEL, normalize_class_level
from modules.student_database.student_database import (
    STUDENT_HEADERS, read_master_students, commit_master_transaction, normalize_admission,
    normalize_student_level, normalize_student_arm, normalize_sex, validate_student_class, full_name,
)
from modules.student_database.admission_manager import load_available_numbers, save_available_numbers, get_highest_known_admission_number, format_admission_number

MAX_IMPORT_ROWS = 2500
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".docx", ".pdf", ".txt", ".text", ".md"}
CANONICAL = list(STUDENT_HEADERS)

HEADER_ALIASES = {
    "Admission_number": ["admission number","admission no","admission","admissionnumber","admissionno","adm no","adm number","student id","studentid","id"],
    "Last_name": ["last name","lastname","surname","family name","familyname","last_name","last"],
    "First_name": ["first name","firstname","given name","givenname","forename","first_name","first"],
    "Other_names": ["other names","other name","othernames","middle name","middle names","middlename","middlenames","other_names"],
    "Phone": ["phone","phone number","phonenumber","mobile","mobile number","mobilenumber","telephone","tel","contact","contact number"],
    "Sex": ["sex","gender"],
    "Class": ["class","class arm","classarm","arm","class_arm","student class"],
    "Class_category": ["class category","classcategory","class level","classlevel","level","class_category","category"],
}


def clean(value):
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def normalize_header(value):
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "", value)


ALIAS_INDEX = {canonical:{normalize_header(alias) for alias in aliases + [canonical]} for canonical, aliases in HEADER_ALIASES.items()}


def deterministic_header_mapping(headers):
    headers = [clean(h) for h in headers]
    normalized = {header:normalize_header(header) for header in headers if clean(header)}
    mapping, confidence, used = {}, {}, set()

    # Disambiguate the very common "Class" + "Class Arm" pair before the generic mapper.
    # In that layout, "Class" normally means the level/category and "Class Arm" is the exact arm.
    arm_preferred = next((header for header, key in normalized.items() if key in {"classarm","arm","studentclassarm"}), None)
    category_explicit = next((header for header, key in normalized.items() if key in {"classcategory","classlevel","level","category"}), None)
    plain_class = next((header for header, key in normalized.items() if key == "class"), None)
    if arm_preferred:
        mapping["Class"] = arm_preferred; confidence["Class"] = 1.0; used.add(arm_preferred)
        category_source = category_explicit or plain_class
        if category_source and category_source not in used:
            mapping["Class_category"] = category_source; confidence["Class_category"] = 1.0; used.add(category_source)
    elif category_explicit:
        mapping["Class_category"] = category_explicit; confidence["Class_category"] = 1.0; used.add(category_explicit)

    for canonical, aliases in ALIAS_INDEX.items():
        if canonical in mapping: continue
        exact = next((header for header, key in normalized.items() if key in aliases and header not in used), None)
        if exact:
            mapping[canonical] = exact; confidence[canonical] = 1.0; used.add(exact); continue
        best_header, best_score = "", 0.0
        for header, key in normalized.items():
            if header in used or not key: continue
            score = max((SequenceMatcher(None, key, alias).ratio() for alias in aliases), default=0.0)
            if score > best_score: best_header, best_score = header, score
        if best_header and best_score >= 0.84:
            mapping[canonical] = best_header; confidence[canonical] = round(best_score, 3); used.add(best_header)
    return mapping, confidence


def _matrix_to_records(matrix):
    rows = [[clean(v) for v in row] for row in matrix if any(clean(v) for v in row)]
    if not rows: return [], []
    headers = rows[0]
    while headers and not clean(headers[-1]): headers.pop()
    records = []
    for values in rows[1:MAX_IMPORT_ROWS + 1]:
        if not any(clean(v) for v in values): continue
        record = {headers[i]:clean(values[i]) if i < len(values) else "" for i in range(len(headers)) if clean(headers[i])}
        if record: records.append(record)
    return headers, records


def _parse_delimited_text(text):
    text = clean(text)
    if not text: return [], []
    lines = [line for line in text.splitlines() if clean(line)]
    if not lines: return [], []
    candidates = [",", "\t", ";", "|"]
    delimiter = max(candidates, key=lambda d: lines[0].count(d))
    if lines[0].count(delimiter) < 1: return [], []
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    return _matrix_to_records(list(reader))


def _extract_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:6000], delimiters=",\t;|")
        reader = csv.reader(io.StringIO(text), dialect)
        return _matrix_to_records(list(reader)), text
    except Exception:
        return _parse_delimited_text(text), text


def _extract_xlsx(data):
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        matrix = list(worksheet.iter_rows(values_only=True))[:MAX_IMPORT_ROWS + 20]
        return _matrix_to_records(matrix), ""
    finally: workbook.close()


def _extract_docx(data):
    try:
        from docx import Document
    except Exception as error: raise ValueError(f"DOCX import requires python-docx: {error}")
    doc = Document(io.BytesIO(data))
    best = None
    for table in doc.tables:
        matrix = [[clean(cell.text) for cell in row.cells] for row in table.rows]
        if len(matrix) >= 2 and (best is None or len(matrix) * max(1,len(matrix[0])) > len(best) * max(1,len(best[0]))): best = matrix
    text = "\n".join(clean(p.text) for p in doc.paragraphs if clean(p.text))
    if best: return _matrix_to_records(best), text
    parsed = _parse_delimited_text(text)
    return parsed, text


def _extract_pdf(data):
    reader = None
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
        except Exception as error: raise ValueError(f"PDF import requires pypdf (or PyPDF2): {error}")
    text = "\n".join(clean(page.extract_text()) for page in reader.pages if clean(page.extract_text()))
    if not text: raise ValueError("No selectable text was found in this PDF. Scanned/image-only PDFs need OCR before import.")
    return _parse_delimited_text(text), text


def extract_upload(file_bytes, filename):
    match = re.search(r"(\.[A-Za-z0-9]+)$", clean(filename)); ext = match.group(1).lower() if match else ""
    if ext not in SUPPORTED_EXTENSIONS: raise ValueError("Unsupported file type. Use CSV, XLSX/XLSM, DOCX, text/notes, Markdown or a text-based PDF.")
    if ext == ".csv": (headers, records), text = _extract_csv(file_bytes)
    elif ext in {".xlsx", ".xlsm"}: (headers, records), text = _extract_xlsx(file_bytes)
    elif ext == ".docx": (headers, records), text = _extract_docx(file_bytes)
    elif ext == ".pdf": (headers, records), text = _extract_pdf(file_bytes)
    else:
        text = file_bytes.decode("utf-8-sig", errors="replace"); headers, records = _parse_delimited_text(text)
    return {"extension":ext, "headers":headers, "records":records[:MAX_IMPORT_ROWS], "text":text[:70000]}


def _canonical_row(source, mapping): return {field:clean(source.get(mapping.get(field,""))) if mapping.get(field) else "" for field in CANONICAL}


def _normalize_phone(value): return re.sub(r"\s+", "", clean(value))


def _normalize_sex(value):
    normalized = normalize_sex(value)
    return normalized if normalized in {"Male","Female"} else ""


def _resolve_class(row, default_level="", default_arm=""):
    raw_level = clean(row.get("Class_category")); raw_arm = clean(row.get("Class")); default_level = clean(default_level); default_arm = clean(default_arm)
    level = normalize_student_level(raw_level or raw_arm or default_level)
    auto_keys = {"AUTO","AUTOMATIC","BALANCED","AUTO_SCIENCE","SCIENCE_AUTO","AUTO_ART_COMMERCIAL","AUTO_ARTS_COMMERCIAL","ART_COMMERCIAL_AUTO"}
    candidate_auto = (raw_arm or default_arm).upper().replace(" ", "_")
    auto_arm = "AUTO_SCIENCE" if candidate_auto in {"AUTO_SCIENCE","SCIENCE_AUTO"} else "AUTO_ART_COMMERCIAL" if candidate_auto in {"AUTO_ART_COMMERCIAL","AUTO_ARTS_COMMERCIAL","ART_COMMERCIAL_AUTO"} else "AUTO" if candidate_auto in {"AUTO","AUTOMATIC","BALANCED"} else ""
    arm = normalize_student_arm(raw_arm, level) if level and raw_arm and not auto_arm else ""
    if not arm and level and default_arm and not auto_arm: arm = normalize_student_arm(default_arm, level)
    return level, arm, auto_arm


def _existing_duplicate_index():
    result = set()
    for student in read_master_students():
        key = (normalize_header(student.get("Last_name")), normalize_header(student.get("First_name")), normalize_header(student.get("Other_names")), re.sub(r"\D", "", clean(student.get("Phone"))))
        result.add(key)
    return result


def analyze_student_upload(file_bytes, filename, default_class_category="", default_class_arm="", use_ai=True):
    extracted = extract_upload(file_bytes, filename); headers, records, text = extracted["headers"], extracted["records"], extracted["text"]
    mapping, confidence = deterministic_header_mapping(headers); ai_meta = {"provider":"Rules only", "model":"", "note":"Deterministic header mapping used."}
    required_missing = [field for field in ["Last_name","First_name"] if field not in mapping]
    unresolved_useful = len(mapping) < min(5, len(headers)) if headers else True
    if use_ai and headers and (required_missing or unresolved_useful):
        try:
            from modules.student_database.promotion_ai import map_student_headers
            ai_result = map_student_headers(headers, records[:5]); used_sources = set(mapping.values())
            for canonical, source in (ai_result.get("mapping") or {}).items():
                if canonical not in mapping and source not in used_sources: mapping[canonical] = source; used_sources.add(source)
            ai_meta = {"provider":ai_result.get("provider","Rules only"), "model":ai_result.get("model",""), "note":ai_result.get("note","")}
        except Exception as error: ai_meta = {"provider":"Rules only", "model":"", "note":f"AI mapping unavailable; rules were used. {error}"}
    if not records and use_ai and text:
        try:
            from modules.student_database.promotion_ai import parse_unstructured_student_text
            ai_result = parse_unstructured_student_text(text); records = ai_result.get("rows") or []
            if records:
                headers = CANONICAL; mapping = {field:field for field in CANONICAL}; ai_meta = {"provider":ai_result.get("provider","Rules only"), "model":ai_result.get("model",""), "note":ai_result.get("note","")}
        except Exception as error: ai_meta = {"provider":"Rules only", "model":"", "note":f"Unstructured AI extraction unavailable. {error}"}
    if not records: raise ValueError("No student rows could be extracted. For PDF/DOCX notes without a table, enable AI Assist or convert the list to CSV/XLSX.")

    existing_keys = _existing_duplicate_index(); analyzed = []; source_seen = set(); stats = {"rows":0,"valid":0,"needs_attention":0,"possible_duplicates":0}
    for index, source in enumerate(records[:MAX_IMPORT_ROWS], start=1):
        canonical = _canonical_row(source, mapping); canonical["Phone"] = _normalize_phone(canonical.get("Phone")); canonical["Sex"] = _normalize_sex(canonical.get("Sex"))
        level, arm, auto_arm = _resolve_class(canonical, default_class_category, default_class_arm); canonical["Class_category"] = level; canonical["Class"] = arm or (auto_arm if level and auto_arm else "")
        errors, warnings = [], []
        if not canonical["Last_name"]: errors.append("Last name is required.")
        if not canonical["First_name"]: errors.append("First name is required.")
        if canonical["Sex"] not in {"Male","Female"}: errors.append("Sex must be assigned as Male or Female.")
        if not level: errors.append("Class level must be assigned.")
        if level and not canonical["Class"]: errors.append("Class arm must be assigned or set to Automatic.")
        if level and arm:
            try: validate_student_class(level, arm)
            except ValueError as error: errors.append(str(error))
        if level.startswith("SS") and canonical.get("Class") == "AUTO": warnings.append("Automatic / Balanced may place this student in either a Science or Arts / Commercial arm. Choose an exact arm or stream-specific automatic option if the stream matters.")
        duplicate_key = (normalize_header(canonical.get("Last_name")), normalize_header(canonical.get("First_name")), normalize_header(canonical.get("Other_names")), re.sub(r"\D", "", canonical.get("Phone","")))
        if duplicate_key in source_seen: warnings.append("Possible duplicate within this upload.")
        source_seen.add(duplicate_key)
        if duplicate_key in existing_keys and any(duplicate_key): warnings.append("A similar active student already exists; review before importing.")
        if "?" in canonical.get("Phone",""): warnings.append("Phone contains unknown characters; it will be preserved for manual correction.")
        canonical["_row"] = index; canonical["_valid"] = not errors; canonical["_errors"] = errors; canonical["_warnings"] = warnings; canonical["_include"] = True
        analyzed.append(canonical); stats["rows"] += 1; stats["valid"] += 1 if not errors else 0; stats["needs_attention"] += 1 if errors else 0; stats["possible_duplicates"] += 1 if warnings else 0
    return {"filename":clean(filename), "extension":extracted["extension"], "source_headers":headers, "mapping":mapping, "mapping_confidence":confidence, "rows":analyzed, "stats":stats, "ai":ai_meta, "required_columns":["Last_name","First_name"], "canonical_headers":CANONICAL, "supported_extensions":sorted(SUPPORTED_EXTENSIONS)}


def _choose_balanced_arm(level, counts, stream=""):
    arms = list(CLASS_ARMS_BY_LEVEL.get(level, [])); stream = clean(stream).upper()
    if stream == "SCIENCE": arms = [arm for arm in arms if any(tag in arm for tag in ("_GOLD","_SILVER","_DIAMOND"))]
    elif stream == "ART_COMMERCIAL": arms = [arm for arm in arms if arm.endswith("_B/C")]
    if not arms: raise ValueError(f"No configured class arms exist for {level} with the requested stream.")
    return min(arms, key=lambda arm:(counts.get(arm,0), arms.index(arm)))


def _next_batch_admission(available_pool, active, assigned, counter):
    while available_pool:
        candidate = normalize_admission(available_pool.pop(0))
        if candidate and candidate not in active and candidate not in assigned: return candidate, "recycled", counter
    number = counter
    while True:
        number += 1; candidate = format_admission_number(number)
        if candidate not in active and candidate not in assigned and candidate not in set(available_pool): return candidate, "new", number


def commit_student_import(rows, mode="add"):
    mode = clean(mode).lower()
    if mode not in {"add","upsert"}: raise ValueError("Import mode must be add or upsert.")
    source_rows = [dict(row) for row in (rows or []) if row.get("_include", True)]
    if not source_rows: raise ValueError("No student rows were selected for import.")
    if len(source_rows) > MAX_IMPORT_ROWS: raise ValueError(f"A maximum of {MAX_IMPORT_ROWS} students can be imported at once.")

    existing = read_master_students(); active_map = {normalize_admission(row.get("Admission_number")):dict(row) for row in existing if normalize_admission(row.get("Admission_number"))}; active = set(active_map)
    available_pool = list(load_available_numbers()); original_available = list(available_pool); assigned = set(); counter = get_highest_known_admission_number()
    arm_counts = {level:{arm:0 for arm in CLASS_ARMS_BY_LEVEL.get(level,[])} for level in CLASS_ARMS_BY_LEVEL}
    for student in existing:
        level = normalize_student_level(student.get("Class_category") or student.get("Class")); arm = normalize_student_arm(student.get("Class"), level)
        if level in arm_counts and arm in arm_counts[level]: arm_counts[level][arm] += 1

    added, updated, admission_sources, normalized_new = [], [], {}, []
    for idx, row in enumerate(source_rows, start=1):
        last_name, first_name = clean(row.get("Last_name")), clean(row.get("First_name")); other_names, phone = clean(row.get("Other_names")), _normalize_phone(row.get("Phone")); sex = _normalize_sex(row.get("Sex"))
        if not last_name or not first_name: raise ValueError(f"Row {idx}: Last_name and First_name are required.")
        if sex not in {"Male","Female"}: raise ValueError(f"Row {idx} ({last_name} {first_name}): Sex must be Male or Female.")
        level = normalize_student_level(row.get("Class_category") or row.get("Class"))
        if not level: raise ValueError(f"Row {idx} ({last_name} {first_name}): valid Class_category is required.")
        raw_arm = clean(row.get("Class")); auto_key = raw_arm.upper().replace(" ", "_")
        auto_stream = "SCIENCE" if auto_key in {"AUTO_SCIENCE","SCIENCE_AUTO"} else "ART_COMMERCIAL" if auto_key in {"AUTO_ART_COMMERCIAL","AUTO_ARTS_COMMERCIAL","ART_COMMERCIAL_AUTO"} else ""
        arm = "" if auto_key in {"AUTO","AUTOMATIC","BALANCED","AUTO_SCIENCE","SCIENCE_AUTO","AUTO_ART_COMMERCIAL","AUTO_ARTS_COMMERCIAL","ART_COMMERCIAL_AUTO"} else normalize_student_arm(raw_arm, level)
        if not arm: arm = _choose_balanced_arm(level, arm_counts, auto_stream)
        validate_student_class(level, arm); arm_counts[level][arm] = arm_counts[level].get(arm,0) + 1

        imported_admission = normalize_admission(row.get("Admission_number"))
        if mode == "upsert" and imported_admission and imported_admission in active_map:
            student = {"Admission_number":imported_admission, "Last_name":last_name, "First_name":first_name, "Other_names":other_names, "Phone":phone, "Sex":sex, "Class":arm, "Class_category":level}
            active_map[imported_admission] = student; updated.append(student); continue

        admission, source, counter = _next_batch_admission(available_pool, active, assigned, counter); assigned.add(admission); admission_sources[admission] = source
        student = {"Admission_number":admission, "Last_name":last_name, "First_name":first_name, "Other_names":other_names, "Phone":phone, "Sex":sex, "Class":arm, "Class_category":level}
        normalized_new.append(student); added.append(student)

    proposed = list(active_map.values()) + normalized_new
    try:
        transaction = commit_master_transaction(proposed, "bulk_student_import")
        save_available_numbers(available_pool)
    except Exception:
        try: save_available_numbers(original_available)
        except Exception: pass
        raise
    return {"success":True, "count":len(added)+len(updated), "added_count":len(added), "updated_count":len(updated), "added":added, "updated":updated, "admission_sources":admission_sources, "backup":transaction.get("backup"), "total_students":transaction.get("total_students"), "class_counts":transaction.get("class_counts"), "verification":transaction.get("verification")}


def canonical_csv_template():
    stream = io.StringIO(); writer = csv.DictWriter(stream, fieldnames=CANONICAL); writer.writeheader(); writer.writerow({"Last_name":"Example", "First_name":"Student", "Other_names":"", "Phone":"", "Sex":"Female", "Class":"JSS1A", "Class_category":"JSS1"}); return stream.getvalue()
