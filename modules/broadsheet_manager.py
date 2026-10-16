# MODULE: Broadsheet Manager — Phase 8 term + cumulative class-wide performance analysis
# PURPOSE: Build summary broadsheets, annual cumulative broadsheets, analytics and multi-sheet Excel exports from permanent Academic History.

import io
from functools import wraps

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.academic_history import ensure_history_bootstrap, read_academic_records
from modules.academic_records import academic_term_label, normalize_academic_session, normalize_academic_term
from modules.cumulative_results import TERM_LABELS, build_class_cumulative, cumulative_config, grade_for_score, ordinal
from modules.student_lookup import normalize_admission_number
from modules.grading_system import grading_system_for_class, grading_system_ready
from modules.school_structure import CLASS_LABELS

broadsheet_bp = Blueprint("broadsheet_bp", __name__)


def clean(value): return str(value or "").strip()
def clean_upper(value): return clean(value).upper()
def safe_float(value):
    try: return round(float(value), 2) if value not in (None, "") else None
    except (TypeError, ValueError): return None


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if clean(session.get("user_type")).lower() != "admin": return jsonify({"success": False, "message": "Administrator access is required."}), 403
        return view(*args, **kwargs)
    return wrapped


def _record_rows(academic_session, term, class_level, class_arm=""):
    ensure_history_bootstrap(); academic_session = normalize_academic_session(academic_session); term = normalize_academic_term(term); class_level = clean_upper(class_level); class_arm = clean_upper(class_arm); rows = []
    for row in read_academic_records():
        if academic_session and clean(row.get("session")) != academic_session: continue
        if term and normalize_academic_term(row.get("term")) != term: continue
        if class_level and clean_upper(row.get("class_level")) != class_level: continue
        if class_arm and clean_upper(row.get("class_arm")) != class_arm: continue
        rows.append(row)
    return rows


def _subject_key(row): return clean_upper(row.get("subject_key") or row.get("subject"))

def _subject_maximum(row):
    row = row if isinstance(row, dict) else {}; value = safe_float(row.get("total_max")); value = value if value is not None else safe_float(row.get("maximum")); value = value if value is not None else safe_float(row.get("max_score")); return value if value and value > 0 else 100.0

def _grading_payload(class_level):
    system = grading_system_for_class(class_level); return {"system": system, "ready": grading_system_ready(system)}


def _subject_catalog(records):
    catalog = {}
    for record in records:
        report = record.get("report") if isinstance(record.get("report"), dict) else {}
        for subject in report.get("subjects") or []:
            if not isinstance(subject, dict): continue
            key = _subject_key(subject)
            if key and key not in catalog: catalog[key] = {"name": clean(subject.get("subject")) or key.title(), "maximum": _subject_maximum(subject)}
    return [{"key": key, "name": value.get("name"), "maximum": value.get("maximum")} for key, value in sorted(catalog.items(), key=lambda item: item[1].get("name", "").lower())]


def _term_student(record, subjects):
    report = record.get("report") if isinstance(record.get("report"), dict) else {}; mapping, maxima = {}, {}
    for subject in report.get("subjects") or []:
        if not isinstance(subject, dict): continue
        key = _subject_key(subject)
        if not key: continue
        mapping[key] = safe_float(subject.get("total")); maxima[key] = _subject_maximum(subject)
    subject_keys = [subject.get("key") for subject in subjects]; average = safe_float(record.get("average")); provisional = safe_float(report.get("provisional_average")); attendance = safe_float(record.get("attendance_percentage")); academic_complete = bool(record.get("academic_complete")); display_average = average if average is not None else provisional; subject_count = int(record.get("subject_count") or len(subject_keys)); total_score = safe_float(record.get("total_score")); total_score = total_score if total_score is not None else safe_float(report.get("total_score")); maximum_score = safe_float(record.get("maximum_score")); maximum_score = maximum_score if maximum_score is not None else (round(sum(float(maxima.get(key) or 100) for key in subject_keys), 2) if subject_keys else None); score_percentage = round(total_score / maximum_score * 100, 2) if total_score is not None and maximum_score else display_average; class_level = clean_upper(record.get("class_level")); grading = _grading_payload(class_level); grade = clean(record.get("grade")); grade = grade if grade and grade != "-" else grade_for_score(display_average, class_level)[0]
    return {"record_id": clean(record.get("record_id")), "student_key": clean(record.get("student_key")), "admission_number": normalize_admission_number(record.get("admission_number")), "full_name": clean(record.get("full_name")), "sex": clean(record.get("sex")), "class_level": class_level, "class_arm": clean(record.get("class_arm")), "class_label": CLASS_LABELS.get(class_level, class_level), "grading_system": grading.get("system"), "grading_ready": grading.get("ready"), "scores": {key: mapping.get(key) for key in subject_keys}, "score_maxima": {key: maxima.get(key) for key in subject_keys}, "total_score": total_score, "maximum_score": maximum_score, "score_percentage": score_percentage, "average": average, "provisional_average": provisional, "display_average": display_average, "grade": grade or "-", "position": int(record.get("position") or 0), "position_text": clean(record.get("position_text")) or "--", "out_of": int(record.get("out_of") or 0), "attendance_percentage": attendance, "academic_complete": academic_complete, "subjects_complete": int(record.get("subjects_complete") or 0), "subject_count": subject_count, "source_mode": clean(record.get("source_mode") or "auto"), "generated_at": clean(record.get("generated_at"))}


def _grade_distribution(students, field="display_average", class_level=""):
    grading = _grading_payload(class_level)
    if not grading.get("ready"): return {"Pending": sum(1 for student in students if safe_float(student.get(field)) is not None)}
    output = {}
    for student in students:
        score = safe_float(student.get(field))
        if score is None: continue
        grade = grade_for_score(score, class_level)[0]; output[grade] = output.get(grade, 0) + 1
    return output


def _ranked_extremes(students, field="display_average"):
    ranked = [row for row in students if row.get("academic_complete") and safe_float(row.get(field)) is not None]; ranked.sort(key=lambda row: (-float(row.get(field)), clean(row.get("full_name")).lower()))
    top = ranked[:5]; bottom = list(reversed(ranked[-5:])) if ranked else []
    return {"top": [{"admission_number": row.get("admission_number"), "full_name": row.get("full_name"), "average": safe_float(row.get(field)), "position_text": row.get("position_text")} for row in top], "bottom": [{"admission_number": row.get("admission_number"), "full_name": row.get("full_name"), "average": safe_float(row.get(field)), "position_text": row.get("position_text")} for row in bottom]}


def _term_subject_stats(students, subjects, class_level=""):
    output = []; grading_ready = _grading_payload(class_level).get("ready")
    for subject in subjects:
        key = subject.get("key"); maximum = safe_float(subject.get("maximum")) or 100.0; values = [safe_float(student.get("scores", {}).get(key)) for student in students]; values = [value for value in values if value is not None]; percentages = [round(value / maximum * 100, 2) for value in values] if maximum else []; passed = sum(1 for value in percentages if value >= 40) if grading_ready else None; failed = (len(percentages) - passed) if grading_ready and passed is not None else None
        output.append({"subject_key": key, "subject": subject.get("name"), "maximum": maximum, "students_scored": len(values), "average_score": round(sum(values) / len(values), 2) if values else None, "average": round(sum(percentages) / len(percentages), 2) if percentages else None, "highest_score": max(values) if values else None, "lowest_score": min(values) if values else None, "highest": max(percentages) if percentages else None, "lowest": min(percentages) if percentages else None, "passed": passed, "failed": failed, "pass_rate": round(passed / len(percentages) * 100, 2) if grading_ready and percentages and passed is not None else None, "grading_ready": grading_ready})
    return output


def build_term_broadsheet(academic_session, term, class_level, class_arm=""):
    academic_session = normalize_academic_session(academic_session); term = normalize_academic_term(term); level = clean_upper(class_level); grading = _grading_payload(level); rows = _record_rows(academic_session, term, level, class_arm); subjects = _subject_catalog(rows); students = [_term_student(record, subjects) for record in rows]; students.sort(key=lambda row: clean(row.get("full_name")).lower())
    complete_students = [row for row in students if row.get("academic_complete") and safe_float(row.get("average")) is not None]; ranked_values = sorted([float(row.get("average")) for row in complete_students], reverse=True)
    for row in students:
        if row.get("academic_complete") and safe_float(row.get("average")) is not None:
            position = ranked_values.index(float(row.get("average"))) + 1; row["position"], row["position_text"], row["out_of"] = position, ordinal(position), len(complete_students)
    averages = [safe_float(row.get("average")) for row in complete_students]; averages = [value for value in averages if value is not None]; subject_stats = _term_subject_stats(students, subjects, level); best_subject = max([row for row in subject_stats if row.get("average") is not None], key=lambda row: row.get("average"), default=None); focus_subject = min([row for row in subject_stats if row.get("average") is not None], key=lambda row: row.get("average"), default=None); extremes = _ranked_extremes(students, "average")
    return {"session": academic_session, "term": term, "term_label": academic_term_label(term), "class_level": level, "class_label": CLASS_LABELS.get(level, level), "class_arm": clean_upper(class_arm), "grading_system": grading.get("system"), "grading_ready": grading.get("ready"), "count": len(students), "complete_students": len(complete_students), "partial_students": max(0, len(students) - len(complete_students)), "coverage_percentage": round(len(complete_students) / len(students) * 100, 2) if students else 0, "class_average": round(sum(averages) / len(averages), 2) if averages else None, "highest_average": max(averages) if averages else None, "lowest_average": min(averages) if averages else None, "subjects": subjects, "students": students, "subject_stats": subject_stats, "grade_distribution": _grade_distribution(students, "average", level), "top_students": extremes.get("top", []), "bottom_students": extremes.get("bottom", []), "best_subject": best_subject, "focus_subject": focus_subject}


def _annual_subject_stats(students, class_level=""):
    catalog = {}; grading_ready = _grading_payload(class_level).get("ready")
    for student in students:
        for subject in student.get("subjects") or []:
            key = clean_upper(subject.get("subject_key") or subject.get("subject")); name = clean(subject.get("subject")) or key.title(); entry = catalog.setdefault(key, {"subject_key": key, "subject": name, "available": [], "final": []})
            value = safe_float(subject.get("annual_average"))
            if value is None: continue
            entry["available"].append(value)
            if subject.get("complete"): entry["final"].append(value)
    output = []
    for entry in catalog.values():
        values = entry.get("final") or entry.get("available") or []; passed = sum(1 for value in values if value >= 40) if grading_ready else None
        output.append({"subject_key": entry.get("subject_key"), "subject": entry.get("subject"), "students_scored": len(values), "complete_scores": len(entry.get("final") or []), "average": round(sum(values) / len(values), 2) if values else None, "highest": max(values) if values else None, "lowest": min(values) if values else None, "passed": passed, "failed": len(values) - passed if grading_ready and passed is not None else None, "pass_rate": round(passed / len(values) * 100, 2) if grading_ready and values and passed is not None else None, "provisional": not bool(entry.get("final")) and bool(entry.get("available")), "grading_ready": grading_ready})
    output.sort(key=lambda row: clean(row.get("subject")).lower()); return output


def build_cumulative_broadsheet(academic_session, class_level, class_arm=""):
    level = clean_upper(class_level); grading = _grading_payload(level); payload = build_class_cumulative(academic_session, level, class_arm); students = payload.get("students") or []; subject_stats = _annual_subject_stats(students, level); complete = [student for student in students if student.get("complete") and safe_float(student.get("annual_average")) is not None]; ranked = sorted(complete, key=lambda student: (-float(student.get("annual_average")), clean(student.get("profile", {}).get("full_name")).lower())); top = ranked[:5]; bottom = list(reversed(ranked[-5:])) if ranked else []; distribution = _grade_distribution([{"display_average": student.get("annual_average")} for student in complete], "display_average", level)
    term_completion = {}
    for term in ("FIRST", "SECOND", "THIRD"): term_completion[term] = sum(1 for student in students if (student.get("terms", {}).get(term) or {}).get("average") is not None)
    averages = [safe_float(student.get("annual_average")) for student in complete]; averages = [value for value in averages if value is not None]; best_subject = max([row for row in subject_stats if row.get("average") is not None], key=lambda row: row.get("average"), default=None); focus_subject = min([row for row in subject_stats if row.get("average") is not None], key=lambda row: row.get("average"), default=None)
    payload.update({"grading_system": grading.get("system"), "grading_ready": grading.get("ready"), "subject_stats": subject_stats, "grade_distribution": distribution, "top_students": [{"admission_number": student.get("profile", {}).get("admission_number"), "student_key": student.get("profile", {}).get("student_key"), "full_name": student.get("profile", {}).get("full_name"), "average": student.get("annual_average"), "position_text": student.get("annual_position_text")} for student in top], "bottom_students": [{"admission_number": student.get("profile", {}).get("admission_number"), "student_key": student.get("profile", {}).get("student_key"), "full_name": student.get("profile", {}).get("full_name"), "average": student.get("annual_average"), "position_text": student.get("annual_position_text")} for student in bottom], "term_completion": term_completion, "highest_average": max(averages) if averages else None, "lowest_average": min(averages) if averages else None, "best_subject": best_subject, "focus_subject": focus_subject})
    return payload


def build_broadsheet_bundle(academic_session, term, class_level, class_arm=""):
    term_payload = build_term_broadsheet(academic_session, term, class_level, class_arm); cumulative_payload = build_cumulative_broadsheet(academic_session, class_level, class_arm)
    level = clean_upper(class_level); return {"session": normalize_academic_session(academic_session), "term": normalize_academic_term(term), "term_label": academic_term_label(term), "class_level": level, "class_label": CLASS_LABELS.get(level, level), "class_arm": clean_upper(class_arm), "term_broadsheet": term_payload, "cumulative_broadsheet": cumulative_payload}


def broadsheet_config():
    config = cumulative_config(); config["terms"] = [{"key": key, "label": value} for key, value in TERM_LABELS.items()]; return config


def _safe_sheet_title(value):
    text = "".join(char for char in clean(value) if char not in "[]:*?/\\")[:31]
    return text or "Sheet"


def _style_header(sheet, row, fill="173F76"):
    white, thin = "FFFFFF", Side(style="thin", color="D8E4EF")
    for cell in sheet[row]: cell.font = Font(bold=True, color=white); cell.fill = PatternFill("solid", fgColor=fill); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)


def _auto_width(sheet, maximum=32):
    for column_cells in sheet.columns:
        length = max((len(str(cell.value or "")) for cell in column_cells), default=0); sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 9), maximum)


def _write_overview_sheet(workbook, bundle):
    term_data, annual = bundle.get("term_broadsheet") or {}, bundle.get("cumulative_broadsheet") or {}; sheet = workbook.active; sheet.title = "Overview"; class_name = bundle.get("class_arm") or bundle.get("class_level") or "Class"
    sheet.merge_cells("A1:H1"); sheet["A1"] = "EPITOME MODEL ISLAMIC SCHOOLS — BROADSHEET ANALYTICS"; sheet["A1"].font = Font(bold=True, size=16, color="173F76"); sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.append(["Session", bundle.get("session"), "Class", class_name, "Selected Term", bundle.get("term_label"), "Students", term_data.get("count")]); sheet.append([]); sheet.append(["Metric", "Term", "Annual", "", "Subject", "Average", "Pass Rate", "Status"]); _style_header(sheet, 4)
    sheet.append(["Class Average", term_data.get("class_average"), annual.get("class_average"), "", (term_data.get("best_subject") or {}).get("subject"), (term_data.get("best_subject") or {}).get("average"), (term_data.get("best_subject") or {}).get("pass_rate"), "Strongest Term Subject"])
    sheet.append(["Complete Students", term_data.get("complete_students"), annual.get("complete_students"), "", (annual.get("best_subject") or {}).get("subject"), (annual.get("best_subject") or {}).get("average"), (annual.get("best_subject") or {}).get("pass_rate"), "Strongest Annual Subject"])
    sheet.append(["Coverage %", term_data.get("coverage_percentage"), annual.get("coverage_percentage"), "", (term_data.get("focus_subject") or {}).get("subject"), (term_data.get("focus_subject") or {}).get("average"), (term_data.get("focus_subject") or {}).get("pass_rate"), "Term Focus Subject"])
    sheet.append(["Highest Average", term_data.get("highest_average"), annual.get("highest_average"), "", (annual.get("focus_subject") or {}).get("subject"), (annual.get("focus_subject") or {}).get("average"), (annual.get("focus_subject") or {}).get("pass_rate"), "Annual Focus Subject"]); sheet.freeze_panes = "A5"; sheet.sheet_view.showGridLines = False; _auto_width(sheet)


def _write_term_sheet(workbook, data):
    sheet = workbook.create_sheet(_safe_sheet_title(data.get("term_label") or "Term Broadsheet")); subjects = data.get("subjects") or []; headers = ["S/N", "Admission No.", "Student"] + [f"{subject.get('name')} /{int(subject.get('maximum') or 100)}" for subject in subjects] + ["Total Score", "Maximum", "Average %", "Grade", "Position", "Attendance %", "Status"]; sheet.append(headers); _style_header(sheet, 1); keys = [subject.get("key") for subject in subjects]
    for index, student in enumerate(data.get("students") or [], 1): sheet.append([index, student.get("admission_number"), student.get("full_name")] + [f"{student.get('scores', {}).get(key)} / {student.get('score_maxima', {}).get(key)}" if student.get("scores", {}).get(key) is not None and student.get("score_maxima", {}).get(key) else "" for key in keys] + [student.get("total_score"), student.get("maximum_score"), student.get("average"), student.get("grade"), student.get("position_text"), student.get("attendance_percentage"), "Complete" if student.get("academic_complete") else "Partial"])
    sheet.freeze_panes = "D2"; sheet.sheet_view.showGridLines = False; _auto_width(sheet, 20)


def _annual_subject_catalog(data):
    catalog = {}
    for student in data.get("students") or []:
        for subject in student.get("subjects") or []:
            key = clean_upper(subject.get("subject_key") or subject.get("subject")); catalog.setdefault(key, clean(subject.get("subject")) or key.title())
    return [{"key": key, "name": name} for key, name in sorted(catalog.items(), key=lambda item: item[1].lower())]


def _write_annual_sheet(workbook, data):
    sheet = workbook.create_sheet("Cumulative Broadsheet"); subjects = _annual_subject_catalog(data); headers = ["S/N", "Admission No.", "Student", "1st Term Score", "1st %", "2nd Term Score", "2nd %", "3rd Term Score", "3rd %"] + [f"{subject.get('name')} Annual Score / %" for subject in subjects] + ["Annual Score", "Annual Max", "Annual %", "Grade", "Position", "Status"]; sheet.append(headers); _style_header(sheet, 1); keys = [subject.get("key") for subject in subjects]
    for index, student in enumerate(data.get("students") or [], 1):
        profile, terms = student.get("profile") or {}, student.get("terms") or {}; subject_map = {clean_upper(row.get("subject_key") or row.get("subject")): row for row in student.get("subjects") or []}; status = "Complete" if student.get("complete") else f"Partial {student.get('terms_available', 0)}/3"; ratio = lambda term: f"{terms.get(term, {}).get('score')} / {terms.get(term, {}).get('maximum')}" if terms.get(term, {}).get("score") is not None and terms.get(term, {}).get("maximum") else ""; subject_value = lambda key: f"{subject_map.get(key, {}).get('annual_score')} / {subject_map.get(key, {}).get('annual_max')} ({subject_map.get(key, {}).get('annual_average')}%)" if subject_map.get(key, {}).get("annual_score") is not None and subject_map.get(key, {}).get("annual_max") else ""
        sheet.append([index, profile.get("admission_number"), profile.get("full_name"), ratio("FIRST"), terms.get("FIRST", {}).get("average"), ratio("SECOND"), terms.get("SECOND", {}).get("average"), ratio("THIRD"), terms.get("THIRD", {}).get("average")] + [subject_value(key) for key in keys] + [student.get("annual_score"), student.get("annual_max"), student.get("annual_score_percentage") if student.get("annual_score_percentage") is not None else student.get("annual_average"), student.get("annual_grade"), student.get("annual_position_text"), status])
    sheet.freeze_panes = "D2"; sheet.sheet_view.showGridLines = False; _auto_width(sheet, 24)


def _write_subject_sheet(workbook, bundle):
    sheet = workbook.create_sheet("Subject Analytics"); sheet.append(["Scope", "Subject", "Students Scored", "Average", "Highest", "Lowest", "Passed", "Failed", "Pass Rate %", "Status"]); _style_header(sheet, 1)
    for scope, rows in ((bundle.get("term_label") or "Term", (bundle.get("term_broadsheet") or {}).get("subject_stats") or []), ("Annual", (bundle.get("cumulative_broadsheet") or {}).get("subject_stats") or [])):
        for row in rows: sheet.append([scope, row.get("subject"), row.get("students_scored"), row.get("average"), row.get("highest"), row.get("lowest"), row.get("passed"), row.get("failed"), row.get("pass_rate"), "Provisional" if row.get("provisional") else "Final"])
    sheet.freeze_panes = "A2"; sheet.sheet_view.showGridLines = False; _auto_width(sheet, 24)


def build_broadsheet_workbook(bundle):
    workbook = Workbook(); _write_overview_sheet(workbook, bundle); _write_term_sheet(workbook, bundle.get("term_broadsheet") or {}); _write_annual_sheet(workbook, bundle.get("cumulative_broadsheet") or {}); _write_subject_sheet(workbook, bundle); return workbook


@broadsheet_bp.route("/api/broadsheets/config")
@admin_required
def api_broadsheet_config(): return jsonify({"success": True, **broadsheet_config()})


@broadsheet_bp.route("/api/broadsheets/data")
@admin_required
def api_broadsheet_data():
    academic_session, term, class_level, class_arm = request.args.get("session"), request.args.get("term"), request.args.get("class_level"), request.args.get("class_arm")
    if not normalize_academic_session(academic_session) or not normalize_academic_term(term) or not clean(class_level): return jsonify({"success": False, "message": "Select an academic session, term and class."}), 400
    return jsonify({"success": True, **build_broadsheet_bundle(academic_session, term, class_level, class_arm)})


@broadsheet_bp.route("/api/broadsheets/export")
@admin_required
def api_broadsheet_export():
    academic_session, term, class_level, class_arm = request.args.get("session"), request.args.get("term"), request.args.get("class_level"), request.args.get("class_arm")
    if not normalize_academic_session(academic_session) or not normalize_academic_term(term) or not clean(class_level): return jsonify({"success": False, "message": "Select an academic session, term and class before export."}), 400
    bundle = build_broadsheet_bundle(academic_session, term, class_level, class_arm); workbook = build_broadsheet_workbook(bundle); output = io.BytesIO(); workbook.save(output); output.seek(0); class_name = clean(class_arm or class_level or "Class").replace(" ", "_"); session_name = clean(academic_session).replace("/", "-")
    return send_file(output, as_attachment=True, download_name=f"{class_name}_{session_name}_Broadsheets.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
