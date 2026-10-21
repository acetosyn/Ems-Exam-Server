# MODULE: Cumulative Results — Annual First/Second/Third term cumulative result engine
# PURPOSE: Build printable cumulative student sheets and class-wide annual summaries from permanent Academic History records.

import io
from functools import wraps

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.academic_history import ensure_history_bootstrap, read_academic_records
from modules.class_config import CLASS_ARMS_BY_LEVEL, STUDENT_DATABASE_CLASSES, TAHFEEZ_CLASSES
from modules.academic_records import academic_term_label, normalize_academic_session, normalize_academic_term
from modules.student_lookup import normalize_admission_number
from modules.grading_system import grade_score_for_class, grading_system_for_class, grading_system_ready
from modules.school_structure import CLASS_LABELS

cumulative_results_bp = Blueprint("cumulative_results_bp", __name__)

TERM_ORDER = {"FIRST": 1, "SECOND": 2, "THIRD": 3}
TERM_LABELS = {"FIRST": "First Term", "SECOND": "Second Term", "THIRD": "Third Term"}

def clean(value): return str(value or "").strip()
def clean_upper(value): return clean(value).upper()
def safe_float(value):
    try: return round(float(value), 2) if value not in (None, "") else None
    except (TypeError, ValueError): return None

def round2(value): return round(float(value), 2) if value is not None else None

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if clean(session.get("user_type")).lower() != "admin": return jsonify({"success": False, "message": "Administrator access is required."}), 403
        return view(*args, **kwargs)
    return wrapped


def grade_for_score(score, class_level=""):
    score = safe_float(score)
    if score is None: return "-", "Pending"
    return grade_score_for_class(score, class_level)

def _grading_payload(class_level):
    system = grading_system_for_class(class_level); return {"system": system, "ready": grading_system_ready(system)}

def _subject_maximum(row):
    row = row if isinstance(row, dict) else {}; value = safe_float(row.get("total_max")); value = value if value is not None else safe_float(row.get("maximum")); value = value if value is not None else safe_float(row.get("max_score")); return value if value and value > 0 else 100.0

def ordinal(number):
    try: number = int(number)
    except (TypeError, ValueError): return "--"
    if number <= 0: return "--"
    if 10 <= number % 100 <= 20: suffix = "th"
    else: suffix = {1: "st", 2: "nd", 3: "rd"}.get(number % 10, "th")
    return f"{number}{suffix}"


def _class_display(record): return clean(record.get("class_arm")) or clean(record.get("class_level"))

def _subject_key(row): return clean_upper(row.get("subject_key") or row.get("subject"))

def _term_record_map(records): return {normalize_academic_term(row.get("term")): row for row in records if normalize_academic_term(row.get("term")) in TERM_ORDER}


def _record_score_bundle(record):
    record = record if isinstance(record, dict) else {}; report = record.get("report") if isinstance(record.get("report"), dict) else {}; subjects = report.get("subjects") if isinstance(report.get("subjects"), list) else []; total = safe_float(record.get("total_score")); total = total if total is not None else safe_float(report.get("total_score")); subject_count = int(record.get("subject_count") or report.get("subject_count") or len(subjects)); maximum = safe_float(record.get("maximum_score")); maximum = maximum if maximum is not None else (round(sum(_subject_maximum(row) for row in subjects if isinstance(row, dict)), 2) if subjects else None); maximum = maximum if maximum is not None else (round(subject_count * 100, 2) if subject_count else None); percentage = round(total / maximum * 100, 2) if total is not None and maximum else safe_float(record.get("score_percentage")); percentage = percentage if percentage is not None else safe_float(record.get("average")); return {"score": total, "maximum": maximum, "percentage": percentage, "subject_count": subject_count}

def _records_for_scope(academic_session, class_level="", class_arm=""):
    ensure_history_bootstrap(); academic_session = normalize_academic_session(academic_session); class_level = clean_upper(class_level); class_arm = clean_upper(class_arm)
    rows = []
    for row in read_academic_records():
        if clean(row.get("record_variant") or "standard").lower() != "standard": continue
        if academic_session and clean(row.get("session")) != academic_session: continue
        if class_level and clean_upper(row.get("class_level")) != class_level: continue
        if class_arm and clean_upper(row.get("class_arm")) != class_arm: continue
        rows.append(row)
    return rows


def _attendance_from_record(record):
    report = record.get("report") if isinstance(record.get("report"), dict) else {}; attendance = report.get("attendance") if isinstance(report.get("attendance"), dict) else {}
    return {"days_open": int(attendance.get("days_open") or attendance.get("days_school_open") or 0), "present": int(attendance.get("present") or attendance.get("days_present") or 0), "absent": int(attendance.get("absent") or attendance.get("days_absent") or 0), "late": int(attendance.get("late") or attendance.get("days_late") or 0), "sick": int(attendance.get("sick") or attendance.get("days_sick") or 0), "excused": int(attendance.get("excused") or attendance.get("days_excused") or 0), "present_credit": safe_float(attendance.get("present_credit") or attendance.get("attendance_score"))}


def _combine_attendance(term_records):
    summary = {"days_open": 0, "present": 0, "absent": 0, "late": 0, "sick": 0, "excused": 0, "present_credit": 0.0}
    any_attendance = False
    for record in term_records.values():
        attendance = _attendance_from_record(record)
        if attendance["days_open"] or attendance["present"] or attendance["absent"] or attendance["late"] or attendance["sick"] or attendance["excused"]: any_attendance = True
        for key in ("days_open", "present", "absent", "late", "sick", "excused"): summary[key] += int(attendance.get(key) or 0)
        summary["present_credit"] += float(attendance.get("present_credit") or 0)
    summary["attendance_percentage"] = round(summary["present_credit"] / summary["days_open"] * 100, 2) if summary["days_open"] else None; summary["available"] = any_attendance
    return summary


def _subject_rows(term_records, class_level=""):
    subjects = {}; term_subjects = {}
    for term, record in term_records.items():
        report = record.get("report") if isinstance(record.get("report"), dict) else {}; mapping = {}
        for row in report.get("subjects") or []:
            if not isinstance(row, dict): continue
            key = _subject_key(row)
            if not key: continue
            mapping[key] = row; subjects.setdefault(key, clean(row.get("subject")) or key.title())
        term_subjects[term] = mapping
    output = []
    for key, name in subjects.items():
        totals, maxima = {}, {}
        for term in TERM_ORDER:
            row = term_subjects.get(term, {}).get(key) or {}; total = safe_float(row.get("total")); totals[term] = total; maxima[term] = _subject_maximum(row) if total is not None else None
        available_terms = [term for term in TERM_ORDER if totals.get(term) is not None and maxima.get(term)]
        annual_score = round(sum(float(totals[term]) for term in available_terms), 2) if available_terms else None; annual_max = round(sum(float(maxima[term]) for term in available_terms), 2) if available_terms else None; annual_average = round(annual_score / annual_max * 100, 2) if annual_score is not None and annual_max else None; grade, comment = grade_for_score(annual_average, class_level)
        output.append({"subject_key": key, "subject": name, "first_term": totals.get("FIRST"), "second_term": totals.get("SECOND"), "third_term": totals.get("THIRD"), "first_term_max": maxima.get("FIRST"), "second_term_max": maxima.get("SECOND"), "third_term_max": maxima.get("THIRD"), "annual_score": annual_score, "annual_max": annual_max, "annual_average": annual_average, "grade": grade, "comment": comment, "terms_available": len(available_terms), "complete": len(available_terms) == 3})
    output.sort(key=lambda row: row.get("subject", "").lower()); return output


def _student_payload(student_key, records):
    records = sorted(records, key=lambda row: TERM_ORDER.get(normalize_academic_term(row.get("term")), 0)); term_records = _term_record_map(records); latest = records[-1] if records else {}; class_level = clean_upper(latest.get("class_level")); grading = _grading_payload(class_level); term_averages = {term: safe_float(term_records.get(term, {}).get("average")) for term in TERM_ORDER}; available_averages = [value for value in term_averages.values() if value is not None]; annual_average = round(sum(available_averages) / len(available_averages), 2) if available_averages else None; grade, comment = grade_for_score(annual_average, class_level); subject_rows = _subject_rows(term_records, class_level); complete_terms = sum(1 for term in TERM_ORDER if term in term_records and safe_float(term_records[term].get("average")) is not None)
    profile = {"student_key": student_key, "admission_number": normalize_admission_number(latest.get("admission_number")), "full_name": clean(latest.get("full_name")), "sex": clean(latest.get("sex")), "session": clean(latest.get("session")), "class_level": class_level, "class_arm": clean(latest.get("class_arm")), "class_display": _class_display(latest), "class_label": CLASS_LABELS.get(class_level, _class_display(latest))}
    terms = {}; annual_score = 0.0; annual_max = 0.0; score_terms = 0
    for term in TERM_ORDER:
        row = term_records.get(term); bundle = _record_score_bundle(row) if row else {"score": None, "maximum": None, "percentage": None, "subject_count": 0}
        if bundle.get("score") is not None and bundle.get("maximum"): annual_score += float(bundle["score"]); annual_max += float(bundle["maximum"]); score_terms += 1
        term_grade = clean(row.get("grade")) if row else "-"
        if row and not term_grade: term_grade = grade_for_score(safe_float(row.get("average")), class_level)[0]
        terms[term] = {"available": bool(row), "average": safe_float(row.get("average")) if row else None, "score": bundle.get("score"), "maximum": bundle.get("maximum"), "score_percentage": bundle.get("percentage"), "subject_count": bundle.get("subject_count"), "grade": term_grade or "-", "position": int(row.get("position") or 0) if row else 0, "position_text": clean(row.get("position_text")) if row else "--", "out_of": int(row.get("out_of") or 0) if row else 0, "attendance_percentage": safe_float(row.get("attendance_percentage")) if row else None, "academic_complete": bool(row.get("academic_complete")) if row else False, "record_id": clean(row.get("record_id")) if row else ""}
    annual_score = round(annual_score, 2) if score_terms else None; annual_max = round(annual_max, 2) if score_terms else None; annual_score_percentage = round(annual_score / annual_max * 100, 2) if annual_score is not None and annual_max else None
    return {"profile": profile, "grading_system": grading.get("system"), "grading_ready": grading.get("ready"), "terms": terms, "subjects": subject_rows, "attendance": _combine_attendance(term_records), "annual_score": annual_score, "annual_max": annual_max, "annual_score_percentage": annual_score_percentage, "annual_average": annual_average, "annual_grade": grade, "annual_comment": comment, "terms_available": complete_terms, "score_terms_available": score_terms, "complete": complete_terms == 3, "subjects_complete": sum(1 for row in subject_rows if row.get("complete")), "subject_count": len(subject_rows), "annual_position": 0, "annual_position_text": "--", "annual_out_of": 0}


def build_class_cumulative_from_records(rows, academic_session, class_level="", class_arm=""):
    """Compile cumulative results from an explicit record set.

    Production uses Academic History; Report Test Lab passes an isolated synthetic record set here so trial runs can exercise the same cumulative engine without contaminating official history.
    """
    academic_session = normalize_academic_session(academic_session); level = clean_upper(class_level); arm = clean_upper(class_arm); scoped = []
    for row in rows or []:
        if academic_session and clean(row.get("session")) != academic_session: continue
        if level and clean_upper(row.get("class_level")) != level: continue
        if arm and clean_upper(row.get("class_arm")) != arm: continue
        scoped.append(row)
    grouped = {}
    for row in scoped:
        key = clean(row.get("student_key")) or f"{normalize_admission_number(row.get('admission_number'))}|{clean(row.get('full_name')).lower()}"; grouped.setdefault(key, []).append(row)
    students = [_student_payload(key, records) for key, records in grouped.items()]; students.sort(key=lambda item: clean(item.get("profile", {}).get("full_name")).lower())
    complete_students = [student for student in students if student.get("complete") and student.get("annual_average") is not None]; ranked_values = sorted([student.get("annual_average") for student in complete_students], reverse=True)
    for student in students:
        if not student.get("complete") or student.get("annual_average") is None: continue
        position = ranked_values.index(student.get("annual_average")) + 1; student["annual_position"] = position; student["annual_position_text"] = ordinal(position); student["annual_out_of"] = len(complete_students)
    averages = [student.get("annual_average") for student in complete_students if student.get("annual_average") is not None]; class_average = round(sum(averages) / len(averages), 2) if averages else None
    grading = _grading_payload(level); return {"session": academic_session, "class_level": level, "class_label": CLASS_LABELS.get(level, level), "class_arm": arm, "grading_system": grading.get("system"), "grading_ready": grading.get("ready"), "count": len(students), "complete_students": len(complete_students), "partial_students": max(0, len(students) - len(complete_students)), "class_average": class_average, "coverage_percentage": round(len(complete_students) / len(students) * 100, 2) if students else 0, "students": students}


def build_class_cumulative(academic_session, class_level="", class_arm=""):
    return build_class_cumulative_from_records(_records_for_scope(academic_session, class_level, class_arm), academic_session, class_level, class_arm)


def cumulative_config():
    ensure_history_bootstrap(); records = read_academic_records(); sessions = sorted({clean(row.get("session")) for row in records if clean(row.get("session"))}, key=lambda value: int(value.split("/")[0]) if value.split("/")[0].isdigit() else 0, reverse=True); archived_classes = sorted({clean_upper(row.get("class_level")) for row in records if clean(row.get("class_level"))}); configured = [clean_upper(level) for level in list(STUDENT_DATABASE_CLASSES) + list(TAHFEEZ_CLASSES)]; classes = list(dict.fromkeys(configured + archived_classes)); arms = {}
    for level in classes:
        defaults = [clean_upper(value) for value in CLASS_ARMS_BY_LEVEL.get(level, [])]
        if not defaults and (level in [clean_upper(v) for v in STUDENT_DATABASE_CLASSES] or level in [clean_upper(v) for v in TAHFEEZ_CLASSES]): defaults = [level]
        archived_arms = [clean_upper(row.get("class_arm")) for row in records if clean_upper(row.get("class_level")) == level and clean(row.get("class_arm"))]
        arms[level] = list(dict.fromkeys(defaults + sorted(set(archived_arms))))
    return {"sessions": sessions, "classes": classes, "class_labels": {level: CLASS_LABELS.get(level, level) for level in classes}, "arms": arms, "archived_classes": archived_classes, "term_labels": TERM_LABELS, "grading": {level: _grading_payload(level) for level in classes}}


def _find_student(class_payload, admission_number, student_key=""):
    admission = normalize_admission_number(admission_number); student_key = clean(student_key); candidates = [student for student in class_payload.get("students", []) if normalize_admission_number(student.get("profile", {}).get("admission_number")) == admission]
    if student_key: candidates = [student for student in candidates if clean(student.get("profile", {}).get("student_key")) == student_key]
    if not candidates: return None, False
    return (candidates[0], len(candidates) > 1 and not student_key)


def _style_workbook_sheet(sheet, title_columns):
    navy, soft, green, white = "173F76", "EAF2FB", "E9F8F1", "FFFFFF"; thin = Side(style="thin", color="D8E4EF")
    sheet.freeze_panes = "A8"; sheet.sheet_view.showGridLines = False
    for cell in sheet[1]: cell.font = Font(bold=True, size=16, color=navy); cell.alignment = Alignment(horizontal="center")
    for cell in sheet[7]: cell.font = Font(bold=True, color=white); cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)
    for row in sheet.iter_rows(min_row=8):
        for cell in row: cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in enumerate(title_columns, 1): sheet.column_dimensions[get_column_letter(column)].width = width


def build_class_workbook(class_payload):
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Cumulative Results"; class_name = class_payload.get("class_arm") or class_payload.get("class_level") or "Class"
    sheet.merge_cells("A1:M1"); sheet["A1"] = "EPITOME MODEL ISLAMIC SCHOOLS — CUMULATIVE RESULT SUMMARY"; sheet.append(["Session", class_payload.get("session"), "Class", class_name, "Students", class_payload.get("count"), "Complete", class_payload.get("complete_students"), "Class Average", class_payload.get("class_average"), "Coverage %", class_payload.get("coverage_percentage"), ""]); sheet.append([]); sheet.append(["Real scores are shown with their maximum possible scores. Percentages remain available for comparison."]); sheet.merge_cells("A4:M4"); sheet.append([]); sheet.append([]); sheet.append(["S/N", "Admission No.", "Student", "1st Term Score", "1st %", "2nd Term Score", "2nd %", "3rd Term Score", "3rd %", "Annual Score", "Annual %", "Grade", "Position / Status"])
    for index, student in enumerate(class_payload.get("students") or [], 1):
        profile, terms = student.get("profile") or {}, student.get("terms") or {}; status = student.get("annual_position_text") if student.get("complete") else f"Partial ({student.get('terms_available', 0)}/3)"
        ratio = lambda term: f"{terms.get(term, {}).get('score')} / {terms.get(term, {}).get('maximum')}" if terms.get(term, {}).get("score") is not None and terms.get(term, {}).get("maximum") else ""
        annual_ratio = f"{student.get('annual_score')} / {student.get('annual_max')}" if student.get("annual_score") is not None and student.get("annual_max") else ""
        sheet.append([index, profile.get("admission_number"), profile.get("full_name"), ratio("FIRST"), terms.get("FIRST", {}).get("average"), ratio("SECOND"), terms.get("SECOND", {}).get("average"), ratio("THIRD"), terms.get("THIRD", {}).get("average"), annual_ratio, student.get("annual_average"), student.get("annual_grade"), status])
    _style_workbook_sheet(sheet, [7, 16, 28, 18, 10, 18, 10, 18, 10, 20, 11, 9, 18]); return workbook


def build_student_workbook(student):
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Cumulative Sheet"; profile, terms = student.get("profile") or {}, student.get("terms") or {}; class_name = profile.get("class_arm") or profile.get("class_level")
    sheet.merge_cells("A1:I1"); sheet["A1"] = "EPITOME MODEL ISLAMIC SCHOOLS — STUDENT CUMULATIVE RESULT"; sheet.append(["Student", profile.get("full_name"), "Admission No.", profile.get("admission_number"), "Class", class_name, "Session", profile.get("session"), ""]); sheet.append(["1st Term Score", f"{terms.get('FIRST', {}).get('score')} / {terms.get('FIRST', {}).get('maximum')}" if terms.get('FIRST', {}).get('score') is not None else "", "1st %", terms.get("FIRST", {}).get("average"), "2nd Term Score", f"{terms.get('SECOND', {}).get('score')} / {terms.get('SECOND', {}).get('maximum')}" if terms.get('SECOND', {}).get('score') is not None else "", "2nd %", terms.get("SECOND", {}).get("average"), ""]); sheet.append(["3rd Term Score", f"{terms.get('THIRD', {}).get('score')} / {terms.get('THIRD', {}).get('maximum')}" if terms.get('THIRD', {}).get('score') is not None else "", "3rd %", terms.get("THIRD", {}).get("average"), "Annual Score", f"{student.get('annual_score')} / {student.get('annual_max')}" if student.get('annual_score') is not None else "", "Annual %", student.get("annual_average"), student.get("annual_grade")]); sheet.append([]); sheet.append([]); sheet.append(["S/N", "Subject", "1st Term", "2nd Term", "3rd Term", "Annual Score", "Annual Max", "Annual %", "Grade"])
    for index, row in enumerate(student.get("subjects") or [], 1): sheet.append([index, row.get("subject"), f"{row.get('first_term')} / {row.get('first_term_max')}" if row.get("first_term") is not None and row.get("first_term_max") else "", f"{row.get('second_term')} / {row.get('second_term_max')}" if row.get("second_term") is not None and row.get("second_term_max") else "", f"{row.get('third_term')} / {row.get('third_term_max')}" if row.get("third_term") is not None and row.get("third_term_max") else "", row.get("annual_score"), row.get("annual_max"), row.get("annual_average"), row.get("grade")])
    _style_workbook_sheet(sheet, [7, 28, 14, 14, 14, 14, 14, 12, 10]); return workbook


@cumulative_results_bp.route("/api/cumulative-results/config")
@admin_required
def api_cumulative_config(): return jsonify({"success": True, **cumulative_config()})


@cumulative_results_bp.route("/api/cumulative-results/class")
@admin_required
def api_cumulative_class():
    academic_session = request.args.get("session"); class_level = request.args.get("class_level"); class_arm = request.args.get("class_arm")
    if not normalize_academic_session(academic_session) or not clean(class_level): return jsonify({"success": False, "message": "Select an academic session and class."}), 400
    return jsonify({"success": True, **build_class_cumulative(academic_session, class_level, class_arm)})


@cumulative_results_bp.route("/api/cumulative-results/student/<admission_number>")
@admin_required
def api_cumulative_student(admission_number):
    academic_session, class_level, class_arm, student_key = request.args.get("session"), request.args.get("class_level"), request.args.get("class_arm"), request.args.get("student_key")
    payload = build_class_cumulative(academic_session, class_level, class_arm); student, ambiguous = _find_student(payload, admission_number, student_key)
    if ambiguous: return jsonify({"success": False, "message": "This admission number belongs to more than one historical identity. Select the correct student record first."}), 409
    if not student: return jsonify({"success": False, "message": "No cumulative academic record was found for this student in the selected class/session."}), 404
    return jsonify({"success": True, "student": student, "class_average": payload.get("class_average"), "class_count": payload.get("count"), "complete_count": payload.get("complete_students")})


@cumulative_results_bp.route("/api/cumulative-results/export/class")
@admin_required
def api_cumulative_export_class():
    payload = build_class_cumulative(request.args.get("session"), request.args.get("class_level"), request.args.get("class_arm")); workbook = build_class_workbook(payload); output = io.BytesIO(); workbook.save(output); output.seek(0); class_name = clean(payload.get("class_arm") or payload.get("class_level") or "Class").replace(" ", "_"); session_name = clean(payload.get("session") or "Session").replace("/", "-")
    return send_file(output, as_attachment=True, download_name=f"{class_name}_{session_name}_Cumulative_Results.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@cumulative_results_bp.route("/api/cumulative-results/export/student/<admission_number>")
@admin_required
def api_cumulative_export_student(admission_number):
    payload = build_class_cumulative(request.args.get("session"), request.args.get("class_level"), request.args.get("class_arm")); student, ambiguous = _find_student(payload, admission_number, request.args.get("student_key"))
    if ambiguous: return jsonify({"success": False, "message": "Select the correct historical student identity before export."}), 409
    if not student: return jsonify({"success": False, "message": "Student cumulative record was not found."}), 404
    workbook = build_student_workbook(student); output = io.BytesIO(); workbook.save(output); output.seek(0); profile = student.get("profile") or {}; safe_name = "_".join(clean(profile.get("full_name") or admission_number).split())
    return send_file(output, as_attachment=True, download_name=f"{safe_name}_Cumulative_Result.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
