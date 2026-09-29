# MODULE: Academic History — Permanent student term-record archive and history APIs
# PURPOSE: Preserve generated academic records by student identity across promotion, graduation and admission-number reuse.

import hashlib
import io
import json
import threading
from datetime import datetime, timezone
from functools import wraps
from pathlib import Path

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.academic_records import academic_term_label, normalize_academic_session, normalize_academic_term
from modules.grading_system import grade_score
from modules.student_lookup import find_student_by_admission, normalize_admission_number
from modules.result_sync import queue_emis_event_safely

academic_history_bp = Blueprint("academic_history_bp", __name__)

BASE_DIR = Path(__file__).resolve().parent.parent
HISTORY_DIR = BASE_DIR / "static" / "data" / "academic_history"
RECORD_FILE = HISTORY_DIR / "academic_records.jsonl"
REVISION_FILE = HISTORY_DIR / "record_revisions.jsonl"
SNAPSHOT_FILE = BASE_DIR / "static" / "data" / "report_sheets" / "generated_reports.jsonl"
HISTORY_LOCK = threading.RLock()
TERM_ORDER = {"FIRST": 1, "SECOND": 2, "THIRD": 3}


def clean(value): return str(value or "").strip()
def clean_upper(value): return clean(value).upper()
def now_iso(): return datetime.now(timezone.utc).isoformat(timespec="seconds")
def safe_float(value):
    try: return round(float(value), 2) if value not in (None, "") else None
    except (TypeError, ValueError): return None


def _score_bundle(record):
    record = record if isinstance(record, dict) else {}; report = record.get("report") if isinstance(record.get("report"), dict) else {}; subjects = report.get("subjects") if isinstance(report.get("subjects"), list) else []; subject_count = int(record.get("subject_count") or report.get("subject_count") or len(subjects)); total = safe_float(record.get("total_score")); total = total if total is not None else safe_float(report.get("total_score")); maximum = safe_float(record.get("maximum_score")); maximum = maximum if maximum is not None else (round(subject_count * 100, 2) if subject_count else None); percentage = round(total / maximum * 100, 2) if total is not None and maximum else safe_float(record.get("average")); return {"score": total, "maximum": maximum, "percentage": percentage, "subject_count": subject_count}

def _score_ratio(score, maximum): return f"{score} / {maximum}" if score is not None and maximum else ""


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if clean(session.get("user_type")).lower() != "admin": return jsonify({"success": False, "message": "Administrator access is required."}), 403
        return view(*args, **kwargs)
    return wrapped


def ensure_history_dir(): HISTORY_DIR.mkdir(parents=True, exist_ok=True); return HISTORY_DIR

def _normalize_name(value): return " ".join(clean(value).lower().split())

def student_identity_key(admission_number, full_name=""):
    admission = normalize_admission_number(admission_number); name = _normalize_name(full_name)
    return hashlib.sha256(f"{admission}|{name}".encode("utf-8")).hexdigest()[:20]


def academic_record_key(student_key, academic_session, term, class_level, class_arm):
    raw = "|".join([clean(student_key), normalize_academic_session(academic_session), normalize_academic_term(term), clean_upper(class_level), clean_upper(class_arm)])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]


def _read_jsonl(path):
    if not path.exists(): return []
    rows = []
    with path.open("r", encoding="utf-8") as file:
        for line in file:
            try:
                item = json.loads(line)
                if isinstance(item, dict): rows.append(item)
            except Exception: continue
    return rows


def _write_jsonl(path, rows):
    ensure_history_dir(); temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8") as file:
        for row in rows: file.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
    temp.replace(path)


def read_academic_records():
    with HISTORY_LOCK: return _read_jsonl(RECORD_FILE)


def read_revisions():
    with HISTORY_LOCK: return _read_jsonl(REVISION_FILE)


def _snapshot_to_record(snapshot):
    snapshot = snapshot if isinstance(snapshot, dict) else {}; report = snapshot.get("report") if isinstance(snapshot.get("report"), dict) else {}
    admission = normalize_admission_number(report.get("admission_number") or snapshot.get("admission_number")); full_name = clean(report.get("full_name") or snapshot.get("full_name"))
    academic_session = normalize_academic_session(report.get("session") or snapshot.get("session")); term = normalize_academic_term(report.get("term") or snapshot.get("term")); class_level = clean_upper(report.get("class_level") or snapshot.get("class_level")); class_arm = clean_upper(report.get("class_arm") or snapshot.get("class_arm"))
    if not admission or not academic_session or not term or not class_level: return None
    student_key = student_identity_key(admission, full_name); record_id = academic_record_key(student_key, academic_session, term, class_level, class_arm)
    attendance = report.get("attendance") if isinstance(report.get("attendance"), dict) else {}; subjects = report.get("subjects") if isinstance(report.get("subjects"), list) else []
    average = safe_float(report.get("average")); total_score = safe_float(report.get("total_score")); maximum_score = round(len(subjects) * 100, 2) if subjects else None; score_percentage = round(total_score / maximum_score * 100, 2) if total_score is not None and maximum_score else average; attendance_percentage = safe_float(attendance.get("attendance_percentage"))
    return {
        "record_id": record_id, "student_key": student_key, "admission_number": admission, "full_name": full_name, "sex": clean(report.get("sex")),
        "session": academic_session, "term": term, "term_label": academic_term_label(term), "class_level": class_level, "class_arm": class_arm, "result_year": clean(report.get("result_year")),
        "snapshot_id": clean(snapshot.get("snapshot_id")), "result_name": clean(snapshot.get("result_name") or report.get("result_name")), "generated_at": clean(snapshot.get("generated_at")) or now_iso(), "generated_by": clean(snapshot.get("generated_by")) or "Staff",
        "source_mode": clean(report.get("source_mode") or snapshot.get("source_mode") or "auto"), "academic_complete": bool(report.get("academic_complete")), "subject_count": len(subjects), "subjects_complete": int(report.get("subjects_complete") or sum(1 for row in subjects if isinstance(row, dict) and row.get("complete"))),
        "total_score": total_score, "maximum_score": maximum_score, "score_percentage": score_percentage, "average": average, "grade": clean(report.get("grade")) or "-", "position": int(report.get("position") or 0), "position_text": clean(report.get("position_text")) or "--", "out_of": int(report.get("out_of") or 0), "attendance_percentage": attendance_percentage,
        "teacher_remark": clean(report.get("teacher_remark")), "principal_remark": clean(report.get("principal_remark")), "form_teacher": clean(report.get("form_teacher")), "revision_count": 1, "archived_at": now_iso(), "report": report,
    }


def _append_revision(previous, replaced_by):
    if not previous: return
    revision = {"revision_id": hashlib.sha256(f"{previous.get('record_id')}|{previous.get('snapshot_id')}|{previous.get('archived_at')}".encode("utf-8")).hexdigest()[:24], "record_id": previous.get("record_id"), "student_key": previous.get("student_key"), "admission_number": previous.get("admission_number"), "session": previous.get("session"), "term": previous.get("term"), "class_level": previous.get("class_level"), "class_arm": previous.get("class_arm"), "snapshot_id": previous.get("snapshot_id"), "archived_at": previous.get("archived_at"), "replaced_at": now_iso(), "replaced_by_snapshot_id": clean((replaced_by or {}).get("snapshot_id")), "record": previous}
    with REVISION_FILE.open("a", encoding="utf-8") as file: file.write(json.dumps(revision, ensure_ascii=False, separators=(",", ":")) + "\n")


def upsert_academic_record(record, queue_sync=True, preserve_revision=True):
    if not isinstance(record, dict) or not clean(record.get("record_id")): raise ValueError("A valid academic history record is required.")
    ensure_history_dir(); record = dict(record); record["archived_at"] = now_iso()
    with HISTORY_LOCK:
        rows = _read_jsonl(RECORD_FILE); index = {clean(row.get("record_id")): i for i, row in enumerate(rows)}; rid = clean(record.get("record_id")); previous = rows[index[rid]] if rid in index else None
        if previous:
            same_snapshot = clean(previous.get("snapshot_id")) == clean(record.get("snapshot_id")) and clean(record.get("snapshot_id"))
            if same_snapshot: return {"record": previous, "created": False, "updated": False, "stale": False}
            incoming_stamp, previous_stamp = clean(record.get("generated_at")), clean(previous.get("generated_at"))
            if incoming_stamp and previous_stamp and incoming_stamp < previous_stamp: return {"record": previous, "created": False, "updated": False, "stale": True}
            record["revision_count"] = int(previous.get("revision_count") or 1) + 1
            if preserve_revision: _append_revision(previous, record)
            rows[index[rid]] = record
        else: rows.append(record)
        _write_jsonl(RECORD_FILE, rows)
    if queue_sync: queue_emis_event_safely("academic_history", "record_upsert", {"record": record}, entity_key=f"history|{rid}")
    return {"record": record, "created": previous is None, "updated": previous is not None}


def archive_report_snapshots(snapshots, queue_sync=True):
    results = []
    for snapshot in snapshots or []:
        record = _snapshot_to_record(snapshot)
        if not record: continue
        results.append(upsert_academic_record(record, queue_sync=queue_sync))
    return results


def rebuild_archive_from_snapshots(queue_sync=False):
    snapshots = _read_jsonl(SNAPSHOT_FILE); snapshots.sort(key=lambda row: clean(row.get("generated_at")))
    created = updated = skipped = 0
    for snapshot in snapshots:
        record = _snapshot_to_record(snapshot)
        if not record: skipped += 1; continue
        result = upsert_academic_record(record, queue_sync=False)
        created += int(result.get("created")); updated += int(result.get("updated")); skipped += int(not result.get("created") and not result.get("updated"))
    records = read_academic_records()
    if queue_sync:
        for record in records: queue_emis_event_safely("academic_history", "record_upsert", {"record": record}, entity_key=f"history|{clean(record.get('record_id'))}")
    return {"snapshots": len(snapshots), "created": created, "updated": updated, "skipped": skipped, "records": len(records), "sync_records_queued": len(records) if queue_sync else 0}


def ensure_history_bootstrap():
    if RECORD_FILE.exists() and RECORD_FILE.stat().st_size > 0: return
    if SNAPSHOT_FILE.exists() and SNAPSHOT_FILE.stat().st_size > 0:
        try: rebuild_archive_from_snapshots(queue_sync=False)
        except Exception as error: print("[EMIS ACADEMIC HISTORY] Initial backfill warning:", error)


def _session_sort(value):
    try: return int(clean(value).split("/")[0])
    except Exception: return 0


def _record_sort(row): return (_session_sort(row.get("session")), TERM_ORDER.get(normalize_academic_term(row.get("term")), 0), clean(row.get("generated_at")))


def _summary_record(record):
    summary = {key: record.get(key) for key in ("record_id", "student_key", "admission_number", "full_name", "sex", "session", "term", "term_label", "class_level", "class_arm", "result_year", "snapshot_id", "result_name", "generated_at", "source_mode", "academic_complete", "subject_count", "subjects_complete", "total_score", "maximum_score", "score_percentage", "average", "grade", "position", "position_text", "out_of", "attendance_percentage", "revision_count")}; bundle = _score_bundle(record); summary["total_score"], summary["maximum_score"], summary["score_percentage"] = bundle.get("score"), bundle.get("maximum"), bundle.get("percentage"); summary["score_ratio"] = _score_ratio(bundle.get("score"), bundle.get("maximum")); return summary


def _student_profile(admission, records):
    active = None
    try: active = find_student_by_admission(admission)
    except Exception: active = None
    latest = max(records, key=_record_sort) if records else {}; latest_name = clean(latest.get("full_name")); active_name = clean((active or {}).get("full_name")); same_identity = bool(active and (not latest_name or _normalize_name(active_name) == _normalize_name(latest_name)))
    return {"admission_number": admission, "full_name": latest_name or active_name, "sex": clean(latest.get("sex") or ((active or {}).get("sex") if same_identity else "")), "current_class": clean(((active or {}).get("class_arm") or (active or {}).get("class_level")) if same_identity else (latest.get("class_arm") or latest.get("class_level"))), "active": same_identity, "student_key": clean(latest.get("student_key"))}


def dashboard_summary(records):
    students = {(clean(row.get("student_key")), normalize_admission_number(row.get("admission_number"))) for row in records}; sessions = sorted({clean(row.get("session")) for row in records if clean(row.get("session"))}, key=_session_sort, reverse=True); complete = sum(1 for row in records if row.get("academic_complete")); averages = [safe_float(row.get("average")) for row in records]; averages = [value for value in averages if value is not None]
    return {"students": len(students), "term_records": len(records), "sessions": len(sessions), "session_values": sessions, "complete_records": complete, "partial_records": max(0, len(records) - complete), "average": round(sum(averages) / len(averages), 2) if averages else None, "revisions": len(read_revisions())}


def search_history(query_text="", academic_session="", class_level="", term="", limit=200):
    ensure_history_bootstrap(); token = clean(query_text).lower(); academic_session = normalize_academic_session(academic_session); class_level = clean_upper(class_level); term = normalize_academic_term(term); rows = read_academic_records()
    filtered = []
    for row in rows:
        haystack = f"{row.get('admission_number','')} {row.get('full_name','')} {row.get('class_level','')} {row.get('class_arm','')} {row.get('session','')}".lower()
        if token and token not in haystack: continue
        if academic_session and clean(row.get("session")) != academic_session: continue
        if class_level and clean_upper(row.get("class_level")) != class_level: continue
        if term and normalize_academic_term(row.get("term")) != term: continue
        filtered.append(row)
    filtered.sort(key=_record_sort, reverse=True); filtered = filtered[:max(1, min(int(limit or 200), 500))]
    return [_summary_record(row) for row in filtered]


def student_history(admission_number, student_key=""):
    ensure_history_bootstrap(); admission = normalize_admission_number(admission_number); requested_key = clean(student_key); rows = [row for row in read_academic_records() if normalize_admission_number(row.get("admission_number")) == admission]
    identities = {}
    for row in rows:
        key = clean(row.get("student_key")); identities.setdefault(key, {"student_key": key, "full_name": clean(row.get("full_name")), "sex": clean(row.get("sex")), "records": 0}); identities[key]["records"] += 1
    if requested_key: rows = [row for row in rows if clean(row.get("student_key")) == requested_key]
    elif len(identities) == 1: requested_key = next(iter(identities)); rows = [row for row in rows if clean(row.get("student_key")) == requested_key]
    rows.sort(key=_record_sort, reverse=True); profile = _student_profile(admission, rows); sessions = []
    for session_value in sorted({clean(row.get("session")) for row in rows}, key=_session_sort, reverse=True):
        period_rows = [row for row in rows if clean(row.get("session")) == session_value]; period_rows.sort(key=lambda row: TERM_ORDER.get(normalize_academic_term(row.get("term")), 0)); averages = [safe_float(row.get("average")) for row in period_rows]; averages = [value for value in averages if value is not None]; bundles = [_score_bundle(row) for row in period_rows]; annual_score = round(sum(float(item["score"]) for item in bundles if item.get("score") is not None), 2) if any(item.get("score") is not None for item in bundles) else None; annual_max = round(sum(float(item["maximum"]) for item in bundles if item.get("score") is not None and item.get("maximum")), 2) if annual_score is not None else None; sessions.append({"session": session_value, "class_level": clean(period_rows[-1].get("class_level")) if period_rows else "", "class_arm": clean(period_rows[-1].get("class_arm")) if period_rows else "", "terms": [_summary_record(row) for row in period_rows], "annual_score": annual_score, "annual_max": annual_max, "annual_score_percentage": round(annual_score / annual_max * 100, 2) if annual_score is not None and annual_max else None, "annual_average": round(sum(averages) / len(averages), 2) if averages else None, "terms_available": len(averages)})
    return {"profile": profile, "identities": list(identities.values()), "ambiguous_identity": len(identities) > 1 and not requested_key, "records": [_summary_record(row) for row in rows], "sessions": sessions}



def _transcript_subject_key(row): return clean_upper(row.get("subject_key") or row.get("subject"))

def _transcript_subject_rows(period_rows):
    subject_names, term_maps = {}, {}
    for record in period_rows:
        term = normalize_academic_term(record.get("term")); report = record.get("report") if isinstance(record.get("report"), dict) else {}; mapping = {}
        for row in report.get("subjects") or []:
            if not isinstance(row, dict): continue
            key = _transcript_subject_key(row)
            if not key: continue
            mapping[key] = row; subject_names.setdefault(key, clean(row.get("subject")) or key.title())
        term_maps[term] = mapping
    rows = []
    for key, name in subject_names.items():
        values = {term: safe_float((term_maps.get(term, {}).get(key) or {}).get("total")) for term in TERM_ORDER}; available = [value for value in values.values() if value is not None]; annual_score = round(sum(available), 2) if available else None; annual_max = len(available) * 100 if available else None; annual_average = round(annual_score / annual_max * 100, 2) if annual_score is not None and annual_max else None; grade, comment = grade_score(annual_average, "REPORT_SECONDARY")
        rows.append({"subject_key": key, "subject": name, "first_term": values.get("FIRST"), "second_term": values.get("SECOND"), "third_term": values.get("THIRD"), "first_term_max": 100 if values.get("FIRST") is not None else None, "second_term_max": 100 if values.get("SECOND") is not None else None, "third_term_max": 100 if values.get("THIRD") is not None else None, "annual_score": annual_score, "annual_max": annual_max, "annual_average": annual_average, "grade": grade, "comment": comment, "terms_available": len(available), "complete": len(available) == 3})
    rows.sort(key=lambda row: clean(row.get("subject")).lower()); return rows


def build_student_transcript(admission_number, student_key=""):
    ensure_history_bootstrap(); admission = normalize_admission_number(admission_number); requested_key = clean(student_key); all_rows = [row for row in read_academic_records() if normalize_admission_number(row.get("admission_number")) == admission]
    identities = {}
    for row in all_rows:
        key = clean(row.get("student_key")); identities.setdefault(key, {"student_key": key, "full_name": clean(row.get("full_name")), "sex": clean(row.get("sex")), "records": 0}); identities[key]["records"] += 1
    if requested_key: rows = [row for row in all_rows if clean(row.get("student_key")) == requested_key]
    elif len(identities) == 1: requested_key = next(iter(identities)); rows = [row for row in all_rows if clean(row.get("student_key")) == requested_key]
    else: rows = all_rows
    if len(identities) > 1 and not requested_key: return {"ambiguous_identity": True, "identities": list(identities.values()), "profile": {"admission_number": admission}, "sessions": []}
    rows.sort(key=_record_sort); profile = _student_profile(admission, rows); sessions = []
    for session_value in sorted({clean(row.get("session")) for row in rows if clean(row.get("session"))}, key=_session_sort):
        period_rows = [row for row in rows if clean(row.get("session")) == session_value]; period_rows.sort(key=lambda row: TERM_ORDER.get(normalize_academic_term(row.get("term")), 0)); term_map = {normalize_academic_term(row.get("term")): row for row in period_rows if normalize_academic_term(row.get("term")) in TERM_ORDER}; term_averages = {term: safe_float((term_map.get(term) or {}).get("average")) for term in TERM_ORDER}; available = [value for value in term_averages.values() if value is not None]; annual_average = round(sum(available) / len(available), 2) if available else None; annual_grade, annual_comment = grade_score(annual_average, "REPORT_SECONDARY"); subject_rows = _transcript_subject_rows(period_rows); attendance_values = [safe_float(row.get("attendance_percentage")) for row in period_rows]; attendance_values = [value for value in attendance_values if value is not None]; term_payload = {}; session_score = 0.0; session_max = 0.0; score_terms = 0
        for term in TERM_ORDER:
            row = term_map.get(term); bundle = _score_bundle(row) if row else {"score": None, "maximum": None, "percentage": None}
            if bundle.get("score") is not None and bundle.get("maximum"): session_score += float(bundle["score"]); session_max += float(bundle["maximum"]); score_terms += 1
            term_payload[term] = {"available": bool(row), "score": bundle.get("score"), "maximum": bundle.get("maximum"), "score_percentage": bundle.get("percentage"), "average": term_averages.get(term), "grade": clean((row or {}).get("grade")) or "-", "academic_complete": bool((row or {}).get("academic_complete"))}
        session_score = round(session_score, 2) if score_terms else None; session_max = round(session_max, 2) if score_terms else None
        sessions.append({"session": session_value, "class_level": clean(period_rows[-1].get("class_level")) if period_rows else "", "class_arm": clean(period_rows[-1].get("class_arm")) if period_rows else "", "terms": term_payload, "terms_available": len(available), "complete": len(available) == 3, "annual_score": session_score, "annual_max": session_max, "annual_score_percentage": round(session_score / session_max * 100, 2) if session_score is not None and session_max else None, "annual_average": annual_average, "annual_grade": annual_grade, "annual_comment": annual_comment, "attendance_average": round(sum(attendance_values) / len(attendance_values), 2) if attendance_values else None, "subjects": subject_rows, "subject_count": len(subject_rows), "subjects_complete": sum(1 for row in subject_rows if row.get("complete"))})
    annual_values = [safe_float(item.get("annual_average")) for item in sessions]; annual_values = [value for value in annual_values if value is not None]; overall_average = round(sum(annual_values) / len(annual_values), 2) if annual_values else None; overall_grade, overall_comment = grade_score(overall_average, "REPORT_SECONDARY"); overall_score = round(sum(float(item.get("annual_score")) for item in sessions if item.get("annual_score") is not None), 2) if any(item.get("annual_score") is not None for item in sessions) else None; overall_max = round(sum(float(item.get("annual_max")) for item in sessions if item.get("annual_score") is not None and item.get("annual_max")), 2) if overall_score is not None else None; record_ids = "|".join(clean(row.get("record_id")) for row in rows); reference = hashlib.sha256(f"{requested_key}|{record_ids}".encode("utf-8")).hexdigest()[:12].upper() if rows else ""; complete_sessions = sum(1 for item in sessions if item.get("complete")); subject_keys = {_transcript_subject_key(subject) for row in rows for subject in ((row.get("report") or {}).get("subjects") or []) if isinstance(subject, dict) and _transcript_subject_key(subject)}
    return {"ambiguous_identity": False, "identities": list(identities.values()), "profile": profile, "student_key": requested_key or clean(profile.get("student_key")), "sessions": sessions, "summary": {"session_count": len(sessions), "term_records": len(rows), "complete_sessions": complete_sessions, "partial_sessions": max(0, len(sessions) - complete_sessions), "distinct_subjects": len(subject_keys), "overall_score": overall_score, "overall_max": overall_max, "overall_score_percentage": round(overall_score / overall_max * 100, 2) if overall_score is not None and overall_max else None, "overall_average": overall_average, "overall_grade": overall_grade, "overall_comment": overall_comment, "archive_status": "Complete" if sessions and complete_sessions == len(sessions) else "Partial", "first_session": sessions[0].get("session") if sessions else "", "latest_session": sessions[-1].get("session") if sessions else ""}, "transcript_reference": f"EMIS-TR-{reference}" if reference else "", "generated_at": now_iso()}


def _transcript_sheet_style(sheet, header_row=1):
    navy, blue, pale, white, line = "173F76", "2563EB", "EEF5FD", "FFFFFF", "D7E3F0"; thin = Side(style="thin", color=line); sheet.sheet_view.showGridLines = False
    for cell in sheet[header_row]: cell.font = Font(bold=True, color=white); cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row: cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_transcript_workbook(transcript):
    workbook = Workbook(); overview = workbook.active; overview.title = "Transcript Overview"; profile, summary = transcript.get("profile") or {}, transcript.get("summary") or {}
    overview.merge_cells("A1:J1"); overview["A1"] = "EPITOME MODEL ISLAMIC SCHOOLS — STUDENT ACADEMIC TRANSCRIPT"; overview["A1"].font = Font(bold=True, size=16, color="173F76"); overview["A1"].alignment = Alignment(horizontal="center")
    overview.append(["Student", profile.get("full_name"), "Admission No.", profile.get("admission_number"), "Sex", profile.get("sex"), "Reference", transcript.get("transcript_reference"), "", ""]); overview.append(["Sessions", summary.get("session_count"), "Term Records", summary.get("term_records"), "Overall Score", _score_ratio(summary.get("overall_score"), summary.get("overall_max")), "Overall %", summary.get("overall_average"), "Grade", summary.get("overall_grade")]); overview.append(["Archive Status", summary.get("archive_status"), "First Session", summary.get("first_session"), "Latest Session", summary.get("latest_session"), "Generated", transcript.get("generated_at"), "", ""]); overview.append([]); overview.append(["Session", "Class", "1st Term Score", "1st %", "2nd Term Score", "2nd %", "3rd Term Score", "3rd %", "Annual Score", "Annual %"]); _transcript_sheet_style(overview, 6)
    for period in transcript.get("sessions") or []:
        terms = period.get("terms") or {}; overview.append([period.get("session"), period.get("class_arm") or period.get("class_level"), _score_ratio((terms.get("FIRST") or {}).get("score"), (terms.get("FIRST") or {}).get("maximum")), (terms.get("FIRST") or {}).get("average"), _score_ratio((terms.get("SECOND") or {}).get("score"), (terms.get("SECOND") or {}).get("maximum")), (terms.get("SECOND") or {}).get("average"), _score_ratio((terms.get("THIRD") or {}).get("score"), (terms.get("THIRD") or {}).get("maximum")), (terms.get("THIRD") or {}).get("average"), _score_ratio(period.get("annual_score"), period.get("annual_max")), period.get("annual_average")])
    for column, width in enumerate([16, 20, 18, 10, 18, 10, 18, 10, 20, 12], 1): overview.column_dimensions[get_column_letter(column)].width = width
    for index, period in enumerate(transcript.get("sessions") or [], 1):
        title = f"{index}_{clean(period.get('session')).replace('/', '-')}_{clean(period.get('class_level'))}"[:31]; sheet = workbook.create_sheet(title); sheet.merge_cells("A1:I1"); sheet["A1"] = f"{period.get('session')} — {period.get('class_arm') or period.get('class_level')}"; sheet["A1"].font = Font(bold=True, size=14, color="173F76"); sheet["A1"].alignment = Alignment(horizontal="center"); terms = period.get("terms") or {}; sheet.append(["1st Term", _score_ratio((terms.get("FIRST") or {}).get("score"), (terms.get("FIRST") or {}).get("maximum")), "2nd Term", _score_ratio((terms.get("SECOND") or {}).get("score"), (terms.get("SECOND") or {}).get("maximum")), "3rd Term", _score_ratio((terms.get("THIRD") or {}).get("score"), (terms.get("THIRD") or {}).get("maximum")), "Annual", _score_ratio(period.get("annual_score"), period.get("annual_max")), ""]); sheet.append(["Annual Average %", period.get("annual_average"), "Annual Grade", period.get("annual_grade"), "Attendance %", period.get("attendance_average"), "Terms", period.get("terms_available"), ""]); sheet.append([]); sheet.append(["S/N", "Subject", "1st Term /100", "2nd Term /100", "3rd Term /100", "Annual Score", "Annual Max", "Annual %", "Grade"]); _transcript_sheet_style(sheet, 5)
        for row_index, row in enumerate(period.get("subjects") or [], 1): sheet.append([row_index, row.get("subject"), row.get("first_term"), row.get("second_term"), row.get("third_term"), row.get("annual_score"), row.get("annual_max"), row.get("annual_average"), row.get("grade")])
        for column, width in enumerate([7, 28, 14, 14, 14, 14, 14, 12, 10], 1): sheet.column_dimensions[get_column_letter(column)].width = width
    return workbook


def get_record(record_id):
    ensure_history_bootstrap(); rid = clean(record_id)
    return next((row for row in read_academic_records() if clean(row.get("record_id")) == rid), None)


def apply_academic_history_sync_event(action, payload, event=None):
    if not isinstance(payload, dict): raise ValueError("Academic History sync payload must be an object.")
    action = clean(action).lower()
    if action == "record_upsert":
        record = payload.get("record")
        if not isinstance(record, dict) or not clean(record.get("record_id")): raise ValueError("record_upsert requires a valid academic history record.")
        result = upsert_academic_record(record, queue_sync=False, preserve_revision=True)
        return {"record_id": record.get("record_id"), "created": result.get("created"), "updated": result.get("updated")}
    raise ValueError(f"Unsupported Academic History sync action: {action}")


@academic_history_bp.route("/api/academic-history/config")
@admin_required
def api_history_config():
    ensure_history_bootstrap(); records = read_academic_records(); classes = sorted({clean_upper(row.get("class_level")) for row in records if clean(row.get("class_level"))}); return jsonify({"success": True, "dashboard": dashboard_summary(records), "classes": classes, "terms": ["FIRST", "SECOND", "THIRD"]})


@academic_history_bp.route("/api/academic-history/search")
@admin_required
def api_history_search():
    rows = search_history(request.args.get("q"), request.args.get("session"), request.args.get("class_level"), request.args.get("term"), request.args.get("limit", 200)); return jsonify({"success": True, "count": len(rows), "records": rows})


@academic_history_bp.route("/api/academic-history/student/<admission_number>")
@admin_required
def api_student_history(admission_number):
    history = student_history(admission_number, request.args.get("student_key")); return jsonify({"success": True, **history})


@academic_history_bp.route("/api/academic-history/record/<record_id>")
@admin_required
def api_history_record(record_id):
    record = get_record(record_id)
    if not record: return jsonify({"success": False, "message": "Academic history record was not found."}), 404
    return jsonify({"success": True, "record": record})


@academic_history_bp.route("/api/academic-history/rebuild", methods=["POST"])
@admin_required
def api_history_rebuild():
    result = rebuild_archive_from_snapshots(queue_sync=True); return jsonify({"success": True, "message": f"Academic archive rebuilt from {result['snapshots']} saved report snapshot(s).", **result, "dashboard": dashboard_summary(read_academic_records())})



@academic_history_bp.route("/api/academic-history/transcript/<admission_number>")
@admin_required
def api_student_transcript(admission_number):
    transcript = build_student_transcript(admission_number, request.args.get("student_key"));
    if transcript.get("ambiguous_identity"): return jsonify({"success": False, "message": "This admission number has more than one historical student identity. Select the correct student identity before compiling a transcript.", **transcript}), 409
    if not transcript.get("sessions"): return jsonify({"success": False, "message": "No archived academic records were found for this student."}), 404
    return jsonify({"success": True, **transcript})


@academic_history_bp.route("/api/academic-history/transcript/export/<admission_number>")
@admin_required
def api_student_transcript_export(admission_number):
    transcript = build_student_transcript(admission_number, request.args.get("student_key"));
    if transcript.get("ambiguous_identity"): return jsonify({"success": False, "message": "Select the correct historical student identity before transcript export."}), 409
    if not transcript.get("sessions"): return jsonify({"success": False, "message": "No archived academic records were found for this student."}), 404
    workbook = build_transcript_workbook(transcript); output = io.BytesIO(); workbook.save(output); output.seek(0); profile = transcript.get("profile") or {}; safe_name = "_".join(clean(profile.get("full_name") or admission_number).split()) or normalize_admission_number(admission_number)
    return send_file(output, as_attachment=True, download_name=f"{safe_name}_Academic_Transcript.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@academic_history_bp.route("/api/academic-history/export/<admission_number>")
@admin_required
def api_history_export(admission_number):
    student_key = clean(request.args.get("student_key")); history = student_history(admission_number, student_key); records = history.get("records") or []
    if history.get("ambiguous_identity"): return jsonify({"success": False, "message": "This admission number has more than one historical student identity. Select the student identity before export."}), 409
    if not records: return jsonify({"success": False, "message": "No archived academic records were found for this student."}), 404
    profile = history.get("profile") or {}; workbook = Workbook(); summary_sheet = workbook.active; summary_sheet.title = "Academic History"
    summary_sheet.append(["EMIS STUDENT ACADEMIC HISTORY"]); summary_sheet.append(["Student", profile.get("full_name")]); summary_sheet.append(["Admission Number", profile.get("admission_number")]); summary_sheet.append(["Sex", profile.get("sex")]); summary_sheet.append([]); summary_sheet.append(["Session", "Term", "Class", "Score", "Max", "Average %", "Grade", "Position", "Out Of", "Attendance %", "Status"])
    for row in records: summary_sheet.append([row.get("session"), row.get("term_label") or row.get("term"), row.get("class_arm") or row.get("class_level"), row.get("total_score"), row.get("maximum_score"), row.get("average"), row.get("grade"), row.get("position_text"), row.get("out_of"), row.get("attendance_percentage"), "Complete" if row.get("academic_complete") else "Partial"])
    summary_sheet["A1"].font = Font(bold=True, size=16, color="173F76"); summary_sheet.merge_cells("A1:K1")
    for cell in summary_sheet[6]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173F76"); cell.alignment = Alignment(horizontal="center")
    for width, column in zip([16, 16, 18, 12, 12, 12, 10, 12, 10, 14, 12], range(1, 12)): summary_sheet.column_dimensions[get_column_letter(column)].width = width
    detail_sheet = workbook.create_sheet("Subject Records"); detail_sheet.append(["Session", "Term", "Class", "Subject", "CA", "Exam", "Total", "Grade", "Position", "Comment"])
    for item in read_academic_records():
        if normalize_admission_number(item.get("admission_number")) != normalize_admission_number(admission_number): continue
        if student_key and clean(item.get("student_key")) != student_key: continue
        report = item.get("report") if isinstance(item.get("report"), dict) else {}
        for subject in report.get("subjects") or []:
            if not isinstance(subject, dict): continue
            detail_sheet.append([item.get("session"), item.get("term_label") or item.get("term"), item.get("class_arm") or item.get("class_level"), subject.get("subject"), subject.get("ca_total"), subject.get("exam"), subject.get("total"), subject.get("grade"), subject.get("position_text"), subject.get("comment")])
    for cell in detail_sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173F76"); cell.alignment = Alignment(horizontal="center")
    for column in range(1, 11): detail_sheet.column_dimensions[get_column_letter(column)].width = 18 if column < 4 else 14
    output = io.BytesIO(); workbook.save(output); output.seek(0); safe_name = "_".join(clean(profile.get("full_name") or admission_number).split()) or normalize_admission_number(admission_number)
    return send_file(output, as_attachment=True, download_name=f"{safe_name}_Academic_History.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
