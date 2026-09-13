# modules/excel_manager.py

from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook

from modules.class_config import SUPPORTED_CLASSES, normalize_class_level, normalize_class_arm


BASE_DIR = Path(__file__).resolve().parent.parent
RESULTS_BASE = BASE_DIR / "RESULTS"
VALID_TERMS = ("FIRST", "SECOND", "THIRD")


# =========================================================
# GENERAL HELPERS
# =========================================================

def ensure_folder(path: Path):
    path.mkdir(parents=True, exist_ok=True)
    return path


def normalize_class_category(class_category):
    level = normalize_class_level(class_category)
    return level if level in SUPPORTED_CLASSES else "UNKNOWN"


def normalize_result_class_arm(class_arm, fallback_level=""):
    arm = normalize_class_arm(class_arm, fallback_level)
    if arm: return arm

    level = normalize_class_category(fallback_level or class_arm)
    return level if level != "UNKNOWN" else "UNKNOWN"


def is_jss_class(class_category):
    level = normalize_class_category(class_category)
    return level != "UNKNOWN" and level.startswith("JSS")


def is_ss_class(class_category):
    level = normalize_class_category(class_category)
    return level != "UNKNOWN" and level.startswith("SS")


# =========================================================
# TERM HELPERS
# =========================================================

def normalize_result_term(value):
    raw = " ".join(str(value or "").strip().upper().replace("_", " ").replace("-", " ").split())

    aliases = {
        "FIRST": "FIRST", "FIRST TERM": "FIRST", "TERM 1": "FIRST", "TERM ONE": "FIRST",
        "1": "FIRST", "1ST": "FIRST", "1ST TERM": "FIRST",

        "SECOND": "SECOND", "SECOND TERM": "SECOND", "TERM 2": "SECOND", "TERM TWO": "SECOND",
        "2": "SECOND", "2ND": "SECOND", "2ND TERM": "SECOND",

        "THIRD": "THIRD", "THIRD TERM": "THIRD", "TERM 3": "THIRD", "TERM THREE": "THIRD",
        "3": "THIRD", "3RD": "THIRD", "3RD TERM": "THIRD",
    }

    return aliases.get(raw, "")


def result_term_label(value):
    return {"FIRST": "FIRST TERM", "SECOND": "SECOND TERM", "THIRD": "THIRD TERM"}.get(normalize_result_term(value), "")


# =========================================================
# SUBJECT NORMALIZATION
# =========================================================

def normalize_subject(subject):
    if not subject: return "Unknown"

    key = str(subject).strip().lower()

    mapping = {
        "biology": "Biology", "chemistry": "Chemistry",
        "civic education": "Civic_education", "civic": "Civic_education",
        "computer science": "Computer_studies", "computer studies": "Computer_studies", "computer": "Computer_studies",
        "economics": "Economics",
        "english language": "English_language", "english": "English_language",

        "financial accounting": "Financial_accounting", "financial account": "Financial_accounting",
        "accounting": "Financial_accounting", "accounts": "Financial_accounting", "account": "Financial_accounting",

        "geography": "Geography", "government": "Government",
        "literature-in-english": "Literature", "literature in english": "Literature", "literature": "Literature",
        "mathematics": "Mathematics", "maths": "Mathematics",
        "physics": "Physics",

        "technical drawing": "Technical", "technical": "Technical",
        "yoruba language": "Yoruba_language", "yoruba": "Yoruba_language",
        "history": "History",

        "irk": "IRK", "irs": "IRS", "cca": "CCA",
        "arabic language": "Arabic_language", "arabic": "Arabic_language",
        "business studies": "Business_studies",

        "poise": "Poise",
        "islamiyyah": "Islamiyyah", "islamiyah": "Islamiyyah",

        "hort & crop production": "Hort_crop_production",
        "hort and crop production": "Hort_crop_production",
        "horticulture and crop production": "Hort_crop_production",

        "digital tech.": "Digital_tech", "digital tech": "Digital_tech", "digital technology": "Digital_tech",
        "inter science": "Inter_science", "integrated science": "Inter_science",
        "garment making": "Garment_making",

        "soc. & cit. std": "Soc_cit_std", "soc & cit std": "Soc_cit_std",
        "social and citizenship studies": "Soc_cit_std",

        "p.h.e": "PHE", "p.h.e.": "PHE", "phe": "PHE", "physical health education": "PHE",
        "bst": "BST",
        "national value": "National_value", "national values": "National_value",
        "pvs": "PVS",

        "hausa language": "Hausa_language", "hausa": "Hausa_language",

        "cit & her. std": "Cit_her_std", "cit & her std": "Cit_her_std", "cit and her std": "Cit_her_std",
        "citizenship and heritage studies": "Cit_her_std",
        "heritage and citizenship studies": "Cit_her_std",
        "heritage & citizenship studies": "Cit_her_std",
        "civic and heritage studies": "Cit_her_std",

        "commerce": "Commerce",
        "marketing": "Marketing", "marketting": "Marketing",

        "further mathematics": "Further_mathematics", "further maths": "Further_mathematics",
        "agricultural science": "Agricultural_science", "agriculture": "Agricultural_science",
    }

    if key in mapping: return mapping[key]

    return str(subject).strip().replace("&", "and").replace(".", "").replace("-", "_").replace("/", "_").replace(" ", "_").capitalize()


# =========================================================
# RESULT PATHS
# =========================================================

def get_excel_path(class_category, subject, year, term=""):
    year = str(year or datetime.now().year).strip()
    class_level = normalize_class_category(class_category)
    subject_folder = normalize_subject(subject)
    normalized_term = normalize_result_term(term)
    root = RESULTS_BASE / year / "CLASS" / class_level

    if normalized_term: return root / normalized_term / subject_folder / "results.xlsx"
    return root / subject_folder / "results.xlsx"


def get_preferred_excel_path(class_category, subject, year, term=""):
    class_level = normalize_class_category(class_category)
    normalized_term = normalize_result_term(term)

    return get_excel_path(class_level, subject, year, normalized_term) if normalized_term else get_excel_path(class_level, subject, year)


def find_existing_excel_path(class_category, subject, year, term=""):
    class_level = normalize_class_category(class_category)
    normalized_term = normalize_result_term(term)

    if normalized_term:
        term_path = get_excel_path(class_level, subject, year, normalized_term)

        if term_path.exists():
            return term_path

        legacy_path = get_excel_path(class_level, subject, year)

        if is_jss_class(class_level) and legacy_path.exists():
            return legacy_path

        if is_ss_class(class_level) and legacy_path.exists() and excel_file_has_term(legacy_path, normalized_term):
            return legacy_path

        return term_path

    return get_excel_path(class_level, subject, year)


# =========================================================
# EXCEL HEADERS
# =========================================================

EXPECTED_HEADERS = [
    "Student Name", "Admission No", "Class Level", "Class Arm", "Subject", "Term",
    "Score (%)", "Correct", "Total", "Status", "Time Taken", "Submitted At",
    "Session", "Submission Status", "Termination Reason",
]


HEADER_ALIASES = {
    "Student Name": {"Student Name", "Name", "Full Name", "Full_Name"},
    "Admission No": {"Admission No", "Admission Number", "Admission_number", "Admission No.", "Admission"},
    "Class Level": {"Class Level", "Class Category", "Class_category", "ClassLevel"},
    "Class Arm": {"Class Arm", "Class", "Class Name", "Class_arm"},
    "Subject": {"Subject"},
    "Term": {"Term", "Exam Term", "Academic Term", "Term Name"},
    "Score (%)": {"Score (%)", "Score", "Percentage", "Percent"},
    "Correct": {"Correct", "Correct Answers"},
    "Total": {"Total", "Total Questions"},
    "Status": {"Status", "Result", "Result Status"},
    "Time Taken": {"Time Taken", "Time", "Duration"},
    "Submitted At": {"Submitted At", "Submitted", "Date", "Timestamp"},
    "Session": {"Session", "Academic Session", "Academic_session"},
    "Submission Status": {"Submission Status", "Exam Status", "Submission_status"},
    "Termination Reason": {"Termination Reason", "Security Reason", "Termination_reason"},
}


def normalize_header_name(value):
    text = str(value or "").strip()

    for canonical, aliases in HEADER_ALIASES.items():
        if text in aliases:
            return canonical

    return text


# =========================================================
# HEADER REPAIR
# =========================================================

def repair_missing_headers(ws):
    if ws.max_row < 1 or ws.max_column < 1:
        for index, header in enumerate(EXPECTED_HEADERS, 1):
            ws.cell(1, index).value = header

        return EXPECTED_HEADERS.copy()

    raw_headers = [ws.cell(1, index).value for index in range(1, ws.max_column + 1)]
    normalized_headers = [normalize_header_name(value) for value in raw_headers]
    header_matches = sum(1 for value in normalized_headers if value in EXPECTED_HEADERS)

    if header_matches < 2:
        ws.insert_rows(1)

        legacy_headers = [
            "Student Name", "Admission No", "Class Level", "Class Arm", "Subject",
            "Score (%)", "Correct", "Total", "Status", "Time Taken", "Submitted At",
        ]

        for index, header in enumerate(legacy_headers, 1):
            ws.cell(1, index).value = header

        normalized_headers = legacy_headers.copy()

    else:
        for index, header in enumerate(normalized_headers, 1):
            ws.cell(1, index).value = header

    existing = set(normalized_headers)

    for header in EXPECTED_HEADERS:
        if header in existing:
            continue

        ws.cell(1, ws.max_column + 1).value = header
        normalized_headers.append(header)
        existing.add(header)

    return normalized_headers


# =========================================================
# ROW MAPPING
# =========================================================

def map_result_row(row, headers=None):
    values = list(row or [])

    if not any(value is not None and str(value).strip() != "" for value in values):
        return None

    if headers:
        normalized_headers = [normalize_header_name(header) for header in headers]
        row_dict = {header: values[index] if index < len(values) else None for index, header in enumerate(normalized_headers) if header}

        for header in EXPECTED_HEADERS:
            row_dict.setdefault(header, None)

    else:
        if len(values) >= len(EXPECTED_HEADERS):
            mapped = values[:len(EXPECTED_HEADERS)]

        elif len(values) >= 12:
            mapped = values[:12] + [None] * max(0, len(EXPECTED_HEADERS) - 12)

        elif len(values) == 11:
            mapped = [
                values[0], values[1], values[2], values[3], values[4], None,
                values[5], values[6], values[7], values[8], values[9], values[10],
            ]

            mapped += [None] * (len(EXPECTED_HEADERS) - len(mapped))

        elif len(values) == 10:
            mapped = [
                values[0], values[1], values[2], values[2], values[3], None,
                values[4], values[5], values[6], values[7], values[8], values[9],
            ]

            mapped += [None] * (len(EXPECTED_HEADERS) - len(mapped))

        elif len(values) == 9:
            mapped = [
                values[0], values[1], values[2], values[2], values[3], None,
                values[4], values[5], values[6], values[7], None, values[8],
            ]

            mapped += [None] * (len(EXPECTED_HEADERS) - len(mapped))

        else:
            mapped = values[:len(EXPECTED_HEADERS)] + [None] * max(0, len(EXPECTED_HEADERS) - len(values))

        row_dict = dict(zip(EXPECTED_HEADERS, mapped))

    raw_level = row_dict.get("Class Level") or row_dict.get("Class Arm")
    raw_arm = row_dict.get("Class Arm") or raw_level

    class_level = normalize_class_category(raw_level)
    class_arm = normalize_result_class_arm(raw_arm, class_level)

    if class_level != "UNKNOWN":
        row_dict["Class Level"] = class_level

    if class_arm != "UNKNOWN":
        row_dict["Class Arm"] = class_arm

    term = normalize_result_term(row_dict.get("Term"))

    row_dict["Term"] = term
    row_dict["Term Label"] = result_term_label(term) if term else ""

    return row_dict


def excel_file_has_term(excel_path: Path, term):
    normalized_term = normalize_result_term(term)

    if not normalized_term or not excel_path.exists():
        return False

    try:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb.active

        headers = [normalize_header_name(value) for value in next(ws.iter_rows(values_only=True), [])]

        if "Term" not in headers:
            wb.close()
            return False

        index = headers.index("Term")

        for row in ws.iter_rows(min_row=2, values_only=True):
            if index < len(row) and normalize_result_term(row[index]) == normalized_term:
                wb.close()
                return True

        wb.close()

    except Exception:
        return False

    return False


# =========================================================
# WORKBOOK / APPEND
# =========================================================

def create_results_workbook(excel_path: Path):
    ensure_folder(excel_path.parent)

    wb = Workbook()
    ws = wb.active

    ws.title = "Results"
    ws.append(EXPECTED_HEADERS)

    wb.save(excel_path)


def append_result_to_excel(result):
    raw_class = (
        result.get("class_level")
        or result.get("class_category")
        or result.get("class_name")
        or result.get("class")
        or result.get("class_arm")
    )

    class_level = normalize_class_category(raw_class)

    class_arm = normalize_result_class_arm(
        result.get("class_arm")
        or result.get("class")
        or result.get("class_name")
        or raw_class,
        class_level,
    )

    subject = str(result.get("subject") or "UNKNOWN").strip()
    year = str(result.get("year") or datetime.now().year).strip()

    raw_term = str(result.get("term") or "").strip()
    term = normalize_result_term(raw_term)

    if class_level == "UNKNOWN":
        raise ValueError(f"Invalid result class: {raw_class}")

    if is_jss_class(class_level) and not term:
        raise ValueError(f"Term is required when saving JSS result: class={class_level}, subject={subject}, year={year}")

    if raw_term and not term:
        raise ValueError(f"Invalid result term: {raw_term}")

    excel_path = get_preferred_excel_path(class_level, subject, year, term)

    if not excel_path.exists():
        create_results_workbook(excel_path)

    wb = load_workbook(excel_path)
    ws = wb.active

    headers = repair_missing_headers(ws)

    try:
        explicit_score = result.get("score_percentage", result.get("percentage", result.get("Score (%)")))

        if explicit_score not in (None, ""):
            score_percent = int(round(float(str(explicit_score).replace("%", "").strip())))

        elif result.get("total") not in (None, "", 0, "0") and result.get("correct") not in (None, ""):
            score_percent = round((float(result.get("correct")) / float(result.get("total"))) * 100)

        else:
            score_percent = int(round(float(str(result.get("score", 0) or 0).replace("%", "").strip())))

        score_percent = max(0, min(100, score_percent))

    except (TypeError, ValueError, ZeroDivisionError):
        score_percent = 0

    raw_status = str(result.get("status") or "").strip().upper()
    status = str(result.get("result_status") or raw_status).strip().upper()

    if status not in {"PASS", "FAIL"}:
        status = "PASS" if score_percent >= 50 else "FAIL"

    submission_status = str(
        result.get("submission_status")
        or result.get("exam_status")
        or (raw_status if raw_status not in {"PASS", "FAIL", ""} else "COMPLETED")
    ).strip().upper()

    termination_reason = str(result.get("termination_reason") or result.get("security_reason") or "").strip()

    time_taken = result.get("time_taken") or result.get("timeTaken") or result.get("Time Taken") or 0
    submitted_at = result.get("submitted_at") or result.get("submittedAt") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    admission_number = result.get("admission_number") or result.get("admission_no") or result.get("student_id") or ""
    full_name = result.get("full_name") or result.get("student_name") or result.get("name") or ""

    row_data = {
        "Student Name": full_name,
        "Admission No": admission_number,
        "Class Level": class_level,
        "Class Arm": class_arm,
        "Subject": subject.upper(),
        "Term": term,
        "Score (%)": f"{score_percent}%",
        "Correct": result.get("correct", 0),
        "Total": result.get("total", 0),
        "Status": status,
        "Time Taken": time_taken,
        "Submitted At": submitted_at,
        "Session": result.get("academic_session") or result.get("session") or "",
        "Submission Status": submission_status,
        "Termination Reason": termination_reason,
    }

    ws.append([row_data.get(header, "") for header in headers])

    wb.save(excel_path)

    return True


# =========================================================
# READ RESULTS
# =========================================================

def read_results(class_category, subject, year, term=""):
    class_level = normalize_class_category(class_category)
    normalized_term = normalize_result_term(term)

    paths = []

    if normalized_term:
        paths.append((get_excel_path(class_level, subject, year, normalized_term), normalized_term))

    paths.append((get_excel_path(class_level, subject, year), ""))

    seen = set()

    for excel_path, folder_term in paths:
        if excel_path in seen:
            continue

        seen.add(excel_path)

        if not excel_path.exists():
            continue

        try:
            wb = load_workbook(excel_path)
            ws = wb.active

            headers = repair_missing_headers(ws)
            rows = list(ws.iter_rows(values_only=True))

            wb.save(excel_path)

        except Exception as error:
            print(f"RESULT READ ERROR [{excel_path}]:", error)
            continue

        if len(rows) < 2:
            return []

        results = []

        for row in rows[1:]:
            item = map_result_row(row, headers)

            if not item:
                continue

            row_term = normalize_result_term(item.get("Term"))

            if folder_term and not row_term:
                row_term = folder_term
                item["Term"] = folder_term
                item["Term Label"] = result_term_label(folder_term)

            if normalized_term:
                if folder_term:
                    if row_term and row_term != normalized_term:
                        continue

                elif is_ss_class(class_level):
                    # SS root-level result = General / legacy.
                    # Do not pretend old blank SS results belong to a selected term.
                    if row_term != normalized_term:
                        continue

                elif is_jss_class(class_level):
                    # Preserve old flat JSS compatibility.
                    if row_term and row_term != normalized_term:
                        continue

            results.append(item)

        if results or folder_term:
            return results

    return []


# =========================================================
# RESULT FILE ITERATOR
# =========================================================

def iter_class_result_files(class_category):
    class_level = normalize_class_category(class_category)

    if not RESULTS_BASE.exists() or class_level == "UNKNOWN":
        return

    for year_folder in RESULTS_BASE.iterdir():
        if not year_folder.is_dir():
            continue

        class_root = year_folder / "CLASS" / class_level

        if not class_root.exists():
            continue

        # FIRST / SECOND / THIRD — JSS + SS.
        for term_folder in class_root.iterdir():
            if not term_folder.is_dir():
                continue

            term = normalize_result_term(term_folder.name)

            if not term:
                continue

            for subject_folder in term_folder.iterdir():
                excel_file = subject_folder / "results.xlsx"

                if subject_folder.is_dir() and excel_file.exists():
                    yield {
                        "year": year_folder.name,
                        "term": term,
                        "subject_folder": subject_folder.name,
                        "path": excel_file,
                    }

        # Root General / legacy — JSS + SS.
        for subject_folder in class_root.iterdir():
            if not subject_folder.is_dir() or normalize_result_term(subject_folder.name):
                continue

            excel_file = subject_folder / "results.xlsx"

            if excel_file.exists():
                yield {
                    "year": year_folder.name,
                    "term": "",
                    "subject_folder": subject_folder.name,
                    "path": excel_file,
                }


# =========================================================
# STUDENT HISTORY
# =========================================================

def _read_student_history(identifier, class_category, by_admission=False):
    results = []

    target = (
        str(identifier or "").strip().lower()
        if by_admission
        else str(identifier or "").strip().upper()
    )

    if not target:
        return results

    for info in iter_class_result_files(class_category) or []:
        try:
            wb = load_workbook(info["path"])
            ws = wb.active

            headers = repair_missing_headers(ws)
            rows = list(ws.iter_rows(values_only=True))

            wb.save(info["path"])

        except Exception:
            continue

        for row in rows[1:]:
            item = map_result_row(row, headers)

            if not item:
                continue

            value = (
                str(item.get("Admission No", "")).strip().lower()
                if by_admission
                else str(item.get("Student Name", "")).strip().upper()
            )

            if value != target:
                continue

            item["Year"] = info["year"]
            item["Subject Folder"] = info["subject_folder"]

            if info["term"] and not item.get("Term"):
                item["Term"] = info["term"]
                item["Term Label"] = result_term_label(info["term"])

            results.append(item)

    results.sort(key=lambda item: str(item.get("Submitted At", "")), reverse=True)

    return results


def read_student_history(full_name, class_category):
    return _read_student_history(full_name, class_category, by_admission=False)


def read_student_history_by_admission(admission_number, class_category):
    return _read_student_history(admission_number, class_category, by_admission=True)