# exam_document_export.py
# EMIS CBT — Teacher-readable PDF / DOCX generator from examination JSON

import io
import os
import re
import base64
from pathlib import Path
from xml.sax.saxutils import escape

from flask import Blueprint, request, jsonify, send_file

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Table, TableStyle, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

from uploads import BASE_DIR, SUPPORTED_CLASSES, safe_filename, valid_year, normalize_class, normalize_term, is_term_aware_class, get_json_path, read_json_file


exam_document_export_bp = Blueprint("exam_document_export_bp", __name__)


# ============================================================
# DESIGN
# ============================================================

BLUE = "#3778B8"
BLUE_SOFT = "#F2F7FC"
BLUE_BORDER = "#D9E8F6"
TEXT = "#273444"
TEXT_SOFT = "#657487"
BORDER = "#E1E7ED"
SURFACE = "#FBFCFE"
WHITE = "#FFFFFF"

DOCX_BLUE = RGBColor(55, 120, 184)
DOCX_TEXT = RGBColor(39, 52, 68)
DOCX_SOFT = RGBColor(101, 116, 135)


# ============================================================
# PDF FONT
# ============================================================

PDF_FONT = "Helvetica"
PDF_FONT_BOLD = "Helvetica-Bold"

FONT_CANDIDATES = [
    Path(BASE_DIR) / "static" / "fonts" / "DejaVuSans.ttf",
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]

FONT_BOLD_CANDIDATES = [
    Path(BASE_DIR) / "static" / "fonts" / "DejaVuSans-Bold.ttf",
    Path("C:/Windows/Fonts/arialbd.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]

try:
    normal_font = next((p for p in FONT_CANDIDATES if p.exists()), None)
    bold_font = next((p for p in FONT_BOLD_CANDIDATES if p.exists()), None)

    if normal_font and bold_font:
        pdfmetrics.registerFont(TTFont("EMISFont", str(normal_font)))
        pdfmetrics.registerFont(TTFont("EMISFontBold", str(bold_font)))
        PDF_FONT, PDF_FONT_BOLD = "EMISFont", "EMISFontBold"
except Exception as exc:
    print(f"[EXAM EXPORT] Font registration warning: {exc}")


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):
    return str(value or "").strip()


def term_display(data, fallback=None):
    label = clean_text(data.get("term_label"))
    if label: return label

    term = normalize_term(data.get("term") or fallback)
    return {"FIRST": "1st Term", "SECOND": "2nd Term", "THIRD": "3rd Term"}.get(term, "Examination")


def format_minutes(minutes):
    try: minutes = int(minutes)
    except (TypeError, ValueError): return ""

    if minutes <= 0: return ""

    hours, mins = divmod(minutes, 60)

    if hours and mins: return f"{hours} hr {mins} mins"
    if hours == 1: return "1 hour"
    if hours: return f"{hours} hours"
    return f"{mins} minutes"


def teacher_exam_title(data):
    title = clean_text(data.get("exam_title"))
    if title: return title

    term = term_display(data)
    return f"{term} Examination" if term != "Examination" else "Examination"


def option_letter(option, index):
    text = clean_text(option)
    match = re.match(r"^\s*([A-Z])[\.\)\:\-]\s*", text, re.I)
    return match.group(1).upper() if match else chr(65 + index)


def normalize_option_text(option, index):
    text = clean_text(option)
    letter = option_letter(text, index)
    text = re.sub(r"^\s*[A-Z][\.\)\:\-]\s*", "", text, count=1, flags=re.I)
    return letter, text


def normalize_correct_option(correct_option):
    value = clean_text(correct_option).upper()
    match = re.match(r"^\s*([A-Z])(?:[\.\)\:\-\s]|$)", value)
    return match.group(1) if match else value[:1]


def is_correct_option(option, index, correct_option):
    correct = normalize_correct_option(correct_option)
    if not correct: return False
    return option_letter(option, index) == correct


def pdf_safe(value):
    text = clean_text(value)
    return escape(text).replace("\n", "<br/>")


def document_filename(data, source_filename, extension):
    subject = clean_text(data.get("subject")) or source_filename.replace(".json", "")
    class_name = clean_text(data.get("class_category") or data.get("class") or "")
    term = normalize_term(data.get("term"))

    safe_subject = re.sub(r"[^A-Za-z0-9_-]+", "_", subject).strip("_")
    safe_class = re.sub(r"[^A-Za-z0-9_-]+", "_", class_name).strip("_")
    safe_term = f"_{term}" if term else ""

    return f"{safe_subject}_{safe_class}{safe_term}_Teacher_Copy.{extension}"


# ============================================================
# PATH / JSON RESOLUTION
# ============================================================

def resolve_diagram_path(diagram):
    diagram = clean_text(diagram)
    if not diagram: return None

    diagram = diagram.replace("\\", "/")
    if diagram.startswith("/"): diagram = diagram[1:]

    full_path = (Path(BASE_DIR) / diagram).resolve()
    base_path = Path(BASE_DIR).resolve()

    try: full_path.relative_to(base_path)
    except ValueError: return None

    return full_path if full_path.exists() and full_path.is_file() else None


def resolve_exam_json(year, filename, class_category, term=None):
    year = clean_text(year)
    filename = safe_filename(filename)
    class_category = normalize_class(class_category)
    term = normalize_term(term)

    if not valid_year(year): return None, "Invalid year"
    if class_category not in SUPPORTED_CLASSES: return None, "Invalid class"
    if not filename or not filename.lower().endswith(".json"): return None, "Invalid JSON filename"
    if is_term_aware_class(class_category) and not term: return None, "Term is required for JSS"

    path = get_json_path(year, class_category, filename, term)
    if not path or not os.path.isfile(path): return None, "Exam JSON not found"

    data, error = read_json_file(path)
    if error: return None, error
    if not isinstance(data, dict): return None, "Invalid examination JSON"

    return data, None


# ============================================================
# DOCX XML HELPERS
# ============================================================

def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="E1E7ED", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        element = tc_mar.find(qn(f"w:{margin}"))

        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)

        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


# ============================================================
# DOCX DESIGN HELPERS
# ============================================================

def configure_docx(document):
    section = document.sections[0]
    section.top_margin, section.bottom_margin = Inches(0.55), Inches(0.55)
    section.left_margin, section.right_margin = Inches(0.65), Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Arial", Pt(10.5), DOCX_TEXT
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1


def add_docx_centered_line(document, text, size=10, bold=False, color=None, before=0, after=0):
    if not clean_text(text): return

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)

    run = p.add_run(clean_text(text))
    run.bold, run.font.name, run.font.size = bold, "Arial", Pt(size)
    if color: run.font.color.rgb = color


def add_docx_teacher_badge(document):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    cell = table.cell(0, 0)
    shade_cell(cell, "F2F7FC")
    set_cell_border(cell, "D9E8F6", "4")
    set_cell_margins(cell, 55, 100, 55, 100)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run("TEACHER COPY")
    run.bold, run.font.name, run.font.size, run.font.color.rgb = True, "Arial", Pt(8), DOCX_BLUE

    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_docx_header(document, data):
    school = clean_text(data.get("school")) or "EPITOME MODEL ISLAMIC SCHOOLS"
    exam_title = teacher_exam_title(data)
    session = clean_text(data.get("session"))
    subject = clean_text(data.get("subject"))
    class_name = clean_text(data.get("class") or data.get("class_category"))
    time_allowed = format_minutes(data.get("time_allowed_minutes"))

    add_docx_centered_line(document, school, 14, True, DOCX_TEXT, after=1)
    add_docx_centered_line(document, exam_title.upper(), 11.5, True, DOCX_BLUE, after=1)

    if session: add_docx_centered_line(document, f"{session} ACADEMIC SESSION", 9.5, False, DOCX_SOFT, after=5)

    add_docx_teacher_badge(document)

    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    labels = [("SUBJECT", subject), ("CLASS", class_name), ("SESSION", session), ("DURATION", time_allowed)]

    for index, (label, value) in enumerate(labels):
        row, col = divmod(index, 2)
        cell = table.cell(row, col)

        cell.width = Inches(3.4)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        shade_cell(cell, "FBFCFE")
        set_cell_border(cell, "E1E7ED", "4")
        set_cell_margins(cell, 95, 130, 95, 130)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)

        label_run = p.add_run(f"{label}\n")
        label_run.bold, label_run.font.name, label_run.font.size, label_run.font.color.rgb = True, "Arial", Pt(7.5), DOCX_BLUE

        value_run = p.add_run(clean_text(value) or "—")
        value_run.font.name, value_run.font.size, value_run.font.color.rgb = "Arial", Pt(10), DOCX_TEXT

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_docx_section_heading(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    cell = table.cell(0, 0)
    shade_cell(cell, "F2F7FC")
    set_cell_border(cell, "D9E8F6", "4")
    set_cell_margins(cell, 90, 125, 90, 125)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run(clean_text(text))
    run.bold, run.font.name, run.font.size, run.font.color.rgb = True, "Arial", Pt(10.5), DOCX_BLUE

    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_docx_instruction(document, text):
    if not clean_text(text): return

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(7)

    run = p.add_run(clean_text(text))
    run.italic, run.font.name, run.font.size, run.font.color.rgb = True, "Arial", Pt(9), DOCX_SOFT


def add_docx_question(document, number, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(3)
    p.paragraph_format.keep_with_next = True

    number_run = p.add_run(f"{number}. ")
    number_run.bold, number_run.font.name, number_run.font.size = True, "Arial", Pt(10.5)

    text_run = p.add_run(clean_text(text))
    text_run.font.name, text_run.font.size, text_run.font.color.rgb = "Arial", Pt(10.5), DOCX_TEXT


def add_docx_diagram(document, diagram):
    path = resolve_diagram_path(diagram)
    if not path: return

    try:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(5)
        p.add_run().add_picture(str(path), width=Inches(5.6))
    except Exception as exc:
        print(f"[EXAM EXPORT] DOCX diagram error [{path}]: {exc}")


def add_docx_options(document, options, correct_option):
    if not isinstance(options, list): return

    for index, option in enumerate(options):
        letter, option_text = normalize_option_text(option, index)

        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.02)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(1.5)

        run = p.add_run(f"{letter}. {option_text}")
        run.font.name, run.font.size, run.font.color.rgb = "Arial", Pt(10), DOCX_TEXT

        if is_correct_option(option, index, correct_option): run.bold = True


# ============================================================
# DOCX GENERATOR
# ============================================================

def generate_docx(data):
    document = Document()
    configure_docx(document)
    add_docx_header(document, data)

    questions = data.get("questions", [])

    if isinstance(questions, list) and questions:
        add_docx_section_heading(document, "SECTION A — OBJECTIVE")
        add_docx_instruction(document, "Choose the correct answer from the options provided. Correct answers are shown in bold in this teacher copy.")

        for index, question in enumerate(questions, start=1):
            question_id = question.get("id", index)
            add_docx_question(document, question_id, question.get("question"))
            add_docx_diagram(document, question.get("diagram"))
            add_docx_options(document, question.get("options", []), question.get("correctOption"))

    essay = data.get("essay")

    if isinstance(essay, dict):
        essay_questions = essay.get("questions", [])

        if isinstance(essay_questions, list) and essay_questions:
            document.add_page_break()
            add_docx_section_heading(document, clean_text(essay.get("title")) or "SECTION B — THEORY / ESSAY")
            add_docx_instruction(document, clean_text(essay.get("instruction")) or "Answer the required questions. Theory answers are not included in this teacher copy.")

            for index, question in enumerate(essay_questions, start=1):
                question_id = question.get("id", index)
                add_docx_question(document, question_id, question.get("question"))
                add_docx_diagram(document, question.get("diagram"))

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer


# ============================================================
# PDF STYLES
# ============================================================

def pdf_styles():
    styles = getSampleStyleSheet()

    return {
        "school": ParagraphStyle("School", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=14, leading=17, textColor=colors.HexColor(TEXT), alignment=TA_CENTER, spaceAfter=2),
        "exam": ParagraphStyle("Exam", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=11.5, leading=14, textColor=colors.HexColor(BLUE), alignment=TA_CENTER, spaceAfter=1),
        "session": ParagraphStyle("Session", parent=styles["Normal"], fontName=PDF_FONT, fontSize=9.3, leading=12, textColor=colors.HexColor(TEXT_SOFT), alignment=TA_CENTER, spaceAfter=5),
        "badge": ParagraphStyle("Badge", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=7.5, leading=9, textColor=colors.HexColor(BLUE), alignment=TA_CENTER),
        "meta_label": ParagraphStyle("MetaLabel", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=7.2, leading=9, textColor=colors.HexColor(BLUE), spaceAfter=1),
        "meta_value": ParagraphStyle("MetaValue", parent=styles["Normal"], fontName=PDF_FONT, fontSize=9.5, leading=12, textColor=colors.HexColor(TEXT)),
        "section": ParagraphStyle("Section", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=10.8, leading=14, textColor=colors.HexColor(BLUE), backColor=colors.HexColor(BLUE_SOFT), borderColor=colors.HexColor(BLUE_BORDER), borderWidth=.5, borderPadding=(6, 8, 6, 8), spaceBefore=8, spaceAfter=7),
        "question": ParagraphStyle("Question", parent=styles["Normal"], fontName=PDF_FONT, fontSize=10.1, leading=14.5, textColor=colors.HexColor(TEXT), spaceBefore=2, spaceAfter=4),
        "option": ParagraphStyle("Option", parent=styles["Normal"], fontName=PDF_FONT, fontSize=9.6, leading=13, leftIndent=9 * mm, textColor=colors.HexColor("#49586A"), spaceAfter=1.5),
        "option_correct": ParagraphStyle("OptionCorrect", parent=styles["Normal"], fontName=PDF_FONT_BOLD, fontSize=9.6, leading=13, leftIndent=9 * mm, textColor=colors.HexColor(TEXT), spaceAfter=1.5),
        "instruction": ParagraphStyle("Instruction", parent=styles["Normal"], fontName=PDF_FONT, fontSize=8.8, leading=12.5, textColor=colors.HexColor(TEXT_SOFT), spaceAfter=7),
        "footer": ParagraphStyle("Footer", parent=styles["Normal"], fontName=PDF_FONT, fontSize=7.5, leading=9, textColor=colors.HexColor("#8B97A5"), alignment=TA_CENTER),
    }


# ============================================================
# PDF HELPERS
# ============================================================

def add_pdf_teacher_badge(story, styles):
    badge = Table([[Paragraph("TEACHER COPY", styles["badge"])]], colWidths=[28 * mm])
    badge.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(BLUE_SOFT)),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor(BLUE_BORDER)),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))

    badge.hAlign = "CENTER"
    story.extend([badge, Spacer(1, 4 * mm)])


def add_pdf_metadata(story, styles, data):
    session = clean_text(data.get("session"))
    subject = clean_text(data.get("subject"))
    class_name = clean_text(data.get("class") or data.get("class_category"))
    time_allowed = format_minutes(data.get("time_allowed_minutes"))

    def box(label, value):
        return [Paragraph(pdf_safe(label), styles["meta_label"]), Paragraph(pdf_safe(value or "—"), styles["meta_value"])]

    metadata = Table([
        [box("SUBJECT", subject), box("CLASS", class_name)],
        [box("SESSION", session), box("DURATION", time_allowed)],
    ], colWidths=[88 * mm, 88 * mm])

    metadata.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(SURFACE)),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor(BORDER)),
        ("INNERGRID", (0, 0), (-1, -1), .35, colors.HexColor("#E7ECF1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))

    story.extend([metadata, Spacer(1, 4 * mm)])


def get_image_size(path, max_width=150 * mm, max_height=80 * mm):
    try:
        reader = ImageReader(str(path))
        width, height = reader.getSize()

        if not width or not height: return max_width, max_height

        ratio = min(max_width / width, max_height / height, 1)
        return width * ratio, height * ratio
    except Exception:
        return max_width, max_height


def add_pdf_diagram(story, diagram):
    path = resolve_diagram_path(diagram)
    if not path: return

    try:
        width, height = get_image_size(path)
        image = Image(str(path), width=width, height=height)
        image.hAlign = "CENTER"
        story.extend([Spacer(1, 2 * mm), image, Spacer(1, 3.5 * mm)])
    except Exception as exc:
        print(f"[EXAM EXPORT] PDF diagram error [{path}]: {exc}")


def pdf_page_footer(canvas, doc):
    canvas.saveState()

    width, _ = A4
    canvas.setStrokeColor(colors.HexColor("#E4E9EF"))
    canvas.setLineWidth(.4)
    canvas.line(14 * mm, 10 * mm, width - 14 * mm, 10 * mm)

    canvas.setFont(PDF_FONT, 7.5)
    canvas.setFillColor(colors.HexColor("#8995A3"))
    canvas.drawString(14 * mm, 6.2 * mm, "EMIS CBT • Teacher Copy")
    canvas.drawRightString(width - 14 * mm, 6.2 * mm, f"Page {doc.page}")

    canvas.restoreState()


# ============================================================
# PDF GENERATOR
# ============================================================

def generate_pdf(data):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm, topMargin=12 * mm, bottomMargin=15 * mm, title=f"{clean_text(data.get('subject'))} Teacher Copy", author="EPITOME MODEL ISLAMIC SCHOOLS")

    styles = pdf_styles()
    story = []

    school = clean_text(data.get("school")) or "EPITOME MODEL ISLAMIC SCHOOLS"
    exam_title = teacher_exam_title(data)
    session = clean_text(data.get("session"))

    story.append(Paragraph(pdf_safe(school), styles["school"]))
    story.append(Paragraph(pdf_safe(exam_title.upper()), styles["exam"]))

    if session: story.append(Paragraph(pdf_safe(f"{session} ACADEMIC SESSION"), styles["session"]))
    else: story.append(Spacer(1, 2 * mm))

    add_pdf_teacher_badge(story, styles)
    add_pdf_metadata(story, styles, data)

    questions = data.get("questions", [])

    if isinstance(questions, list) and questions:
        story.append(Paragraph("SECTION A — OBJECTIVE", styles["section"]))
        story.append(Paragraph("Choose the correct answer from the options provided. Correct answers are shown in bold in this teacher copy.", styles["instruction"]))

        for index, question in enumerate(questions, start=1):
            qid = question.get("id", index)

            story.append(Paragraph(f"<b>{pdf_safe(qid)}.</b> {pdf_safe(question.get('question'))}", styles["question"]))
            add_pdf_diagram(story, question.get("diagram"))

            options = question.get("options", [])
            correct_option = question.get("correctOption")

            if isinstance(options, list):
                for option_index, option in enumerate(options):
                    letter, option_text = normalize_option_text(option, option_index)
                    style = styles["option_correct"] if is_correct_option(option, option_index, correct_option) else styles["option"]
                    story.append(Paragraph(f"{pdf_safe(letter)}. {pdf_safe(option_text)}", style))

            story.append(Spacer(1, 2.8 * mm))

    essay = data.get("essay")

    if isinstance(essay, dict):
        essay_questions = essay.get("questions", [])

        if isinstance(essay_questions, list) and essay_questions:
            story.append(PageBreak())
            story.append(Paragraph(pdf_safe(clean_text(essay.get("title")) or "SECTION B — THEORY / ESSAY"), styles["section"]))

            instruction = clean_text(essay.get("instruction")) or "Answer the required questions. Theory answers are not included in this teacher copy."
            story.append(Paragraph(pdf_safe(instruction), styles["instruction"]))

            for index, question in enumerate(essay_questions, start=1):
                qid = question.get("id", index)

                story.append(Paragraph(f"<b>{pdf_safe(qid)}.</b> {pdf_safe(question.get('question'))}", styles["question"]))
                add_pdf_diagram(story, question.get("diagram"))
                story.append(Spacer(1, 3 * mm))

    doc.build(story, onFirstPage=pdf_page_footer, onLaterPages=pdf_page_footer)

    buffer.seek(0)
    return buffer



# ============================================================
# PDF PREVIEW — JSON / BASE64
# Prevents browser download managers from intercepting preview.
# ============================================================

@exam_document_export_bp.route("/uploads/<year>/<filename>/teacher-copy/preview", methods=["GET"])
def teacher_copy_preview(year, filename):
    class_category = request.args.get("class") or request.args.get("class_category") or ""
    term = request.args.get("term")

    data, error = resolve_exam_json(year, filename, class_category, term)

    if error:
        return jsonify({"success": False, "error": error}), 404

    try:
        output = generate_pdf(data)
        pdf_bytes = output.getvalue()

        if not pdf_bytes:
            return jsonify({"success": False, "error": "Generated PDF is empty"}), 500

        return jsonify({
            "success": True,
            "mime_type": "application/pdf",
            "filename": document_filename(data, filename, "pdf"),
            "pdf_base64": base64.b64encode(pdf_bytes).decode("ascii")
        })

    except Exception as exc:
        print(f"[EXAM EXPORT] Preview generation failed [{year}/{filename}]: {exc}")

        return jsonify({
            "success": False,
            "error": f"Unable to generate teacher-copy preview: {exc}"
        }), 500



# ============================================================
# ROUTE
# ============================================================

@exam_document_export_bp.route("/uploads/<year>/<filename>/teacher-copy", methods=["GET"])
def teacher_copy(year, filename):
    class_category = request.args.get("class") or request.args.get("class_category") or ""
    term = request.args.get("term")
    file_format = clean_text(request.args.get("format")).lower()
    disposition = clean_text(request.args.get("disposition")).lower()

    if file_format not in {"pdf", "docx"}:
        return jsonify({"success": False, "error": "Format must be pdf or docx"}), 400

    data, error = resolve_exam_json(year, filename, class_category, term)

    if error:
        return jsonify({"success": False, "error": error}), 404

    try:
        if file_format == "docx":
            output = generate_docx(data)

            return send_file(
                output,
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                as_attachment=True,
                download_name=document_filename(data, filename, "docx"),
            )

        output = generate_pdf(data)
        as_attachment = disposition != "inline"

        return send_file(
            output,
            mimetype="application/pdf",
            as_attachment=as_attachment,
            download_name=document_filename(data, filename, "pdf"),
        )

    except Exception as exc:
        print(f"[EXAM EXPORT] Generation failed [{year}/{filename}]: {exc}")
        return jsonify({"success": False, "error": f"Unable to generate teacher copy: {exc}"}), 500